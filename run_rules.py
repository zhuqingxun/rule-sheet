# /// script
# requires-python = ">=3.11"
# dependencies = ["pandas>=2.2", "openpyxl>=3.1", "pyyaml>=6.0"]
# ///
"""最小规则解释器: 读 Excel 决策表或 YAML 规则, 比较两份 Excel, 按规则更新左表。

用法:
    uv run --script run_rules.py rules/规则_决策表.xlsx
    uv run --script run_rules.py rules/rules.yaml

选项:
    --ignore-tests   测试用例有失败也继续处理真实数据
    --dry-run        只打印结果, 不写输出文件

单元格条件语法 (FEEL 一元测试子集, 空值参与任何比较都算不满足):
    -  或留空          不关心
    活跃               等于
    > 0.05  != 休眠    比较, 右侧可以是另一个变量名, 如 != 主表.电话
    [0.02..0.05]       闭区间; 用 ( ) 表示开区间
    VIP, 金卡          枚举, 任一满足
    非空 / 为空        判空
    not(...)           取反
输出格:
    -  或留空          不改
    待复核             字面值
    本月.电话          复制变量的值
    =表达式            用公式算
YAML 条件组合:
    当: 下的列表默认全部满足 (AND)
    任一满足: [...]    其中之一成立即可 (OR), 可以和 全部满足 互相嵌套
    不满足: ...        取反 (NOT)
"""
from __future__ import annotations

import argparse
import numbers
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

import pandas as pd
import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
Mask = pd.Series
Namespace = dict[str, Any]
Condition = Callable[[Namespace], Any]

TRUE_WORDS = {"是", "true", "True", "TRUE", "yes"}
FALSE_WORDS = {"否", "false", "False", "FALSE", "no"}
BLANK_WORDS = {"", "-", "—", "任意"}
EXIST_VAR = "存在"          # 右表是否匹配到, 实际变量名是 <右表别名>.存在
BASE_DATE_VAR = "基准日期"


# ---------------------------------------------------------------------------
# 数据结构
# ---------------------------------------------------------------------------
@dataclass
class Rule:
    rule_id: str
    priority: float
    conditions: list[Condition]
    outputs: dict[str, Any]            # 输出列 -> 原始单元格值 / YAML 值
    desc: str = ""


@dataclass
class TestCase:
    name: str
    inputs: dict[str, Any]
    expected: dict[str, Any]           # 期望列 -> 值; None 表示「不变」


@dataclass
class RuleSet:
    match: dict[str, Any]
    derived: list[tuple[str, str]]     # (变量名, 表达式), 按定义顺序求值
    hit_policy: str                    # 优先级 / 唯一
    rules: list[Rule]
    outputs: list[str]
    tests: list[TestCase] = field(default_factory=list)
    source: str = ""
    variables: list[str] = field(default_factory=list)   # 规则引用到的变量, 测试用例没给的视为空值

    @property
    def left_alias(self) -> str:
        return str(self.match.get("左表别名", "主表"))

    @property
    def right_alias(self) -> str:
        return str(self.match.get("右表别名", "本月"))


# ---------------------------------------------------------------------------
# 表达式求值: 把中文变量名换成安全标识符, 再交给 Python eval
# ---------------------------------------------------------------------------
def _days_between(later: Any, earlier: Any) -> Any:
    delta = pd.to_datetime(later) - pd.to_datetime(earlier)
    return delta.dt.days if isinstance(delta, pd.Series) else delta.days


def _notna(x: Any) -> Any:
    return x.notna() if isinstance(x, pd.Series) else pd.notna(x)


def _isna(x: Any) -> Any:
    return x.isna() if isinstance(x, pd.Series) else pd.isna(x)


def _isin(x: Any, *values: Any) -> Any:
    return x.isin(values) if isinstance(x, pd.Series) else x in values


FUNCS: dict[str, Any] = {
    "abs": abs, "round": round, "min": min, "max": max,
    "天数差": _days_between, "days_between": _days_between,
    "非空": _notna, "为空": _isna, "属于": _isin,
    "是": True, "否": False, "true": True, "false": False,
}

_IDENTS: dict[str, str] = {}


def _ident(name: str) -> str:
    if name not in _IDENTS:
        _IDENTS[name] = f"_v{len(_IDENTS)}"
    return _IDENTS[name]


def eval_expr(expr: str, ns: Namespace) -> Any:
    text = expr
    local: dict[str, Any] = {}
    for name in sorted(ns, key=len, reverse=True):          # 长名优先, 防止 金额 吃掉 主表.金额
        pattern = r"(?<![\w.])" + re.escape(name) + r"(?![\w.])"
        text, n = re.subn(pattern, _ident(name), text)
        if n:
            local[_ident(name)] = ns[name]
    local.update(FUNCS)
    try:
        return eval(text, {"__builtins__": {}}, local)      # noqa: S307  样例用途, 规则文件视为可信
    except Exception as exc:
        raise ValueError(f"表达式求值失败: {expr!r} ({exc})") from exc


# ---------------------------------------------------------------------------
# 单元格条件 (FEEL 一元测试子集)
# ---------------------------------------------------------------------------
_RANGE_RE = re.compile(r"^([\[\(])\s*(.+?)\s*\.\.\s*(.+?)\s*([\]\)])$")
_CMP_RE = re.compile(r"^(>=|<=|!=|<>|==|=|>|<)\s*(.+)$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def split_top_level(s: str, sep: str = ",") -> list[str]:
    parts: list[str] = []
    buf: list[str] = []
    depth = 0
    quote: str | None = None
    for ch in s:
        if quote:
            buf.append(ch)
            if ch == quote:
                quote = None
            continue
        if ch in "\"'":
            quote = ch
            buf.append(ch)
            continue
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth -= 1
        if ch == sep and depth == 0:
            parts.append("".join(buf))
            buf = []
        else:
            buf.append(ch)
    parts.append("".join(buf))
    return [p.strip() for p in parts if p.strip()]


def literal(token: str, ns: Namespace, var: pd.Series | None = None) -> Any:
    t = token.strip()
    if len(t) >= 2 and t[0] == t[-1] and t[0] in "\"'":
        return t[1:-1]
    if t in ns:
        return ns[t]
    if var is not None and var.dtype == bool:
        if t in TRUE_WORDS:
            return True
        if t in FALSE_WORDS:
            return False
    try:
        f = float(t)
        return int(f) if f.is_integer() and "." not in t else f
    except ValueError:
        pass
    if _DATE_RE.match(t):
        return pd.Timestamp(t)
    return t


def _both_present(var: pd.Series, rhs: Any) -> Mask:
    ok = var.notna()
    if isinstance(rhs, pd.Series):
        ok = ok & rhs.notna()
    return ok


def unary_test(cell: Any, var: pd.Series, ns: Namespace) -> Mask:
    all_true = pd.Series(True, index=var.index)
    if cell is None:
        return all_true
    if isinstance(cell, bool) or isinstance(cell, numbers.Real):
        return (var == cell) & var.notna()
    if isinstance(cell, (datetime, date)):
        return (pd.to_datetime(var) == pd.Timestamp(cell)) & var.notna()

    s = str(cell).strip()
    if s in BLANK_WORDS:
        return all_true
    if s.startswith("not(") and s.endswith(")"):
        return ~unary_test(s[4:-1], var, ns)
    parts = split_top_level(s)
    if len(parts) > 1:                                        # 枚举: 任一满足
        m = pd.Series(False, index=var.index)
        for p in parts:
            m = m | unary_test(p, var, ns)
        return m
    if s == "非空":
        return var.notna()
    if s == "为空":
        return var.isna()
    if m := _RANGE_RE.match(s):
        lo, hi = literal(m.group(2), ns, var), literal(m.group(3), ns, var)
        lo_ok = (var >= lo) if m.group(1) == "[" else (var > lo)
        hi_ok = (var <= hi) if m.group(4) == "]" else (var < hi)
        return lo_ok & hi_ok & var.notna()
    if m := _CMP_RE.match(s):
        op, rhs = m.group(1), literal(m.group(2), ns, var)
        ok = _both_present(var, rhs)
        if op in ("=", "=="):
            return (var == rhs) & ok
        if op in ("!=", "<>"):
            return (var != rhs) & ok
        if op == ">":
            return (var > rhs) & ok
        if op == ">=":
            return (var >= rhs) & ok
        if op == "<":
            return (var < rhs) & ok
        return (var <= rhs) & ok
    rhs = literal(s, ns, var)
    return (var == rhs) & _both_present(var, rhs)


def cell_condition(var_name: str, cell: Any) -> Condition:
    def cond(ns: Namespace) -> Mask:
        if var_name not in ns:
            raise KeyError(f"规则引用了未定义的变量「{var_name}」")
        return unary_test(cell, ns[var_name], ns)
    return cond


def expr_condition(expr: str) -> Condition:
    return lambda ns: eval_expr(expr, ns)


def _to_mask(res: Any, ns: Namespace) -> Mask:
    index = next(iter(ns.values())).index
    if not isinstance(res, pd.Series):
        res = pd.Series(bool(res), index=index)
    return res.fillna(False).astype(bool)


def all_of(conds: list[Condition]) -> Condition:
    def cond(ns: Namespace) -> Mask:
        m = _to_mask(True, ns)
        for c in conds:
            m = m & _to_mask(c(ns), ns)
        return m
    return cond


def any_of(conds: list[Condition]) -> Condition:
    def cond(ns: Namespace) -> Mask:
        m = _to_mask(False, ns)
        for c in conds:
            m = m | _to_mask(c(ns), ns)
        return m
    return cond


def not_of(inner: Condition) -> Condition:
    return lambda ns: ~_to_mask(inner(ns), ns)


COMBINATORS: dict[str, Callable[[list[Condition]], Condition]] = {"全部满足": all_of, "任一满足": any_of}


def build_condition(node: Any) -> Condition:
    """YAML 条件树: 字符串 = 一条表达式; 列表 = 全部满足; 字典 = 全部满足 / 任一满足 / 不满足, 可嵌套。"""
    if isinstance(node, str):
        return expr_condition(node)
    if isinstance(node, list):
        return all_of([build_condition(n) for n in node])
    if isinstance(node, dict) and len(node) == 1:
        key, val = next(iter(node.items()))
        if key in COMBINATORS:
            items = val if isinstance(val, list) else [val]
            return COMBINATORS[key]([build_condition(n) for n in items])
        if key == "不满足":
            return not_of(build_condition(val))
    raise ValueError(f"无法识别的条件写法: {node!r} (只认 表达式 / 全部满足 / 任一满足 / 不满足)")


def _strings(node: Any) -> list[str]:
    if isinstance(node, str):
        return [node]
    if isinstance(node, list):
        return [s for n in node for s in _strings(n)]
    if isinstance(node, dict):
        return [s for v in node.values() for s in _strings(v)]
    return []


def resolve_output(spec: Any, ns: Namespace) -> Any:
    """返回 None 表示不改。"""
    if spec is None:
        return None
    if isinstance(spec, str):
        s = spec.strip()
        if s in BLANK_WORDS:
            return None
        if s.startswith("="):
            return eval_expr(s[1:], ns)
        if s in ns:
            return ns[s]
        if len(s) >= 2 and s[0] == s[-1] and s[0] in "\"'":
            return s[1:-1]
        return s
    return spec


# ---------------------------------------------------------------------------
# 规则加载: Excel 决策表
# ---------------------------------------------------------------------------
def _rows(ws: Any) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _read_kv(ws: Any) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for row in _rows(ws)[1:]:
        if row and row[0] not in (None, "") and len(row) > 1 and row[1] not in (None, ""):
            out[str(row[0]).strip()] = row[1]
    return out


def _parse_header(grid: list[list[Any]]) -> tuple[list[tuple[int, str, str]], int]:
    """返回 (列索引, 列类型, 变量名) 列表; 表头固定在第 2、3 行。"""
    kinds, names = grid[1], grid[2]
    cols = [(i, str(k).strip(), str(n).strip())
            for i, (k, n) in enumerate(zip(kinds, names)) if k not in (None, "") and n not in (None, "")]
    return cols, 3


def load_excel_rules(path: Path) -> RuleSet:
    wb = load_workbook(path, data_only=True)
    match = _read_kv(wb["匹配设置"])
    derived = [(str(r[0]).strip(), str(r[1]).strip()) for r in _rows(wb["变量定义"])[1:]
               if r and r[0] not in (None, "") and len(r) > 1 and r[1] not in (None, "")]

    grid = _rows(wb["决策表"])
    hit_policy = str(grid[0][1] or "优先级").strip()
    cols, body_start = _parse_header(grid)
    id_col = next(i for i, k, _ in cols if k == "规则ID")
    pri_col = next((i for i, k, _ in cols if k == "优先级"), None)
    desc_col = next((i for i, k, _ in cols if k == "说明"), None)
    inputs = [(i, n) for i, k, n in cols if k == "输入"]
    outputs = [(i, n) for i, k, n in cols if k == "输出"]

    rules: list[Rule] = []
    for order, row in enumerate(grid[body_start:]):
        if len(row) <= id_col or row[id_col] in (None, ""):
            continue
        pri = float(row[pri_col]) if pri_col is not None and row[pri_col] not in (None, "") else 1000.0 + order
        rules.append(Rule(
            rule_id=str(row[id_col]).strip(),
            priority=pri,
            conditions=[cell_condition(name, row[i]) for i, name in inputs],
            outputs={name: row[i] for i, name in outputs},
            desc=str(row[desc_col] or "") if desc_col is not None else "",
        ))

    tests: list[TestCase] = []
    if "测试" in wb.sheetnames:
        tgrid = _rows(wb["测试"])
        tcols, tstart = _parse_header(tgrid)
        case_col = next(i for i, k, _ in tcols if k == "用例")
        t_inputs = [(i, n) for i, k, n in tcols if k == "输入"]
        t_expect = [(i, n) for i, k, n in tcols if k == "期望"]
        for row in tgrid[tstart:]:
            if len(row) <= case_col or row[case_col] in (None, ""):
                continue
            tests.append(TestCase(
                name=str(row[case_col]),
                inputs={n: row[i] for i, n in t_inputs},
                expected={n: row[i] for i, n in t_expect},
            ))

    return RuleSet(match, derived, hit_policy, rules, [n for _, n in outputs], tests,
                   source=path.name, variables=[n for _, n in inputs])


# ---------------------------------------------------------------------------
# 规则加载: YAML
# ---------------------------------------------------------------------------
def load_yaml_rules(path: Path) -> RuleSet:
    doc = yaml.safe_load(path.read_text(encoding="utf-8"))
    rules: list[Rule] = []
    outputs: list[str] = []
    for order, r in enumerate(doc.get("规则", [])):
        conds = r.get("当", [])
        rules.append(Rule(
            rule_id=str(r.get("编号", f"R{order + 1}")),
            priority=float(r.get("优先级", 1000 + order)),
            conditions=[build_condition(conds)] if conds else [],
            outputs=dict(r.get("则", {})),
            desc=str(r.get("名称", "")),
        ))
        for col in r.get("则", {}):
            if col not in outputs:
                outputs.append(col)
    tests = [TestCase(str(t.get("用例", f"T{i + 1}")), dict(t.get("输入", {})), dict(t.get("期望", {})))
             for i, t in enumerate(doc.get("测试", []))]
    derived = [(str(k), str(v)) for k, v in (doc.get("变量") or {}).items()]
    match = dict(doc.get("匹配", {}))
    la, ra = match.get("左表别名", "主表"), match.get("右表别名", "本月")
    texts = [s for r in doc.get("规则", []) for s in _strings(r.get("当"))]
    texts += [str(v) for r in doc.get("规则", []) for v in (r.get("则") or {}).values()]
    texts += [v for _, v in derived]
    name_re = re.compile(rf"(?<![\w.])(?:{re.escape(str(la))}|{re.escape(str(ra))})\.\w+")
    variables = sorted({m.group(0) for t in texts for m in name_re.finditer(t)} | {k for k, _ in derived})
    return RuleSet(match, derived, str(doc.get("命中策略", "优先级")),
                   rules, outputs, tests, source=path.name, variables=variables)


def load_rules(path: Path) -> RuleSet:
    if path.suffix.lower() in (".yaml", ".yml"):
        return load_yaml_rules(path)
    return load_excel_rules(path)


# ---------------------------------------------------------------------------
# 命名空间: 把两张表合并成 主表.* / 本月.* 变量
# ---------------------------------------------------------------------------
def build_namespace(rs: RuleSet, left: pd.DataFrame, right: pd.DataFrame, keys: list[str]) -> Namespace:
    dup = right.duplicated(subset=keys, keep="first")
    if dup.any():
        policy = str(rs.match.get("右表重复键", "取第一条"))
        if policy != "取第一条":
            raise ValueError(f"右表有 {int(dup.sum())} 行重复键, 而「右表重复键」设置为 {policy}")
        print(f"提示: 右表有 {int(dup.sum())} 行重复键, 按「取第一条」处理")
        right = right[~dup]

    merged = left.merge(right, on=keys, how="left", suffixes=("", "__r"), indicator=True)
    merged.index = left.index
    la, ra = rs.left_alias, rs.right_alias
    ns: Namespace = {}
    for col in left.columns:
        ns[f"{la}.{col}"] = merged[col]
    for col in right.columns:
        if col in keys:
            continue
        src = f"{col}__r" if f"{col}__r" in merged.columns else col
        ns[f"{ra}.{col}"] = merged[src]
    ns[f"{ra}.{EXIST_VAR}"] = merged["_merge"].eq("both")
    base = rs.match.get(BASE_DATE_VAR)
    ns[BASE_DATE_VAR] = pd.Timestamp(base) if base else pd.Timestamp(date.today())
    return ns


def add_derived(rs: RuleSet, ns: Namespace, index: pd.Index, *, lenient: bool) -> None:
    for name, expr in rs.derived:
        if name in ns:                                  # 测试用例直接给了值, 不再推导
            continue
        try:
            val = eval_expr(expr, ns)
        except ValueError:
            if not lenient:
                raise
            val = float("nan")
        if not isinstance(val, pd.Series):
            val = pd.Series([val] * len(index), index=index)
        ns[name] = val


# ---------------------------------------------------------------------------
# 规则执行
# ---------------------------------------------------------------------------
def rule_mask(rule: Rule, ns: Namespace, index: pd.Index) -> Mask:
    m = pd.Series(True, index=index)
    for cond in rule.conditions:
        try:
            res = cond(ns)
        except KeyError as exc:
            raise KeyError(f"规则 {rule.rule_id}: {exc.args[0]}") from exc
        if not isinstance(res, pd.Series):
            res = pd.Series(bool(res), index=index)
        m = m & res.fillna(False).astype(bool)
    return m


def apply_rules(rs: RuleSet, ns: Namespace, originals: dict[str, pd.Series],
                labels: pd.Series | None = None) -> tuple[dict[str, pd.Series], pd.Series]:
    index = next(iter(ns.values())).index
    masks = {r.rule_id: rule_mask(r, ns, index) for r in rs.rules}
    ordered = sorted(rs.rules, key=lambda r: r.priority)

    if rs.hit_policy == "唯一":
        hits = sum(masks[r.rule_id].astype(int) for r in rs.rules)
        bad = hits[hits > 1]
        if not bad.empty:
            detail = {str(labels.loc[i]) if labels is not None else f"第{int(i) + 2}行":
                      [r.rule_id for r in ordered if masks[r.rule_id].loc[i]] for i in bad.index}
            raise ValueError(f"命中策略为「唯一」, 但以下行命中多条规则: {detail}")

    results: dict[str, pd.Series] = {}
    for col in rs.outputs:
        cur = originals.get(col)
        if cur is None:
            cur = pd.Series([None] * len(index), index=index, dtype=object)
        cur = cur.astype(object)
        for rule in reversed(ordered):                 # 优先级数字大的先写, 小的最后写并覆盖
            val = resolve_output(rule.outputs.get(col), ns)
            if val is None:
                continue
            if isinstance(val, pd.Series):
                val = val.astype(object)
            cur = cur.mask(masks[rule.rule_id], val)
        results[col] = cur

    hit_ids = pd.Series(["; ".join(r.rule_id for r in ordered if masks[r.rule_id].loc[i]) for i in index], index=index)
    return results, hit_ids


# ---------------------------------------------------------------------------
# 比较与格式化
# ---------------------------------------------------------------------------
def _is_na(x: Any) -> bool:
    if x is None:
        return True
    try:
        return bool(pd.isna(x))
    except (TypeError, ValueError):
        return False


def same(a: Any, b: Any) -> bool:
    if _is_na(a) or _is_na(b):
        return _is_na(a) and _is_na(b)
    if isinstance(a, bool) or isinstance(b, bool):
        return bool(a) == bool(b)
    if isinstance(a, numbers.Real) and isinstance(b, numbers.Real):
        return abs(float(a) - float(b)) < 1e-9
    if isinstance(a, (pd.Timestamp, datetime, date)) or isinstance(b, (pd.Timestamp, datetime, date)):
        try:
            return pd.Timestamp(a) == pd.Timestamp(b)
        except (TypeError, ValueError):
            return False
    return str(a) == str(b)


def fmt(x: Any) -> str:
    if _is_na(x):
        return "空"
    if isinstance(x, (pd.Timestamp, datetime, date)):
        return pd.Timestamp(x).strftime("%Y-%m-%d")
    if isinstance(x, bool):
        return "是" if x else "否"
    if isinstance(x, numbers.Real) and float(x).is_integer():
        return str(int(x))
    return str(x)


# ---------------------------------------------------------------------------
# 测试用例
# ---------------------------------------------------------------------------
def run_tests(rs: RuleSet) -> tuple[int, int]:
    exist_var = f"{rs.right_alias}.{EXIST_VAR}"
    passed = 0
    for tc in rs.tests:
        index = pd.RangeIndex(1)
        ns: Namespace = {}
        for name, val in tc.inputs.items():
            if name == exist_var and isinstance(val, str):
                val = val.strip() in TRUE_WORDS
            if isinstance(val, str) and val.strip() == "":
                val = None
            ns[name] = pd.Series([val], index=index, dtype=object if isinstance(val, str) or val is None else None)
        ns.setdefault(BASE_DATE_VAR, pd.Timestamp(rs.match.get(BASE_DATE_VAR) or date.today()))
        for name in rs.variables:                      # 用例没提到的变量 = 空值
            if name not in ns and not any(name == d for d, _ in rs.derived):
                ns[name] = pd.Series([None], index=index, dtype=object)
        add_derived(rs, ns, index, lenient=True)
        originals = {col: ns.get(f"{rs.left_alias}.{col}", pd.Series([None], index=index, dtype=object)) for col in rs.outputs}
        try:
            results, hit_ids = apply_rules(rs, ns, originals)
        except Exception as exc:                       # noqa: BLE001
            print(f"  [错误] {tc.name}: {exc}")
            continue
        problems = []
        for col in rs.outputs:
            exp = tc.expected.get(col)
            got = results[col].iloc[0]
            want = originals[col].iloc[0] if exp is None or (isinstance(exp, str) and exp.strip() == "") else exp
            if not same(got, want):
                problems.append(f"{col}: 期望 {fmt(want)}, 实际 {fmt(got)}")
        if problems:
            print(f"  [失败] {tc.name} (命中 {hit_ids.iloc[0] or '无'}): " + "; ".join(problems))
        else:
            passed += 1
            print(f"  [通过] {tc.name} (命中 {hit_ids.iloc[0] or '无'})")
    return passed, len(rs.tests)


# ---------------------------------------------------------------------------
# 输出
# ---------------------------------------------------------------------------
def write_xlsx_safely(df: pd.DataFrame, path: Path) -> Path:
    """目标被 Excel 占用时自动换 _v2/_v3 后缀, 不中断。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    candidates = [path] + [path.with_name(f"{path.stem}_v{i}{path.suffix}") for i in range(2, 10)]
    for p in candidates:
        try:
            with pd.ExcelWriter(p, engine="openpyxl") as xw:
                df.to_excel(xw, index=False)
                ws = xw.sheets[next(iter(xw.sheets))]
                for i, col in enumerate(df.columns, start=1):
                    width = max([len(str(col))] + [len(str(v)) for v in df[col].head(50)]) * 1.6 + 2
                    ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 60)
                    if pd.api.types.is_datetime64_any_dtype(df[col]):
                        for r in range(2, len(df) + 2):
                            ws.cell(row=r, column=i).number_format = "yyyy-mm-dd"
                ws.freeze_panes = "B2"
            return p
        except PermissionError:
            continue
    raise PermissionError(f"无法写入 {path} 及其备选文件名, 请关闭占用的 Excel 窗口")


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")           # type: ignore[attr-defined]
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("rules", help="规则文件: .xlsx 决策表 或 .yaml")
    ap.add_argument("--ignore-tests", action="store_true", help="测试失败也继续")
    ap.add_argument("--dry-run", action="store_true", help="只打印, 不写文件")
    args = ap.parse_args()

    rules_path = Path(args.rules)
    if not rules_path.is_absolute():
        rules_path = (Path.cwd() / rules_path) if (Path.cwd() / rules_path).exists() else ROOT / rules_path
    rs = load_rules(rules_path)
    print(f"规则文件: {rs.source}   规则数: {len(rs.rules)}   命中策略: {rs.hit_policy}   输出列: {', '.join(rs.outputs)}")

    if rs.tests:
        print(f"\n== 先跑测试用例 ({len(rs.tests)} 条) ==")
        passed, total = run_tests(rs)
        print(f"测试结果: {passed}/{total} 通过")
        if passed < total and not args.ignore_tests:
            print("有测试未通过, 停止处理真实数据 (加 --ignore-tests 可强制继续)")
            return 1

    left = pd.read_excel(ROOT / str(rs.match["左表"]))
    right = pd.read_excel(ROOT / str(rs.match["右表"]))
    keys = [k.strip() for k in str(rs.match["匹配键"]).split(",")]
    print(f"\n== 处理真实数据 ==\n左表 {len(left)} 行, 右表 {len(right)} 行, 匹配键 {keys}")

    ns = build_namespace(rs, left, right, keys)
    add_derived(rs, ns, left.index, lenient=False)
    originals = {col: left[col] for col in rs.outputs if col in left.columns}
    results, hit_ids = apply_rules(rs, ns, originals, labels=left[keys[0]])

    out = left.copy()
    details: list[str] = []
    for i in left.index:
        parts = []
        for col in rs.outputs:
            old = left[col].loc[i] if col in left.columns else None
            new = results[col].loc[i]
            if not same(old, new):
                parts.append(f"{col}: {fmt(old)} -> {fmt(new)}")
        details.append("; ".join(parts))
    for col in rs.outputs:
        out[col] = results[col]
    out["命中规则"] = hit_ids
    out["变更明细"] = details

    show_cols = keys + ["命中规则", "变更明细"]
    changed = out[out["变更明细"] != ""]
    print(f"\n{len(changed)} 行发生变更:")
    print(changed[show_cols].to_string(index=False) if not changed.empty else "(无)")

    if args.dry_run:
        print("\n--dry-run: 未写文件")
        return 0
    out_path = ROOT / str(rs.match.get("输出", "output/结果.xlsx"))
    written = write_xlsx_safely(out, out_path)
    print(f"\n已写出: {written}")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (ValueError, KeyError, FileNotFoundError) as exc:
        print(f"\n错误: {exc.args[0] if exc.args else exc}")
        sys.exit(2)
