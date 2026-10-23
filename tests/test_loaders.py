"""两种前端的加载: 结构、变量扫描、条件组合。"""

from datetime import date, datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from rule_sheet.conditions import build_condition
from rule_sheet.loaders import load_rules, load_yaml_rules

from .conftest import EXCEL_RULES, YAML_RULES


def test_excel_structure():
    rs = load_rules(EXCEL_RULES)
    assert [r.rule_id for r in rs.rules] == ["R1", "R2", "R3", "R4", "R5"]  # R5 占两行 (或 分支), 仍是一条
    assert rs.outputs == ["状态", "电话", "复核标记", "备注", "提醒"]
    assert rs.hit_policy == "优先级"
    assert len(rs.tests) == 8 and rs.tests[0].name == "T1"
    # 表头输入列在前, 之后是派生变量表达式里引用到的列 (用例没给时置空, 派生变量才能算出来)
    assert rs.variables == [
        "客户.状态", "客户.等级", "偏差率", "本月.存在", "距上次活跃天数", "本月.电话", "客户.电话",
        "客户.上次活跃", "客户.金额", "本月.金额",
    ]  # fmt: skip
    assert [d for d, _ in rs.derived] == ["偏差率", "距上次活跃天数"]
    assert isinstance(rs.settings["基准日期"], (date, datetime))
    assert [(s.name, s.primary, s.keys, s.path) for s in rs.inputs] == [
        ("客户", True, [], "客户主表.xlsx"),
        ("本月", False, [("客户编号", "客户编号")], "本月流水.xlsx"),
    ]


def test_yaml_structure_and_variable_scan():
    rs = load_rules(YAML_RULES)
    assert [r.rule_id for r in rs.rules] == ["R1", "R2", "R3", "R4", "R5"]
    assert rs.outputs == ["复核标记", "备注", "状态", "电话", "提醒"]
    assert len(rs.tests) == 8
    # 嵌套在 任一满足 里的变量也要被扫到, 测试用例没给时才会被置空而不是报未定义
    assert {"本月.电话", "客户.电话", "本月.存在", "偏差率", "距上次活跃天数"} <= set(rs.variables)


def test_nested_combinators_evaluate(ns_two_rows):
    cond = build_condition(
        [
            "本月.存在 == 是",
            {"任一满足": ["为空(本月.电话)", "为空(主表.电话)"]},
        ]
    )
    assert cond(ns_two_rows).tolist() == [False, False, True]
    cond = build_condition({"不满足": {"任一满足": ['主表.状态 == "活跃"', '主表.状态 == "休眠"']}})
    assert cond(ns_two_rows).tolist() == [False, False, True]
    cond = build_condition({"全部满足": "本月.存在 == 是"})  # 单个值也接受
    assert cond(ns_two_rows).tolist() == [True, False, True]


def test_unknown_combinator_rejected():
    with pytest.raises(ValueError, match="无法识别"):
        build_condition({"或者": ["a == 1"]})
    with pytest.raises(ValueError, match="无法识别"):
        build_condition(42)


def test_yaml_minimal_document(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text(
        "规则:\n  - 当:\n      - x > 1\n    则:\n      y: 2\n",
        encoding="utf-8",
    )
    rs = load_yaml_rules(p)
    assert rs.rules[0].rule_id == "R1" and rs.rules[0].priority == 1000
    assert rs.outputs == ["y"] and rs.hit_policy == "优先级"
    mask = rs.rules[0].conditions[0]({"x": pd.Series([0, 2])})
    assert mask.tolist() == [False, True]


def test_yaml_hyphen_names_and_unspaced_subtraction(tmp_path):
    """issue #1 的现场: 字段名含 -, 用例没给该变量时要被扫描到并置空, 而不是被截成 上期.类别 留下裸的 会议。
    同时无空格的减法不能被扫描器整体当成一个变量名吞掉。"""
    p = tmp_path / "r.yaml"
    p.write_text(
        "匹配: {左表别名: 本期, 右表别名: 上期}\n"
        "变量:\n  差额: 本期.金额-上期.金额\n"
        "规则:\n"
        '  - 编号: S1\n    当:\n      - 上期.类别-会议 == "变动通报"\n    则:\n      状态-会议: 02-进展更新\n'
        "  - 编号: S2\n    当:\n      - 差额 > 0\n    则:\n      备注: 涨了\n"
        "测试:\n"
        "  - 用例: 用例没给 上期.类别-会议, 视为空, 不命中\n"
        "    输入: {本期.金额: 10, 上期.金额: 4}\n"
        "    期望: {备注: 涨了}\n"
        "  - 用例: 给了就命中\n"
        "    输入: {上期.类别-会议: 变动通报, 本期.金额: 1, 上期.金额: 4}\n"
        "    期望: {状态-会议: 02-进展更新}\n",
        encoding="utf-8",
    )
    rs = load_yaml_rules(p)
    assert rs.variables == ["上期.类别-会议", "上期.金额", "差额", "本期.金额"]
    from rule_sheet.engine import run_tests

    outcomes = run_tests(rs)
    assert [(o.passed, o.hit, o.error) for o in outcomes] == [(True, "S2", ""), (True, "S1", "")]


def test_yaml_backtick_name_end_to_end(tmp_path):
    """反引号名字从 YAML 文件走完整路径: 以反引号开头的表达式整行要加引号 (YAML 的保留字符),
    扫描到的名字在用例没给时置空、给了就参与求值, 两种情况都不报错。"""
    p = tmp_path / "r.yaml"
    p.write_text(
        "匹配: {左表别名: 本期, 右表别名: 上期}\n"
        "规则:\n"
        '  - 编号: R1\n    当:\n      - "`本期.金额(万元)` > 100"\n    则:\n      标记: 大额\n'
        "测试:\n"
        "  - 用例: 没给, 视为空, 不命中\n    输入: {}\n    期望: {}\n"
        "  - 用例: 给了就命中\n    输入: {本期.金额(万元): 150}\n    期望: {标记: 大额}\n",
        encoding="utf-8",
    )
    rs = load_yaml_rules(p)
    assert rs.variables == ["本期.金额(万元)"]
    from rule_sheet.engine import run_tests

    assert [(o.passed, o.hit, o.error) for o in run_tests(rs)] == [(True, "", ""), (True, "R1", "")]


def test_yaml_rule_without_then_and_policy_validation(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text(
        "命中策略: 首次命中\n规则:\n  - 编号: G\n    当:\n      - x > 1\n    则:\n  - 则: {y: 2}\n",
        encoding="utf-8",
    )
    rs = load_yaml_rules(p)
    assert rs.hit_policy == "首次命中"
    assert rs.rules[0].outputs == {} and rs.rules[1].outputs == {"y": 2} and rs.outputs == ["y"]
    p.write_text("命中策略: 首个\n规则: []\n", encoding="utf-8")
    with pytest.raises(ValueError, match="未知的命中策略「首个」.*首次命中"):
        load_yaml_rules(p)


def test_excel_policy_validation(tmp_path):
    with pytest.raises(ValueError, match="未知的命中策略「first」.*优先级 / 唯一 / 首次命中 / 多条拼接"):
        load_rules(make_workbook(tmp_path / "bad.xlsx", [["命中策略", "first"], ["规则ID"], ["规则ID"]]))


def test_join_policy_validation(tmp_path):
    """多条拼接 下输出格写 清空 没有意义, 报错; 没用这个策略却设了「拼接分隔符」多半是策略名写错, 也报错。"""
    cleared = [*HEADER, ["R1", 1, "x", "-", "清空"]]
    with pytest.raises(ValueError, match="规则 R1 输出「标记」= 清空: 命中策略 多条拼接 只拼接有值的输出"):
        load_rules(make_workbook(tmp_path / "a.xlsx", [["命中策略", "多条拼接"], *cleared]))
    with pytest.raises(ValueError, match="设置了「拼接分隔符」, 但命中策略是「优先级」"):
        load_rules(make_workbook(tmp_path / "b.xlsx", [["拼接分隔符", ";"], *RULE]))
    ok = make_workbook(tmp_path / "c.xlsx", [["命中策略", "多条拼接"], ["拼接分隔符", ";"], *RULE])
    rs = load_rules(ok)
    assert rs.hit_policy == "多条拼接" and rs.settings["拼接分隔符"] == ";"
    p = tmp_path / "r.yaml"
    p.write_text("命中策略: 多条拼接\n规则:\n  - 编号: A\n    则: {y: =空}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="规则 A 输出「y」= =空: 命中策略 多条拼接"):
        load_yaml_rules(p)
    p.write_text("拼接分隔符: '/'\n规则:\n  - 则: {y: 1}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="设置了「拼接分隔符」"):
        load_yaml_rules(p)


def test_join_sep_lives_next_to_hit_policy(tmp_path):
    """0.8: 拼接分隔符 写在决策表左上角 (YAML 顶层), 与命中策略放一起; 旧写法 (「输入」sheet 版式的「设置」/ YAML 设置:)
    仍认一个版本; 两处都写报错; 新版「设置」里写它报错指路; 值留空时不把右边的说明当成分隔符。"""
    join = [["命中策略", "多条拼接"], ["拼接分隔符", " | "], *RULE]
    assert load_rules(make_workbook(tmp_path / "a.xlsx", join)).settings["拼接分隔符"] == " | "  # 不 strip
    blank = [["命中策略", "多条拼接"], ["拼接分隔符", None, "留空 = /"], *RULE]
    assert "拼接分隔符" not in load_rules(make_workbook(tmp_path / "b.xlsx", blank)).settings
    old = make_workbook(tmp_path / "c.xlsx", join[:1] + RULE, settings=[["拼接分隔符", ";"]], legacy=True)
    assert load_rules(old).settings["拼接分隔符"] == ";"
    both = make_workbook(tmp_path / "d.xlsx", join, settings=[["拼接分隔符", ";"]], legacy=True)
    with pytest.raises(ValueError, match="「拼接分隔符」写了两处"):
        load_rules(both)
    misplaced = make_workbook(tmp_path / "e.xlsx", join[:1] + RULE, settings=[["拼接分隔符", ";"]])
    with pytest.raises(ValueError, match="「设置」sheet: 「拼接分隔符」写在决策表左上角「命中策略」下面一行"):
        load_rules(misplaced)
    p = tmp_path / "r.yaml"
    base = "输入: {主表: {匹配: 主表}}\n命中策略: 多条拼接\n规则:\n  - 则: {y: 1}\n"
    p.write_text(base + "拼接分隔符: ';'\n", encoding="utf-8")
    assert load_rules(p).settings["拼接分隔符"] == ";"
    p.write_text(base + "设置: {拼接分隔符: ';'}\n", encoding="utf-8")  # 0.7 写法
    assert load_rules(p).settings["拼接分隔符"] == ";"
    p.write_text(base + "设置: {拼接分隔符: ';'}\n拼接分隔符: '/'\n", encoding="utf-8")
    with pytest.raises(ValueError, match="写了两处 \\(YAML 顶层 和「设置」\\)"):
        load_rules(p)


def test_settings_sheet_holds_settings_and_inputs(tmp_path):
    """0.8: 「设置」上方设置项 (标签 / 值 / 说明, 标题行与没填值的行跳过), 下方输入表 (按「名称」「匹配」那一行定位);
    旧版「输入」sheet 仍认; 两处都有输入表报错; 都没有报错指路。"""
    from openpyxl import Workbook

    settings = [["本规则的设置"], ["基准日期", "2026-10-19", "说明: 测试用"], ["输出", "o.xlsx"], ["备注", None, "x"]]
    rs = load_rules(make_workbook(tmp_path / "a.xlsx", RULE, settings=settings))
    assert rs.settings == {"基准日期": "2026-10-19", "输出": "o.xlsx"} and rs.names == ["主表", "本月"]
    legacy = load_rules(make_workbook(tmp_path / "b.xlsx", RULE, settings=settings[1:3], legacy=True))
    assert legacy.settings == rs.settings and legacy.names == rs.names
    wb = load_workbook(tmp_path / "b.xlsx")  # 旧「设置」的第 1 行不管写什么都是表头, 和 0.7 一样
    wb["设置"]["A1"], wb["设置"]["B1"] = "项目", "取值"
    wb.save(tmp_path / "b.xlsx")
    assert load_rules(tmp_path / "b.xlsx").settings == rs.settings
    wb = load_workbook(make_workbook(tmp_path / "d.xlsx", RULE))
    ws = wb.create_sheet("输入")
    for row in INPUTS:
        ws.append(row)
    wb.save(tmp_path / "d.xlsx")
    with pytest.raises(ValueError, match="「设置」sheet 里已经有输入表, 又有「输入」sheet"):
        load_rules(tmp_path / "d.xlsx")
    wb = Workbook()
    wb.active.title = "设置"
    wb.active.append(["基准日期", "2026-10-19"])
    ws = wb.create_sheet("决策表")
    for row in RULE:
        ws.append(row)
    wb.save(tmp_path / "e.xlsx")
    with pytest.raises(ValueError, match="规则文件缺少输入表: 「设置」sheet里要有一行写列名 名称 / 匹配"):
        load_rules(tmp_path / "e.xlsx")


INPUTS = [["名称", "匹配", "文件"], ["主表", "主表", "a.xlsx"], ["本月", "k", "b.xlsx"]]


def make_workbook(
    path,
    decision,
    tests=None,
    derived=None,
    merges=(),
    inputs=None,
    settings=None,
    groups=None,
    groups_header=None,
    legacy=False,
):
    """最小规则工作簿: decision / tests / inputs 是整张 sheet 的行 (含表头), groups 不含表头, merges 是要合并的区域。

    默认 0.8 版式: 「设置」上方设置项行、空一行、下方输入表;
    legacy=True 是 0.7 的「输入」sheet + 只有设置项的「设置」。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if legacy:
        ws.title = "输入"
        for row in inputs or INPUTS:
            ws.append(row)
        if settings:
            ws = wb.create_sheet("设置")
            ws.append(["设置项", "值"])
            for row in settings:
                ws.append(row)
    else:
        ws.title = "设置"
        for row in settings or []:
            ws.append(row)
        if settings:
            ws.append([])
        for row in inputs or INPUTS:
            ws.append(row)
    if derived:
        ws = wb.create_sheet("变量定义")
        ws.append(["变量", "定义"])
        for row in derived:
            ws.append(list(row))
    if groups is not None:
        ws = wb.create_sheet("字段组")
        ws.append(list(groups_header or ["名称", "列名"]))
        for row in groups:
            ws.append(list(row))
    ws = wb.create_sheet("决策表")
    for row in decision:
        ws.append(row)
    for rng in merges:
        ws.merge_cells(rng)
    if tests is not None:
        ws = wb.create_sheet("测试")
        for row in tests:
            ws.append(row)
    wb.save(path)
    return path


HEADER = [["规则ID", "优先级", "输入", "输入", "输出"], ["规则ID", "优先级", "主表.a", "主表.b", "标记"]]
TEST_HEADER = [["用例", "输入", "输入", "期望"], ["用例", "主表.a", "主表.b", "标记"]]


def test_excel_header_found_by_content_not_row_number(tmp_path):
    """issue #6: 表头写在第 1、2 行 (没有元信息行) 也能识别。
    以前第 2 行被当列类型、第一条用例被当列名, 剩下的用例空输入空期望假通过。"""
    from rule_sheet.engine import run_tests

    p = make_workbook(
        tmp_path / "r.xlsx",
        [*HEADER, ["R1", 1, "x", "-", "命中"]],
        [*TEST_HEADER, ["T1", "x", None, "命中"], ["T2", "q", None, "命中"]],
    )
    rs = load_rules(p)
    assert rs.hit_policy == "优先级" and [r.rule_id for r in rs.rules] == ["R1"]
    assert [(t.name, t.inputs["主表.a"]) for t in rs.tests] == [("T1", "x"), ("T2", "q")]
    assert [o.passed for o in run_tests(rs)] == [True, False]  # T2 不该命中, 现在能抓到


def test_excel_policy_label_anywhere_above_header(tmp_path):
    rule = ["R1", 1, "x", "-", "命中"]
    p = make_workbook(tmp_path / "r.xlsx", [["说明", "随便写"], [None, "命中策略:", "首次命中"], *HEADER, rule])
    assert load_rules(p).hit_policy == "首次命中"
    p = make_workbook(tmp_path / "r2.xlsx", [[None, "唯一"], *HEADER, rule])  # 没有标签, 单独一格策略名也认
    assert load_rules(p).hit_policy == "唯一"


@pytest.mark.parametrize(
    ("decision", "msg"),
    [
        ([["规则", "输入"], ["编号", "主表.a"]], "找不到列类型行"),
        ([["规则ID", "条件", "输出"], ["规则ID", "主表.a", "标记"]], "列类型「条件」不认识"),
        ([["规则ID", "输入", "输出"], ["规则ID", None, "标记"]], "B2 格没有写列名"),
        ([["规则ID", "规则ID"], ["规则ID", "又一个"]], "恰有一列类型为「规则ID」"),
        ([["规则ID", "输入"]], "下面没有列名行"),
    ],
)
def test_excel_header_errors(tmp_path, decision, msg):
    with pytest.raises(ValueError, match=msg):
        load_rules(make_workbook(tmp_path / "r.xlsx", decision))


def test_excel_test_sheet_guards(tmp_path):
    """期望列拼错 = 整列被忽略, 输入列拼错 = 填了不起作用, 空 sheet = 什么都没验证: 三种假通过都在加载时报错。"""
    rule = [*HEADER, ["R1", 1, "x", "-", "命中"]]
    kinds = ["用例", "输入", "期望"]
    with pytest.raises(ValueError, match="「测试」sheet: 期望列「标计」不是输出列"):
        load_rules(make_workbook(tmp_path / "a.xlsx", rule, [kinds, ["用例", "主表.a", "标计"], ["T1", "x", "命中"]]))
    with pytest.raises(ValueError, match="输入列「主表.c」不是规则引用到的变量"):
        load_rules(make_workbook(tmp_path / "b.xlsx", rule, [kinds, ["用例", "主表.c", "标记"], ["T1", "x", "命中"]]))
    with pytest.raises(ValueError, match="没有任何用例"):
        load_rules(make_workbook(tmp_path / "c.xlsx", rule, TEST_HEADER))
    # 输出列的原值 (主表.标记)、派生变量及其引用的列、内置变量 都允许作为输入; 派生变量留空 = 按表达式推导
    p = make_workbook(
        tmp_path / "d.xlsx",
        [["规则ID", "输入", "输出"], ["规则ID", "翻倍", "标记"], ["R1", "> 10", "大"]],
        [
            ["用例", "输入", "输入", "输入", "输入", "期望"],
            ["用例", "主表.标记", "主表.n", "翻倍", "本月.存在", "标记"],
            ["T1", "旧", 6, None, "是", "大"],
            ["T2", "旧", 6, 1, "是", None],  # 给了派生变量的值就不再推导: 翻倍 = 1 不命中
        ],
        derived=[("翻倍", "主表.n * 2")],
    )
    rs = load_rules(p)
    assert rs.variables == ["翻倍", "主表.n"]
    from rule_sheet.engine import run_tests

    assert [o.passed for o in run_tests(rs)] == [True, True]


def test_excel_scans_variables_referenced_inside_cells(tmp_path):
    """输入格右侧 (!= 上期.b / [0..主表.上限]) 和输出格 (=主表.c * 2 / 本月.d) 引用的变量不一定是表头输入列,
    它们也要进 variables: 测试用例才能给值, 没给的才会置空而不是被当成字面文本比较。"""
    from rule_sheet.engine import run_tests

    p = make_workbook(
        tmp_path / "r.xlsx",
        [
            ["规则ID", "输入", "输入", "输出", "输出"],
            ["规则ID", "主表.a", "主表.n", "标记", "电话"],
            ["R1", "!= 本月.b", "[0..主表.上限]", "=主表.c * 2", "本月.d"],
        ],
        [
            ["用例", "输入", "输入", "输入", "输入", "输入", "输入", "期望", "期望"],
            ["用例", "主表.a", "本月.b", "主表.n", "主表.上限", "主表.c", "本月.d", "标记", "电话"],
            ["T1", "x", "y", 5, 10, 21, "139", 42, "139"],
            ["T2", "x", "x", 5, 10, 21, "139", None, None],  # a == b 不命中
            ["T3", "x", None, 5, 10, 21, "139", None, None],  # 本月.b 没给 = 空, != 不成立, 不能拿字面 "本月.b" 去比
        ],
    )
    rs = load_rules(p)
    assert rs.variables == ["主表.a", "主表.n", "主表.c", "主表.上限", "本月.b", "本月.d"]
    assert [(o.passed, o.hit) for o in run_tests(rs)] == [(True, "R1"), (True, ""), (True, "")]


def test_excel_continuation_rows_are_or_branches(tmp_path):
    """issue #8: 相邻的同 ID 行是同一条规则的「或」分支 (每行各格 AND), 规则 ID 可以合并单元格, 输出以首行为准。"""
    from rule_sheet.engine import apply_rules

    rows = [["R1", 1, "x", "-", "命中"], [None, None, "-", "y", None], ["R2", 2, "-", "-", "兜底"]]
    p = make_workbook(
        tmp_path / "r.xlsx",
        [["命中策略", "首次命中"], *HEADER, *rows],
        merges=("A4:A5",),  # R1 的规则 ID 合并两行, 第二行读出来是 None, 展开后就是 R1
    )
    rs = load_rules(p)
    assert [r.rule_id for r in rs.rules] == ["R1", "R2"] and rs.rules[0].priority == 1
    assert rs.rules[0].branches == []  # 续行输出与首行相同 → 单纯的「或」, 不是分支
    ns = {"主表.a": pd.Series(["x", "q", "q"], dtype=object), "主表.b": pd.Series([None, "y", "n"], dtype=object)}
    res, hit = apply_rules(rs, ns, {})
    assert res["标记"].tolist() == ["命中", "命中", "兜底"] and hit.tolist() == ["R1", "R1", "R2"]


def test_excel_continuation_rows_with_own_outputs_are_branches(tmp_path):
    """issue #20: 续行写了不同的输出 → 各分支有各自输出的规则 (留空的续行输出继承首行); 规则整体计一次, 审计记 R1。"""
    from rule_sheet.engine import apply_rules

    rows = [
        ["R1", 1, "x", "-", "甲"],
        [None, None, "-", "y", "乙"],
        [None, None, "-", "z", None],
        ["R2", 2, "-", "-", "兜底"],
    ]
    p = make_workbook(tmp_path / "r.xlsx", [["命中策略", "首次命中"], *HEADER, *rows], merges=("A4:A6",))
    rs = load_rules(p)
    r1 = rs.rules[0]
    assert (
        r1.outputs == {}
        and r1.conditions == []
        and [b.outputs for b in r1.branches] == [{"标记": "甲"}, {"标记": "乙"}, {"标记": "甲"}]
    )
    ns = {
        "主表.a": pd.Series(["x", "q", "q", "n"], dtype=object),
        "主表.b": pd.Series([None, "y", "z", "w"], dtype=object),
    }
    res, hit = apply_rules(rs, ns, {"标记": pd.Series(["o"] * 4, dtype=object)})
    assert res["标记"].tolist() == ["甲", "乙", "甲", "兜底"] and hit.tolist() == ["R1", "R1", "R1", "R2"]


@pytest.mark.parametrize(
    ("rows", "msg"),
    [
        ([["R1", 1, "x", "-", "命中"], ["R1", 9, "-", "y", None]], "优先级与首行"),
        ([["R1", 1, "x", "-", "命中"], ["R2", 2, "-", "-", "兜底"], ["R1", 3, "-", "y", "命中"]], "重复出现.*相邻"),
    ],
)
def test_excel_continuation_row_errors(tmp_path, rows, msg):
    with pytest.raises(ValueError, match=msg):
        load_rules(make_workbook(tmp_path / "r.xlsx", [*HEADER, *rows]))


def test_yaml_branches(tmp_path):
    """分支: 规则级 当 是各分支共有的条件, 规则级 则 是各分支的默认值, 分支的 则 按列覆盖; 输出列顺序含分支的 则。"""
    p = tmp_path / "r.yaml"
    p.write_text(
        "输入: {主表: {匹配: 主表}}\n命中策略: 首次命中\n规则:\n"
        '  - 编号: R1\n    当: [主表.c == "on"]\n    则: {备注: 默认}\n    分支:\n'
        '      - 当: [主表.a == "x"]\n        则: {标记: 甲}\n'
        '      - 当: [主表.b == "y"]\n        则: {标记: 乙, 备注: 特别}\n'
        "      - 则: {标记: 丙}\n",
        encoding="utf-8",
    )
    rs = load_yaml_rules(p)
    r = rs.rules[0]
    assert rs.outputs == ["备注", "标记"] and r.outputs == {} and len(r.conditions) == 1
    assert [b.outputs for b in r.branches] == [
        {"备注": "默认", "标记": "甲"},
        {"备注": "特别", "标记": "乙"},
        {"备注": "默认", "标记": "丙"},
    ]
    assert [len(b.conditions) for b in r.branches] == [1, 1, 0]
    assert rs.variables == ["主表.a", "主表.b", "主表.c"]  # 分支里的条件也扫
    p.write_text("规则:\n  - 编号: R1\n    分支: {当: [x > 1]}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="规则 R1 的 分支: 要是一个非空列表"):
        load_yaml_rules(p)
    p.write_text("规则:\n  - 编号: R1\n    分支:\n      - 条件: [x > 1]\n", encoding="utf-8")
    with pytest.raises(ValueError, match="规则 R1 的第 1 个分支只认 当: / 则: / 名称:"):
        load_yaml_rules(p)
    p.write_text("命中策略: 多条拼接\n规则:\n  - 编号: R1\n    分支:\n      - 则: {y: 清空}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="规则 R1 输出「y」= 清空: 命中策略 多条拼接"):
        load_yaml_rules(p)


def test_yaml_duplicate_id_and_test_guards(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text("规则:\n  - 编号: A\n    则: {y: 1}\n  - 编号: A\n    则: {y: 2}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="规则编号「A」重复"):
        load_yaml_rules(p)
    base = "规则:\n  - 编号: A\n    当: [x > 1]\n    则: {y: 1}\n测试:\n"
    p.write_text(base, encoding="utf-8")
    with pytest.raises(ValueError, match="没有任何用例"):
        load_yaml_rules(p)
    p.write_text(base + "  - 用例: T1\n    输入: {x: 2}\n    期望: {z: 1}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="测试用例「T1」: 期望列「z」不是输出列"):
        load_yaml_rules(p)
    p.write_text(base + "  - 用例: T1\n    输入: {本月.x: 2}\n    期望: {y: 1}\n", encoding="utf-8")
    with pytest.raises(ValueError, match="测试用例「T1」: 输入列「本月.x」不是规则引用到的变量"):
        load_yaml_rules(p)


# ---------------------------------------------------------------------------
# 字段组: 组名当输出列, 一格表达整组复制 / 整组清空; 加载时展开成逐列, 引擎不知道有字段组
# ---------------------------------------------------------------------------

GROUP_DECISION = [
    ["规则ID", "输入", "输出", "输出"],
    ["规则ID", "主表.a", "带入", "y"],
    ["R1", "x", "本月", "-"],  # 整组从 本月 复制
    ["R2", "q", "本月", "=主表.a"],  # 显式写的 y 优先于字段组里的 y
    ["R3", "z", "清空", "-"],  # 整组清空
    ["R4", "w", "-", 1],  # 字段组格留空 = 整组不改
]


def test_excel_field_groups_expand_to_columns(tmp_path):
    from rule_sheet.engine import apply_rules

    rs = load_rules(make_workbook(tmp_path / "r.xlsx", GROUP_DECISION, groups=[["带入", "x, y, z"]]))
    assert rs.field_groups == {"带入": ["x", "y", "z"]} and rs.outputs == ["x", "y", "z"]
    assert [r.outputs for r in rs.rules] == [
        {"x": "本月.x", "y": "本月.y", "z": "本月.z"},
        {"x": "本月.x", "y": "=主表.a", "z": "本月.z"},
        {"x": "清空", "y": "清空", "z": "清空"},
        {"y": 1},
    ]
    assert rs.variables == ["主表.a", "本月.x", "本月.y", "本月.z"]  # 展开出来的 本月.x 也进变量清单, 用例能给值
    idx = pd.RangeIndex(4)
    ns = {
        "主表.a": pd.Series(["x", "q", "z", "w"], index=idx, dtype=object),
        "本月.x": pd.Series([1, 2, 3, 4], index=idx),
        "本月.y": pd.Series([5, 6, 7, 8], index=idx),
        "本月.z": pd.Series([9, 10, 11, 12], index=idx),
    }
    originals = {c: pd.Series([0] * 4, index=idx, dtype=object) for c in rs.outputs}
    res, hits = apply_rules(rs, ns, originals)
    assert hits.tolist() == ["R1", "R2", "R3", "R4"]
    assert res["x"].tolist() == [1, 2, None, 0] and res["z"].tolist() == [9, 10, None, 0]
    assert res["y"].tolist() == [5, "q", None, 1]


def test_diff_call_adds_both_inputs_columns_to_variables(tmp_path):
    """差异(本期, 上期, 组) 引用到的 两个输入 × 组里的列 都进变量清单; 组名与派生变量重名报错。"""
    p = tmp_path / "r.yaml"
    p.write_text(
        "输入:\n  本期: {匹配: 主表}\n  上期: {匹配: k}\n  花名册: {匹配: k}\n"
        "字段组:\n  对比: [a, b]\n变量:\n  明细: 差异(本期, 上期, 对比)\n"
        "规则:\n  - 编号: C1\n    当: [明细 != 空]\n    则: {差异类型: CHANGED, 明细列: 明细}\n"
        "测试:\n  - 用例: T1\n    输入: {本期.a: 1, 上期.a: 2}\n    期望: {差异类型: CHANGED, 明细列: 'a:2->1'}\n"
        "  - 用例: T2\n    输入: {本期.a: 1, 上期.a: 1}\n    期望: {}\n",
        encoding="utf-8",
    )
    rs = load_yaml_rules(p)
    assert rs.variables == ["上期.a", "上期.b", "明细", "本期.a", "本期.b"]  # 花名册 没被引用, 不进
    from rule_sheet.engine import run_tests

    assert [(o.passed, o.hit) for o in run_tests(rs)] == [(True, "C1"), (True, "")]
    p.write_text(
        "输入:\n  本期: {匹配: 主表}\n字段组:\n  明细: [a]\n变量:\n  明细: 本期.a\n规则: []\n", encoding="utf-8"
    )
    with pytest.raises(ValueError, match="字段组「明细」与变量定义里的派生变量重名"):
        load_yaml_rules(p)


def test_yaml_field_groups(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text(
        "输入:\n  主表: {匹配: 主表}\n  本月: {匹配: k}\n"
        "字段组:\n  带入: [x, y]\n  另一组: a, b\n"  # 列表或逗号分隔的字符串都认
        "规则:\n  - 编号: R1\n    则: {带入: 本月, y: 2}\n  - 编号: R2\n    则: {另一组: 清空, 带入: '-'}\n",
        encoding="utf-8",
    )
    rs = load_yaml_rules(p)
    assert rs.field_groups == {"带入": ["x", "y"], "另一组": ["a", "b"]}
    assert rs.rules[0].outputs == {"x": "本月.x", "y": 2} and rs.rules[1].outputs == {"a": "清空", "b": "清空"}
    assert rs.outputs == ["x", "y", "a", "b"] and "本月.x" in rs.variables
    p.write_text("字段组: {带入: 5}\n规则: []\n", encoding="utf-8")
    with pytest.raises(ValueError, match="「带入」的值要是列名清单"):
        load_yaml_rules(p)


BAD_GROUP_CELL = [*GROUP_DECISION[:2], ["R1", "x", "上月", "-"]]  # 上月 不是输入名
JOIN_GROUP = [["命中策略", "多条拼接"], *GROUP_DECISION]  # R3 整组清空, 在 多条拼接 下没有意义


@pytest.mark.parametrize(
    ("groups", "header", "decision", "msg"),
    [
        (
            [["带入", "x, y"]],
            None,
            BAD_GROUP_CELL,
            "规则 R1 \\(第 3 行\\) 输出「带入」= 上月: 字段组格只能写一个输入名",
        ),
        ([["带入", ""]], None, GROUP_DECISION, "字段组「带入」没有列"),
        ([["带入", "x, x"]], None, GROUP_DECISION, "字段组「带入」里的列「x」重复"),
        ([["带入", "x"], ["带入", "y"]], None, GROUP_DECISION, "字段组「带入」重复"),
        ([["本月", "x"]], None, GROUP_DECISION, "字段组「本月」与输入名重复"),
        ([["带入", "x, 带入"]], None, GROUP_DECISION, "字段组名「带入」同时是某个字段组里的列名"),
        ([["带入", "x"]], ["名称", "字段"], GROUP_DECISION, "「字段组」sheet 第 1 行缺少「列名」列"),
        ([["带入", "x, y"]], None, JOIN_GROUP, "规则 R3 输出「x」= 清空: 命中策略 多条拼接"),
    ],
)
def test_field_group_errors(tmp_path, groups, header, decision, msg):
    with pytest.raises(ValueError, match=msg):
        load_rules(make_workbook(tmp_path / "r.xlsx", decision, groups=groups, groups_header=header))


def test_field_group_copy_is_fine_under_join_policy(tmp_path):
    """多条拼接 只拦 清空; 字段组写输入名展开成的 上期.x 逐列参与拼接, 和手写六列一样。"""
    decision = [
        ["命中策略", "多条拼接"],
        ["拼接分隔符", ";"],
        *GROUP_DECISION[:2],
        ["R1", "x", "本月", "-"],
        ["R2", "x", "-", "附加"],
    ]
    rs = load_rules(make_workbook(tmp_path / "r.xlsx", decision, groups=[["带入", "x, y"]]))
    assert rs.hit_policy == "多条拼接" and [r.outputs for r in rs.rules] == [
        {"x": "本月.x", "y": "本月.y"},
        {"y": "附加"},
    ]
    from rule_sheet.engine import apply_rules

    idx = pd.RangeIndex(1)
    ns = {
        "主表.a": pd.Series(["x"], index=idx, dtype=object),
        "本月.x": pd.Series([1], index=idx),
        "本月.y": pd.Series(["b"], index=idx),
    }
    res, hits = apply_rules(rs, ns, {c: pd.Series([None], index=idx, dtype=object) for c in rs.outputs})
    assert hits.tolist() == ["R1; R2"] and res["x"].tolist() == ["1"] and res["y"].tolist() == ["b;附加"]


# ---------------------------------------------------------------------------
# 输入声明: 规则文件是形参 (名字 + 怎么对齐到主表), 文件只是可选的默认绑定
# ---------------------------------------------------------------------------

RULE = [*HEADER, ["R1", 1, "x", "-", "命中"]]


def test_filter_declaration(tmp_path):
    """issue #30: 输入表「筛选」列 / YAML 筛选: 只对参照表; 主表写了报错; 引用别的输入加载时就报错。"""
    rows = [
        ["名称", "匹配", "筛选"],
        ["主表", "主表", None],
        ["本月", "k", '本月.类别 == "A" 且 本月.日期 <= 基准日期'],
    ]
    rs = load_rules(make_workbook(tmp_path / "a.xlsx", RULE, inputs=rows))
    assert rs.lookups[0].filter == '本月.类别 == "A" 且 本月.日期 <= 基准日期' and rs.primary.filter == ""
    bad_primary = [["名称", "匹配", "筛选"], ["主表", "主表", '主表.a == "x"'], ["本月", "k", None]]
    with pytest.raises(ValueError, match="主表「主表」不能写「筛选」: 输出的行就是主表的行"):
        load_rules(make_workbook(tmp_path / "b.xlsx", RULE, inputs=bad_primary))
    other = [["名称", "匹配", "筛选"], ["主表", "主表", None], ["本月", "k", '主表.a == "x"']]
    with pytest.raises(ValueError, match="输入「本月」的「筛选」引用了「主表.a」; 筛选在对齐到主表之前"):
        load_rules(make_workbook(tmp_path / "c.xlsx", RULE, inputs=other))
    p = tmp_path / "r.yaml"
    p.write_text(
        '输入:\n  主表: {匹配: 主表}\n  本月:\n    匹配: k\n    筛选: 本月.类别 in ("A", "B")\n规则:\n  - 则: {y: 1}\n',
        encoding="utf-8",
    )
    assert load_rules(p).lookups[0].filter == '本月.类别 in ("A", "B")'


def test_excel_inputs_sheet(tmp_path):
    rows = [
        ["说明", "名称", "文件", "匹配", "重复键"],  # 列顺序随意, 多余的列忽略
        ["被更新的表", "主表", "a.xlsx", "主表", None],
        ["参照", "本月", None, "k", "报错"],
        ["另一张", "花名册", "c.xlsx#人员", "工号 = 主表.拟任人工号, 年月", None],
    ]
    rs = load_rules(make_workbook(tmp_path / "r.xlsx", RULE, inputs=rows, settings=[["基准日期", "2026-10-19"]]))
    assert [(s.name, s.primary, s.keys, s.dup_policy, s.path) for s in rs.inputs] == [
        ("主表", True, [], "取第一条", "a.xlsx"),
        ("本月", False, [("k", "k")], "报错", None),
        ("花名册", False, [("工号", "拟任人工号"), ("年月", "年月")], "取第一条", "c.xlsx#人员"),
    ]
    assert rs.names == ["主表", "本月", "花名册"] and rs.settings == {"基准日期": "2026-10-19"}


def test_single_input_may_omit_primary_word(tmp_path):
    rs = load_rules(make_workbook(tmp_path / "r.xlsx", RULE, inputs=[["名称", "匹配"], ["主表", None]]))
    assert rs.primary.name == "主表" and rs.lookups == []


@pytest.mark.parametrize(
    ("inputs", "msg"),
    [
        ([["名称", "匹配"], ["主表", None], ["本月", "k"]], "必须恰有一个输入的「匹配」写 主表 .* 现在有 0 个"),
        ([["名称", "匹配"], ["主表", "主表"], ["本月", "主表"]], "现在有 2 个 \\(主表, 本月\\)"),
        ([["名称", "匹配"], ["主表", "主表"], ["本月", None]], "输入「本月」没有写「匹配」"),
        ([["名称", "匹配"], ["主表", "主表"], ["主表", "k"]], "输入名「主表」重复"),
        ([["名称", "匹配", "重复键"], ["主表", "主表", None], ["本月", "k", "去重"]], "「重复键」只认 取第一条 / 报错"),
        ([["名称", "匹配"], ["主表", "主表"], ["本月", "k"], ["花名册", "工号 = 本月.工号"]], "只能对齐到主表「主表」"),
        ([["名字", "匹配"], ["主表", "主表"]], "缺少输入表: 「设置」sheet里要有一行写列名 名称 / 匹配"),
        ([["名称", "匹配"]], "没有声明任何输入"),
    ],
)
def test_excel_inputs_errors(tmp_path, inputs, msg):
    with pytest.raises(ValueError, match=msg):
        load_rules(make_workbook(tmp_path / "r.xlsx", RULE, inputs=inputs))


def test_legacy_inputs_sheet_header_must_be_row_1(tmp_path):
    with pytest.raises(ValueError, match="「输入」sheet 第 1 行缺少「名称」列"):
        load_rules(make_workbook(tmp_path / "r.xlsx", RULE, inputs=[["名字", "匹配"], ["主表", "主表"]], legacy=True))


def test_settings_sheet_rejects_unknown_keys(tmp_path):
    with pytest.raises(ValueError, match="设置项「輸出」不认识"):
        load_rules(make_workbook(tmp_path / "r.xlsx", RULE, settings=[["輸出", "x.xlsx"]]))


YAML_BASE = """输入:
  主表: {匹配: 主表}
规则:
  - 编号: R1
    当: [主表.a == 1]
    则: {b: x}
测试:
  - 用例: T1
    输入: {主表.a: 1}
    期望: {b: x}
"""


@pytest.mark.parametrize(
    ("old", "new", "msg"),
    [
        ("规则:", "命中测略: 首次命中\n规则:", "YAML 顶层 里的「命中测略」不认识"),  # 否则静默按 优先级 跑
        ("  - 编号: R1", "  - ID: R1", "YAML 第 1 条规则 里的「ID」不认识"),  # 否则静默编成 R1
        ("    期望: {b: x}", "    期待: {b: x}", "YAML 第 1 条测试用例 里的「期待」不认识"),  # 否则期望为空, 假通过
        ("  - 编号: R1\n    当: [主表.a == 1]\n    则: {b: x}", "  - R1", "YAML 第 1 条规则 要写成 键: 值"),
    ],
)
def test_yaml_rejects_unknown_keys(tmp_path, old, new, msg):
    assert old in YAML_BASE
    p = tmp_path / "r.yaml"
    p.write_text(YAML_BASE.replace(old, new), encoding="utf-8")
    with pytest.raises(ValueError, match=msg):
        load_rules(p)
    p.write_text(YAML_BASE, encoding="utf-8")
    assert load_rules(p).rules[0].rule_id == "R1"


def test_legacy_match_sheet_is_converted(tmp_path):
    """0.4 以前的「匹配设置」sheet (左表 / 右表 / 别名 / 匹配键 / 右表重复键) 折算成输入声明, 老文件不用改。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "匹配设置"
    for row in (
        ["设置项", "值"],
        ["左表", "a.xlsx"],
        ["右表", "b.xlsx"],
        ["左表别名", "本期"],
        ["右表别名", "上期"],
        ["匹配键", "岗位编码, 年月"],
        ["右表重复键", "报错"],
        ["基准日期", "2026-10-19"],
        ["输出", "out.xlsx"],
    ):
        ws.append(row)
    ws = wb.create_sheet("决策表")
    for row in [["规则ID", "输入", "输出"], ["规则ID", "本期.a", "标记"], ["R1", "x", "y"]]:
        ws.append(row)
    wb.save(tmp_path / "old.xlsx")
    rs = load_rules(tmp_path / "old.xlsx")
    assert [(s.name, s.primary, s.keys, s.dup_policy, s.path) for s in rs.inputs] == [
        ("本期", True, [], "取第一条", "a.xlsx"),
        ("上期", False, [("岗位编码", "岗位编码"), ("年月", "年月")], "报错", "b.xlsx"),
    ]
    assert rs.settings == {"基准日期": "2026-10-19", "输出": "out.xlsx"}


def test_yaml_inputs_and_settings(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text(
        "输入:\n"
        "  本期: {文件: a.xlsx, 匹配: 主表}\n"
        "  上期: {文件: b.xlsx, 匹配: 岗位编码, 重复键: 报错}\n"
        "  花名册: {匹配: 工号 = 本期.拟任人工号}\n"
        "设置: {输出: out.xlsx}\n"
        "规则:\n  - 当: [本期.x > 1]\n    则: {y: 2}\n",
        encoding="utf-8",
    )
    rs = load_yaml_rules(p)
    assert [(s.name, s.primary, s.keys, s.dup_policy, s.path) for s in rs.inputs] == [
        ("本期", True, [], "取第一条", "a.xlsx"),
        ("上期", False, [("岗位编码", "岗位编码")], "报错", "b.xlsx"),
        ("花名册", False, [("工号", "拟任人工号")], "取第一条", None),
    ]
    assert rs.settings == {"输出": "out.xlsx"} and rs.variables == ["本期.x"]
    p.write_text("输入:\n  本期: a.xlsx\n规则: []\n", encoding="utf-8")  # 单输入: 只写文件名即可
    assert load_yaml_rules(p).primary.name == "本期"
    p.write_text("输入: {本期: {文件: a.xlsx, 匹配: 主表}}\n匹配: {左表: a.xlsx}\n规则: []\n", encoding="utf-8")
    with pytest.raises(ValueError, match="不能同时出现"):
        load_yaml_rules(p)
    p.write_text("输入: {本期: {档案: a.xlsx}}\n规则: []\n", encoding="utf-8")
    with pytest.raises(ValueError, match="「档案」不认识"):
        load_yaml_rules(p)


CONCAT_INPUTS = [
    ["名称", "匹配", "重复键", "排序", "分隔符", "空值占位", "带序号", "文件"],
    ["主表", "主表", None, None, None, None, None, "a.xlsx"],
    ["推荐", "k", "拼接", "时间 降序, 工号", "/", "NA", "是", None],
    ["冻结", "k", "拼接", None, None, None, None, None],  # 四个可选项都留空 = 默认
]


def test_excel_inputs_concat_options(tmp_path):
    rs = load_rules(make_workbook(tmp_path / "r.xlsx", RULE, inputs=CONCAT_INPUTS))
    a, b = rs.lookups
    assert (a.dup_policy, a.sort, a.sep, a.placeholder, a.numbered) == (
        "拼接",
        [("时间", False), ("工号", True)],
        "/",
        "NA",
        True,
    )
    assert (b.dup_policy, b.sort, b.sep, b.placeholder, b.numbered) == ("拼接", [], ";", "", False)


def _with(rows, i, col, value):
    out = [list(r) for r in rows]
    out[i][CONCAT_INPUTS[0].index(col)] = value
    return out


@pytest.mark.parametrize(
    ("inputs", "msg"),
    [
        (_with(CONCAT_INPUTS, 1, "排序", "k"), "主表「主表」不能写「排序」"),
        (
            _with(_with(CONCAT_INPUTS, 2, "重复键", "取第一条"), 2, "分隔符", None),
            "输入「推荐」写了「排序」, 但「重复键」是 取第一条",
        ),
        (_with(CONCAT_INPUTS, 2, "排序", "时间 倒序"), "输入「推荐」的「排序」写法不对"),
        (_with(CONCAT_INPUTS, 2, "带序号", "对"), "「带序号」只认 是 / 否"),
    ],
)
def test_excel_inputs_concat_errors(tmp_path, inputs, msg):
    with pytest.raises(ValueError, match=msg):
        load_rules(make_workbook(tmp_path / "r.xlsx", RULE, inputs=inputs))


def test_inputs_append_option(tmp_path):
    rows = [
        ["名称", "匹配", "独有行"],
        ["主表", "主表", None],
        ["上期", "k", "追加"],
        ["花名册", "k", "忽略"],
        ["清单", "k", None],
    ]
    rs = load_rules(make_workbook(tmp_path / "r.xlsx", RULE, inputs=rows))
    assert [s.append for s in rs.lookups] == [True, False, False]
    with pytest.raises(ValueError, match="输入「上期」的「独有行」只认 忽略 / 追加, 不认识「保留」"):
        load_rules(make_workbook(tmp_path / "a.xlsx", RULE, inputs=[rows[0], rows[1], ["上期", "k", "保留"]]))
    with pytest.raises(ValueError, match="主表「主表」不能写「独有行」"):
        load_rules(make_workbook(tmp_path / "b.xlsx", RULE, inputs=[rows[0], ["主表", "主表", "追加"], rows[2]]))
    p = tmp_path / "r.yaml"
    p.write_text("输入:\n  本期: {匹配: 主表}\n  上期: {匹配: k, 独有行: 追加}\n规则: []\n", encoding="utf-8")
    assert load_yaml_rules(p).lookups[0].append is True


def test_yaml_inputs_concat_options(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text(
        "输入:\n  本期: {匹配: 主表}\n"
        "  推荐: {匹配: k, 重复键: 拼接, 排序: 时间 降序, 空值占位: NA, 带序号: 是}\n规则: []\n",
        encoding="utf-8",
    )
    s = load_yaml_rules(p).lookups[0]
    assert (s.dup_policy, s.sort, s.sep, s.placeholder, s.numbered) == ("拼接", [("时间", False)], ";", "NA", True)
    p.write_text("输入:\n  本期: {匹配: 主表}\n  推荐: {匹配: k, 分隔符: '/'}\n规则: []\n", encoding="utf-8")
    with pytest.raises(ValueError, match="写了「分隔符」, 但「重复键」是 取第一条"):
        load_yaml_rules(p)


def test_input_prefixed_derived_is_a_variable_in_tests(tmp_path):
    """变量定义里名字带输入前缀的 (本月.x) 是「在那张表的行上算」的派生变量: 进变量清单, 用例里直接给值、不推导。"""
    from rule_sheet.engine import run_tests

    p = make_workbook(
        tmp_path / "r.xlsx",
        [["规则ID", "输入", "输出"], ["规则ID", "本月.x", "标记"], ["R1", "非空", "有"]],
        [
            ["用例", "输入", "输入", "期望"],
            ["用例", "本月.x", "本月.a", "标记"],
            ["T1", "pq", None, "有"],
            ["T2", None, "p", None],  # 没给 本月.x 就是空, 不从 a / b 推导
        ],
        derived=[("本月.x", "拼接(本月.a, 本月.b)")],
    )
    rs = load_rules(p)
    assert rs.variables == ["本月.x", "本月.a", "本月.b"] and rs.input_derived == [("本月.x", "拼接(本月.a, 本月.b)")]
    assert [o.passed for o in run_tests(rs)] == [True, True]


BD_RULE = [["规则ID", "输入", "输出"], ["规则ID", "主表.d", "标记"], ["R1", ">= 基准日期", "新"]]
BD_CASES = [
    ["用例", "输入", "期望"],
    ["用例", "主表.d", "标记"],
    ["T1", date(2026, 10, 20), "新"],
    ["T2", date(2026, 10, 19), None],
]
BD_YAML = """输入:
  主表: {文件: a.xlsx, 匹配: 主表}
规则:
  - 编号: R1
    当: [主表.d >= 基准日期]
    则: {标记: 新}
测试:
  基准日期: 2026-10-20
  用例:
    - 用例: T1
      输入: {主表.d: 2026-10-20}
      期望: {标记: 新}
"""


def test_test_sheet_base_date_owns_the_cases(tmp_path):
    """issue #31: 用例的期望按固定日期写, 日期跟着用例走: 「测试」sheet 表头上方 > 「设置」(生产默认) > 当天;
    用例自己的 基准日期 列仍优先。"""
    from rule_sheet.engine import case_base_date, run_tests

    tests = [["基准日期", "2026-10-20", "用例按它写"], *BD_CASES]
    rs = load_rules(make_workbook(tmp_path / "r.xlsx", BD_RULE, tests, settings=[["基准日期", "2030-01-01"]]))
    assert rs.test_base_date == "2026-10-20" and rs.uses_base_date
    assert case_base_date(rs) == (pd.Timestamp("2026-10-20"), "测试的基准日期")
    assert all(o.passed for o in run_tests(rs))  # 「设置」的 2030-01-01 是生产默认, 用例不看它
    own = [
        ["用例", "输入", "输入", "期望"],
        ["用例", "主表.d", "基准日期", "标记"],
        ["T1", date(2026, 10, 20), date(2026, 10, 25), None],  # 这条用例自己的日期: 10-20 < 10-25, 不命中
    ]
    rs = load_rules(make_workbook(tmp_path / "o.xlsx", BD_RULE, [["基准日期", date(2026, 10, 20)], *own]))
    assert all(o.passed for o in run_tests(rs))
    with pytest.raises(ValueError, match="「测试」sheet 表头上方的基准日期「下周一」不是日期"):
        load_rules(make_workbook(tmp_path / "b.xlsx", BD_RULE, [["基准日期", "下周一"], *BD_CASES]))


def test_cases_need_a_fixed_base_date_when_rules_use_it(tmp_path):
    """规则引用了基准日期, 用例却哪里都没有固定的: 按当天算会今天过明天挂, 跑测试时报错; 「设置」有值的老文件照旧;
    不引用基准日期的规则文件不受影响。"""
    from rule_sheet.engine import run_tests

    rs = load_rules(make_workbook(tmp_path / "r.xlsx", BD_RULE, BD_CASES))
    with pytest.raises(ValueError, match="用例 T1, T2 没有固定的基准日期"):
        run_tests(rs)
    old = load_rules(make_workbook(tmp_path / "s.xlsx", BD_RULE, BD_CASES, settings=[["基准日期", "2026-10-20"]]))
    assert all(o.passed for o in run_tests(old))
    plain = load_rules(make_workbook(tmp_path / "p.xlsx", RULE, [*TEST_HEADER, ["T1", "x", None, "命中"]]))
    assert not plain.uses_base_date and all(o.passed for o in run_tests(plain))


def test_yaml_tests_with_base_date(tmp_path):
    """YAML 测试: 写成 {基准日期, 用例: [...]}; 列表写法照旧; 键拼错 / 日期写错加载时报错。"""
    from rule_sheet.engine import run_tests

    p = tmp_path / "r.yaml"
    p.write_text(BD_YAML, encoding="utf-8")
    rs = load_rules(p)
    assert rs.test_base_date == date(2026, 10, 20) and [t.name for t in rs.tests] == ["T1"]
    assert all(o.passed for o in run_tests(rs))
    for old, new, msg in [
        ("  基准日期: 2026-10-20\n", "  基准日其: 2026-10-20\n", "YAML 测试: 里的「基准日其」不认识"),
        ("基准日期: 2026-10-20", "基准日期: 下周", "YAML 测试: 里的基准日期「下周」不是日期"),
        (BD_YAML[BD_YAML.index("  用例:\n") :], "  用例: T1\n", "用例: 要是一个列表"),
    ]:
        p.write_text(BD_YAML.replace(old, new), encoding="utf-8")
        with pytest.raises(ValueError, match=msg):
            load_rules(p)
    head, cases = BD_YAML.split("  基准日期: 2026-10-20\n  用例:\n")
    p.write_text(head + "\n".join(line[2:] for line in cases.splitlines()) + "\n", encoding="utf-8")
    rs = load_rules(p)
    assert rs.test_base_date is None and len(rs.tests) == 1
    with pytest.raises(ValueError, match="没有固定的基准日期"):
        run_tests(rs)


def test_mentions_base_date_follows_name_boundaries():
    from rule_sheet.loaders import mentions_base_date

    assert mentions_base_date(["主表.d >= 基准日期"]) and mentions_base_date(["=基准日期"])
    assert mentions_base_date(["天数差(基准日期, 主表.d)"]) and mentions_base_date(["`基准日期` - 1天"])
    assert not mentions_base_date(['主表.x == "基准日期"', "基准日期前一年 > 1", "主表.基准日期", "包含 x"])


def test_missing_file_and_missing_sheet(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_rules(tmp_path / "nope.xlsx")
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "决策表"
    wb.save(tmp_path / "bad.xlsx")
    with pytest.raises(ValueError, match="缺少输入表"):
        load_rules(tmp_path / "bad.xlsx")
    wb = Workbook()
    wb.active.title = "随便"
    wb.save(tmp_path / "bad2.xlsx")
    with pytest.raises(ValueError, match="缺少「决策表」sheet"):
        load_rules(tmp_path / "bad2.xlsx")
