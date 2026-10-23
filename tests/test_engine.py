"""执行核心: 命名空间构建、命中策略合并、输出格解析、测试用例运行。"""

from datetime import date

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from rule_sheet.conditions import cell_condition, expr_condition
from rule_sheet.engine import (
    _is_date_column,
    add_derived,
    apply_rules,
    build_namespace,
    read_table,
    run_tests,
    split_sheet,
    write_xlsx_safely,
)
from rule_sheet.model import InputSpec, Rule, RuleSet, TestCase
from rule_sheet.style import display_width

# ---------------------------------------------------------------------------
# read_table
# ---------------------------------------------------------------------------


def test_read_table_keeps_text_cells_as_text(tmp_path):
    """pandas 3 默认会把文本格式的 "007" 读成 7、电话读成浮点; 业务数据不能接受。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["编号", "电话", "金额", "日期"])
    ws.append(["007", "13800000001", 100, pd.Timestamp("2026-10-01").to_pydatetime()])
    ws.append(["008", None, 250, None])
    for row in ws.iter_rows(min_row=2, max_col=2):
        for c in row:
            c.number_format = "@"
    p = tmp_path / "t.xlsx"
    wb.save(p)
    df = read_table(p)
    assert df["编号"].tolist() == ["007", "008"]
    assert df["电话"].iloc[0] == "13800000001" and pd.isna(df["电话"].iloc[1])
    assert pd.api.types.is_integer_dtype(df["金额"]) and df["金额"].tolist() == [100, 250]
    assert pd.api.types.is_datetime64_any_dtype(df["日期"])


def inputs_of(*specs: tuple):
    """(名称, 匹配键列表 | None) ...; 第一个是主表。键写成 "本输入的列=主表的列" 或 "同名列"。"""
    out = []
    for i, (name, keys) in enumerate(specs):
        pairs = [tuple(k.split("=")) if "=" in k else (k, k) for k in (keys or [])]
        out.append(InputSpec(name, primary=(i == 0), keys=pairs))
    return out


def make_rs(rules, outputs, hit_policy="优先级", **kw):
    return RuleSet(
        inputs=kw.pop("inputs", inputs_of(("主表", None), ("本月", ["编号"]))),
        settings=kw.pop("settings", {}),
        derived=kw.pop("derived", []),
        hit_policy=hit_policy,
        rules=rules,
        outputs=outputs,
        **kw,
    )


# ---------------------------------------------------------------------------
# build_namespace
# ---------------------------------------------------------------------------


def test_namespace_prefixes_exist_flag_and_collision():
    left = pd.DataFrame({"编号": ["A", "B", "C"], "金额": [1, 2, 3], "备注": ["x", "y", "z"]})
    right = pd.DataFrame({"编号": ["A", "C"], "金额": [10, 30], "电话": ["t1", None]})
    rs = make_rs([], [], settings={"基准日期": "2026-10-19"})
    ns = build_namespace(rs, {"主表": left, "本月": right})
    assert ns["主表.金额"].tolist() == [1, 2, 3]
    assert ns["本月.金额"].tolist()[0] == 10 and pd.isna(ns["本月.金额"].iloc[1])
    assert ns["本月.存在"].tolist() == [True, False, True]
    assert "本月.编号" not in ns  # 参照表的匹配键列不暴露, 值就是主表的
    assert ns["基准日期"] == pd.Timestamp("2026-10-19")


def test_namespace_three_inputs_with_own_keys_and_same_named_columns():
    """N 个参照表各自按自己的键对齐到主表; 参照表之间同名列 (姓名) 互不干扰; 键名可以不同 (工号 = 主表.拟任人工号)。"""
    base = pd.DataFrame({"岗位": ["P1", "P2", "P3"], "拟任人工号": ["E1", "E2", None], "姓名": ["a", "b", "c"]})
    last = pd.DataFrame({"岗位": ["P1", "P3"], "姓名": ["a0", "c0"], "状态": ["新增", "进展"]})
    roster = pd.DataFrame({"工号": ["E1", "E2"], "姓名": ["张三", "李四"], "部门": ["财务", "人事"]})
    rs = make_rs([], [], inputs=inputs_of(("本期", None), ("上期", ["岗位"]), ("花名册", ["工号=拟任人工号"])))
    ns = build_namespace(rs, {"本期": base, "上期": last, "花名册": roster})
    assert ns["本期.姓名"].tolist() == ["a", "b", "c"]
    assert ns["上期.姓名"].tolist()[0] == "a0" and pd.isna(ns["上期.姓名"].iloc[1])
    assert ns["花名册.姓名"].tolist()[:2] == ["张三", "李四"] and pd.isna(ns["花名册.姓名"].iloc[2])
    assert ns["上期.存在"].tolist() == [True, False, True]
    assert ns["花名册.存在"].tolist() == [True, True, False]  # 主表键为空的行匹配不到
    assert "花名册.工号" not in ns and "上期.岗位" not in ns


def test_namespace_single_input_needs_no_join():
    base = pd.DataFrame({"k": [1, 2], "v": ["a", "b"]})
    ns = build_namespace(make_rs([], [], inputs=inputs_of(("表", None))), {"表": base})
    assert set(ns) == {"表.k", "表.v", "表.存在", "表", "基准日期"}  # 主表的 .存在 总在; 输入名本身是 InputRef
    assert ns["表.存在"].tolist() == [True, True]


def test_same_treats_blank_text_as_null():
    """变更判定 (审计列 / 用例期望 / diff) 里空串、空白与空值相同 (issue #18 / #22); 文本仍严格比较。"""
    from rule_sheet.engine import same

    assert same("", None) and same(" ", float("nan")) and same(None, None) and same("", "  ")
    assert not same("", "0") and not same("007", 7) and same(7.0, 7) and same("a", "a")
    assert same(pd.Timestamp("2026-10-20"), date(2026, 10, 20))


def test_lookup_rows_with_null_keys_never_match():
    # pandas merge 会把两边的 NaN 键配成一对, 参照表里键为空的行必须先剔掉, 否则 存在 会假阳性
    base = pd.DataFrame({"k": [1, None], "v": ["a", "b"]})
    lookup = pd.DataFrame({"k": [1, None], "w": ["x", "y"]})
    ns = build_namespace(
        make_rs([], [], inputs=inputs_of(("主表", None), ("本月", ["k"]))), {"主表": base, "本月": lookup}
    )
    assert ns["本月.存在"].tolist() == [True, False]
    assert ns["本月.w"].tolist()[0] == "x" and pd.isna(ns["本月.w"].iloc[1])


def test_duplicate_lookup_keys_policy():
    left = pd.DataFrame({"编号": [1]})
    right = pd.DataFrame({"编号": [1, 1], "v": ["first", "second"]})
    notes: list[str] = []
    ns = build_namespace(make_rs([], []), {"主表": left, "本月": right}, notes)
    assert ns["本月.v"].tolist() == ["first"] and any("按「取第一条」处理" in n for n in notes)
    strict = make_rs([], [])
    strict.lookups[0].dup_policy = "报错"
    with pytest.raises(ValueError, match="输入「本月」有 1 行重复键"):
        build_namespace(strict, {"主表": left, "本月": right})


def test_missing_key_column():
    with pytest.raises(ValueError, match="主表「主表」没有列「编号」"):
        build_namespace(make_rs([], []), {"主表": pd.DataFrame({"a": [1]}), "本月": pd.DataFrame({"编号": [1]})})
    with pytest.raises(ValueError, match="输入「本月」没有匹配键列「编号」"):
        build_namespace(make_rs([], []), {"主表": pd.DataFrame({"编号": [1]}), "本月": pd.DataFrame({"b": [1]})})


def test_build_namespace_runtime_base_date():
    """issue #28: 分步调用者也能给基准日期; 带输入前缀的派生变量 (聚合前在参照表行上算) 同样看到注入值。"""
    rs = make_rs(
        [],
        [],
        settings={"基准日期": date(2026, 10, 19)},
        derived=[("本月.天数", "天数差(基准日期, 本月.日期)")],
    )
    tables = {
        "主表": pd.DataFrame({"编号": [1]}),
        "本月": pd.DataFrame({"编号": [1], "日期": [pd.Timestamp("2026-10-01")]}),
    }
    notes: list[str] = []
    ns = build_namespace(rs, tables, notes, variables={"基准日期": pd.Timestamp("2026-10-31")})
    assert ns["基准日期"] == pd.Timestamp("2026-10-31") and ns["本月.天数"].tolist() == [30]
    assert notes[0] == "基准日期 = 2026-10-31 (运行时给定)"
    ns = build_namespace(rs, tables)  # 不给就是规则文件「设置」里的
    assert ns["基准日期"] == pd.Timestamp("2026-10-19") and ns["本月.天数"].tolist() == [18]
    notes = []
    build_namespace(make_rs([], []), tables, notes)
    assert notes[0].endswith("(规则文件没设, 取运行当天)")
    with pytest.raises(ValueError, match="运行时变量「输出」不认识"):  # 输出 / 拼接分隔符 是设置项, 但不是运行时变量
        build_namespace(rs, tables, variables={"输出": "x.xlsx"})


def test_filter_keeps_rows_before_dedup_and_join():
    """issue #30: 筛选在去重 / 对齐之前: 别的类别里的同名键不算重复 (重复键: 报错 不误报), 也不会被取第一条挑中;
    可以引用带前缀的派生列和 基准日期 (运行时给的也生效); 一行不剩报错; 引用不存在的列报错并指出是哪个输入。"""
    base = pd.DataFrame({"组织": ["X", "A", "B"]})
    mixed = pd.DataFrame(
        {
            "类别": ["销服", "职能", "职能", "销服"],
            "项": ["X", "X", "A", "B"],
            "生效": pd.to_datetime(["2026-01-01", "2026-01-01", "2026-12-01", "2026-01-01"]),
        }
    )
    specs = inputs_of(("主表", None), ("职能", ["项=组织"]))
    specs[1].dup_policy = "报错"
    specs[1].filter = '职能.类别 == "职能" 且 职能.生效 <= 基准日期'
    rs = make_rs([], [], inputs=specs, settings={"基准日期": date(2026, 10, 23)})
    notes: list[str] = []
    ns = build_namespace(rs, {"主表": base, "职能": mixed}, notes)
    assert ns["职能.存在"].tolist() == [True, False, False]  # A 还没生效, B 是销服的
    assert "输入「职能」按筛选保留 1 / 4 行" in notes
    ns = build_namespace(rs, {"主表": base, "职能": mixed}, variables={"基准日期": "2026-12-31"})
    assert ns["职能.存在"].tolist() == [True, True, False]
    specs[1].filter = '职能.标记 == "是"'  # 带前缀的派生列先算, 筛选能用它
    rs.derived = [("职能.标记", '如果(职能.类别 == "职能", "是", "否")')]
    assert build_namespace(rs, {"主表": base, "职能": mixed})["职能.存在"].tolist() == [True, True, False]
    rs.derived = []
    specs[1].filter = '职能.类别 == "职能 "'
    with pytest.raises(ValueError, match='输入「职能」按筛选「职能.类别 == "职能 "」一行都没留下 \\(原有 4 行\\)'):
        build_namespace(rs, {"主表": base, "职能": mixed})
    specs[1].filter = '职能.种类 == "职能"'
    with pytest.raises(ValueError, match='输入「职能」的筛选「职能.种类 == "职能"」求值失败'):
        build_namespace(rs, {"主表": base, "职能": mixed})


def test_concat_policy_aggregates_rows_before_join():
    """重复键: 拼接 (issue #17): 同键的行按排序接成一行 (空值排末尾), 空值占位保位置, 带序号;
    带输入前缀的派生变量在那张表的行上算、发生在聚合之前; 没设占位符且整组全空的列为空。"""
    base = pd.DataFrame({"岗位": ["P1", "P2", "P3"], "现任": ["C1", "C2", None]})
    rec = pd.DataFrame(
        {
            "岗位": ["P1", "P1", "P1", "P2"],
            "工号": ["E1", "E2", "E3", "E4"],
            "部门": ["财务", None, "人事", None],
            "职级": ["22", "23", "21", None],
            "时间": pd.to_datetime(["2026-09-01", "2026-10-01", None, None]),
        }
    )
    frz = pd.DataFrame(
        {
            "工号": ["C1", "C1"],
            "开始": pd.to_datetime(["2026-05-01", "2026-09-01"]),
            "依据": ["纪律", "审计"],
            "备注": [None, None],
        }
    )
    rs = make_rs(
        [],
        [],
        inputs=inputs_of(("岗位", None), ("推荐", ["岗位"]), ("冻结", ["工号=现任"])),
        derived=[("推荐.现岗", '拼接(推荐.部门, "-", 推荐.职级)'), ("有推荐", "推荐.存在")],
    )
    rec_spec, frz_spec = rs.lookups
    rec_spec.dup_policy, rec_spec.sort, rec_spec.placeholder = "拼接", [("时间", False)], "NA"
    frz_spec.dup_policy, frz_spec.sort, frz_spec.numbered, frz_spec.sep = "拼接", [("开始", False)], True, "/"
    assert rs.input_derived == [("推荐.现岗", '拼接(推荐.部门, "-", 推荐.职级)')] and rs.plain_derived == [
        ("有推荐", "推荐.存在")
    ]
    notes: list[str] = []
    ns = build_namespace(rs, {"岗位": base, "推荐": rec, "冻结": frz}, notes)
    assert ns["推荐.工号"].tolist()[:2] == ["E2;E1;E3", "E4"] and pd.isna(ns["推荐.工号"].iloc[2])
    assert ns["推荐.部门"].tolist()[:2] == ["NA;财务;人事", "NA"]
    assert ns["推荐.现岗"].tolist()[:2] == ["-23;财务-22;人事-21", "NA"]  # 聚合前逐行算; 列参数全空为空, 再被占位
    assert ns["推荐.时间"].tolist()[:2] == ["2026-10-01;2026-09-01;NA", "NA"]
    assert ns["推荐.存在"].tolist() == [True, True, False]
    assert ns["冻结.依据"].tolist()[0] == "1.审计/2.纪律" and pd.isna(ns["冻结.依据"].iloc[1])
    assert pd.isna(ns["冻结.备注"].iloc[0])  # 没设占位符, 整组全空 -> 空
    assert ns["冻结.存在"].tolist() == [True, False, False]
    assert sum("按「拼接」接成一行" in n for n in notes) == 2
    add_derived(rs, ns, base.index, lenient=False)
    assert ns["有推荐"].tolist() == [True, True, False] and ns["推荐.现岗"].tolist()[0] == "-23;财务-22;人事-21"


def test_concat_policy_errors():
    base = pd.DataFrame({"岗位": ["P1"], "x": [1]})
    rec = pd.DataFrame({"岗位": ["P1"], "工号": ["E1"]})
    rs = make_rs([], [], inputs=inputs_of(("岗位", None), ("推荐", ["岗位"])))
    rs.lookups[0].dup_policy, rs.lookups[0].sort = "拼接", [("时间", True)]
    with pytest.raises(ValueError, match="输入「推荐」没有排序列「时间」"):
        build_namespace(rs, {"岗位": base, "推荐": rec})
    rs.lookups[0].sort = []
    rs.derived = [("推荐.工号", "推荐.工号")]
    with pytest.raises(ValueError, match="派生变量「推荐.工号」与输入「推荐」的列「工号」同名"):
        build_namespace(rs, {"岗位": base, "推荐": rec})
    rs.derived = [("推荐.y", "岗位.x")]  # 只能引用同一张表的列
    with pytest.raises(ValueError, match="派生变量「推荐.y」在输入「推荐」的行上求值失败"):
        build_namespace(rs, {"岗位": base, "推荐": rec})


def test_append_policy_adds_lookup_only_rows_to_the_row_set():
    """独有行: 追加 (issue #19): 参照表里主表没有的键接到行集末尾, 同名列带入、键列按映射填、主表独有列为空;
    主表 .存在 在这些行上为 否; 其它参照表照常对齐到追加来的行; 多个追加输入按声明顺序, 后面的对齐到已扩的行集。"""
    base = pd.DataFrame({"编码": ["P1", "P2"], "名称": ["a", "b"], "工号": ["E1", "E2"], "本期独有": [1, 2]})
    last = pd.DataFrame(
        {"编码": ["P2", "P3", "P4"], "名称": ["b0", "c0", "d0"], "工号": ["E2", "E3", "E4"], "状态": ["x", "y", "z"]}
    )
    roster = pd.DataFrame({"人员工号": ["E1", "E3"], "姓名": ["张三", "王五"]})
    older = pd.DataFrame(
        {"岗位编码": ["P3", "P5"], "名称": ["c00", "e00"]}
    )  # 键名与主表不同; P3 已被 上期 追加, 不再追加
    rs = make_rs(
        [],
        [],
        inputs=inputs_of(
            ("本期", None), ("上期", ["编码"]), ("花名册", ["人员工号=工号"]), ("上上期", ["岗位编码=编码"])
        ),
    )
    rs.lookups[0].append = rs.lookups[2].append = True
    notes: list[str] = []
    ns = build_namespace(rs, {"本期": base, "上期": last, "花名册": roster, "上上期": older}, notes)
    assert ns["本期.编码"].tolist() == ["P1", "P2", "P3", "P4", "P5"]
    assert ns["本期.名称"].tolist() == ["a", "b", "c0", "d0", "e00"]  # 同名列带入
    assert ns["本期.工号"].tolist()[:4] == ["E1", "E2", "E3", "E4"] and pd.isna(ns["本期.工号"].iloc[4])
    assert ns["本期.本期独有"].tolist()[:2] == [1, 2] and ns["本期.本期独有"].isna().tolist()[2:] == [True] * 3
    assert ns["本期.存在"].tolist() == [True, True, False, False, False]
    assert ns["上期.存在"].tolist() == [False, True, True, True, False] and ns["上期.状态"].tolist()[2] == "y"
    assert ns["花名册.存在"].tolist() == [True, False, True, False, False]  # 追加行带入了 工号, 花名册照常对上
    assert ns["花名册.姓名"].tolist()[2] == "王五"
    assert ns["上上期.存在"].tolist() == [False, False, True, False, True]
    assert "本期.状态" not in ns  # 主表没有的列不会因为追加多出来
    assert [n for n in notes if "按「追加」接在末尾" in n] == [
        "输入「上期」有 2 行主表没有的键, 按「追加」接在末尾",
        "输入「上上期」有 1 行主表没有的键, 按「追加」接在末尾",
    ]
    from rule_sheet.engine import primary_frame

    frame = primary_frame(rs, ns, list(base.columns))
    assert list(frame.columns) == list(base.columns) and len(frame) == 5 and frame["编码"].tolist()[-1] == "P5"


def test_append_with_concat_policy_appends_the_aggregated_row():
    base = pd.DataFrame({"编码": ["P1"]})
    rec = pd.DataFrame({"编码": ["P2", "P2", "P1"], "工号": ["E2", "E3", "E1"]})
    rs = make_rs([], [], inputs=inputs_of(("本期", None), ("推荐", ["编码"])))
    rs.lookups[0].dup_policy, rs.lookups[0].append = "拼接", True
    ns = build_namespace(rs, {"本期": base, "推荐": rec})
    assert ns["本期.编码"].tolist() == ["P1", "P2"] and ns["推荐.工号"].tolist() == ["E1", "E2;E3"]
    assert ns["本期.存在"].tolist() == [True, False]


def test_primary_exists_defaults_to_true_in_cases():
    """用例里主表的 .存在 没给 = 是 (普通行), 写 否 = 追加来的行。"""
    rs = make_rs(
        [Rule("X1", 1, [cell_condition("本期.存在", "否")], {"差异类型": "REMOVED"})],
        ["差异类型"],
        inputs=inputs_of(("本期", None), ("上期", ["编码"])),
        variables=["本期.存在"],
        tests=[
            TestCase("没给 = 普通行", {}, {}),
            TestCase("追加来的行", {"本期.存在": "否"}, {"差异类型": "REMOVED"}),
        ],
    )
    assert [(o.passed, o.hit) for o in run_tests(rs)] == [(True, ""), (True, "X1")]


def test_input_prefixed_derived_is_given_not_computed_in_cases():
    """带输入前缀的派生变量在用例里直接给值: 真实数据里它在聚合前算过, 用例里从聚合后的文本再推导是错的。"""
    rs = make_rs(
        [Rule("R1", 1, [cell_condition("本月.x", "非空")], {"y": "本月.x"})],
        ["y"],
        derived=[("本月.x", "拼接(本月.a, 本月.b)")],
        variables=["本月.x", "本月.a", "本月.b"],
        tests=[
            TestCase("没给 x, 不从 a / b 推导", {"本月.a": "p", "本月.b": "q"}, {}),
            TestCase("给了就用", {"本月.x": "p;q"}, {"y": "p;q"}),
        ],
    )
    assert [(o.passed, o.hit) for o in run_tests(rs)] == [(True, ""), (True, "R1")]


def test_split_sheet(tmp_path):
    plain = tmp_path / "a.xlsx"
    plain.write_bytes(b"")
    hashed = tmp_path / "报表#1.xlsx"
    hashed.write_bytes(b"")
    assert split_sheet(str(plain)) == (plain, None)
    assert split_sheet(f"{plain}#明细") == (plain, "明细")
    assert split_sheet(str(hashed)) == (hashed, None)  # 文件名本身含 #, 整个字符串是存在的文件就不拆
    assert split_sheet("a.xlsx#s", tmp_path) == (plain, "s")  # 相对路径按 base_dir 解析


# ---------------------------------------------------------------------------
# apply_rules
# ---------------------------------------------------------------------------


def test_priority_merges_per_output_column(ns_two_rows):
    r1 = Rule("R1", 1, [cell_condition("主表.状态", "活跃")], {"备注": "来自R1", "电话": "-"})
    r2 = Rule("R2", 2, [cell_condition("主表.状态", "活跃")], {"备注": "来自R2", "电话": "本月.电话"})
    rs = make_rs([r2, r1], ["备注", "电话"])
    results, hit = apply_rules(rs, ns_two_rows, {"电话": ns_two_rows["主表.电话"]})
    assert results["备注"].tolist()[0] == "来自R1"  # 两条都写了, 优先级小的赢
    assert results["电话"].tolist()[0] == "139"  # 只有 R2 写了, 取 R2 (复制变量)
    assert results["电话"].tolist()[1] is None  # 未命中行保持原值 (原值为空)
    assert hit.tolist() == ["R1; R2", "", ""]


def test_priority_lets_independent_output_columns_hit_separately():
    """issue #27: 两组互不相干的规则各管一列 (两条定 岗位类型, 一条定 类别) 放一张表, 优先级 下每列各自取
    第一条写了该列的规则; 写 `-` 的规则对该列不参与也不拦截, 输出列在主表里不存在 (originals 没有) 也一样。
    首次命中 是整行的, 作对照: 类别 永远拿不到 R6。"""
    idx = pd.RangeIndex(3)
    ns = {
        "主表.层级": pd.Series(["总监", "经理", "经理"], index=idx, dtype=object),
        "主表.空岗": pd.Series(["Y", "Y", "N"], index=idx, dtype=object),
    }
    r1 = Rule("R1", 1, [cell_condition("主表.层级", "总监")], {"岗位类型": "直管", "类别": "-"})
    r2 = Rule("R2", 2, [cell_condition("主表.层级", "!= 总监")], {"岗位类型": "非直管", "类别": "-"})
    r6 = Rule("R6", 6, [cell_condition("主表.空岗", "Y")], {"岗位类型": "-", "类别": "冒泡-空岗"})
    outputs = ["岗位类型", "类别"]
    res, hit = apply_rules(make_rs([r1, r2, r6], outputs), ns, {})
    assert res["岗位类型"].tolist() == ["直管", "非直管", "非直管"]
    assert res["类别"].tolist() == ["冒泡-空岗", "冒泡-空岗", None]
    assert hit.tolist() == ["R1; R6", "R2; R6", "R2"]
    res, hit = apply_rules(make_rs([r1, r2, r6], outputs, hit_policy="首次命中"), ns, {})
    assert res["岗位类型"].tolist() == ["直管", "非直管", "非直管"]
    assert res["类别"].tolist() == [None, None, None]
    assert hit.tolist() == ["R1", "R2", "R2"]


def test_output_kinds(ns_two_rows):
    r = Rule(
        "R",
        1,
        [cell_condition("本月.存在", "是")],
        {
            "a": "-",
            "b": '"引号字面值"',
            "c": "=主表.金额 * 2",
            "d": 42,
            "e": "本月.金额",
        },
    )
    rs = make_rs([r], ["a", "b", "c", "d", "e"])
    res, _ = apply_rules(rs, ns_two_rows, {"a": pd.Series(["keep"] * 3)})
    assert res["a"].tolist() == ["keep"] * 3
    assert res["b"].iloc[0] == "引号字面值" and res["b"].iloc[1] is None
    assert res["c"].iloc[0] == 20000.0
    assert res["d"].iloc[2] == 42
    assert res["e"].iloc[2] == 5100.0


def test_output_dangling_variable_is_an_error(ns_two_rows):
    # 想复制 本月.电话 却写成 本月.电洁: 不能把字面 "本月.电洁" 写进结果; 报错带规则 ID 和输出列
    r = Rule("R7", 1, [cell_condition("本月.存在", "是")], {"电话": "本月.电洁", "备注": '"本月.电洁"'})
    with pytest.raises(KeyError, match="规则 R7 输出「电话」: 规则引用了未定义的变量「本月.电洁」"):
        apply_rules(make_rs([r], ["电话", "备注"]), ns_two_rows, {})
    r.outputs = {"备注": '"本月.电洁"'}  # 加引号 = 明确要当文字
    res, _ = apply_rules(make_rs([r], ["备注"]), ns_two_rows, {})
    assert res["备注"].iloc[0] == "本月.电洁"


def test_clear_output_writes_null(ns_two_rows):
    """输出格 清空 把目标列写成空值 (issue #2); - 仍是不改; 加引号的 "清空" 是普通文字。"""
    r = Rule("R", 1, [cell_condition("本月.存在", "是")], {"a": "清空", "b": "置空", "c": "-", "d": '"清空"'})
    rs = make_rs([r], ["a", "b", "c", "d"])
    keep = {k: pd.Series(["keep"] * 3) for k in "abcd"}
    res, _ = apply_rules(rs, ns_two_rows, keep)
    assert res["a"].tolist() == [None, "keep", None]
    assert res["b"].tolist() == [None, "keep", None]
    assert res["c"].tolist() == ["keep"] * 3
    assert res["d"].tolist() == ["清空", "keep", "清空"]


def test_null_keyword_in_outputs_tests_and_derived(ns_two_rows):
    """关键字 空: 输出格 (空 / =空) 等价于 清空; 期望格写 空 = 期望为空; 输入写 空 = 空值; 派生变量算出 空 = 空列。"""
    r = Rule("R", 1, [cell_condition("本月.存在", "是")], {"a": "空", "b": "=空", "c": "= 空"})
    keep = {k: pd.Series(["keep"] * 3) for k in "abc"}
    res, _ = apply_rules(make_rs([r], ["a", "b", "c"]), ns_two_rows, keep)
    assert res["a"].tolist() == res["b"].tolist() == res["c"].tolist() == [None, "keep", None]

    r1 = Rule("R1", 1, [cell_condition("主表.状态", "空")], {"备注": "没状态"})
    r2 = Rule("R2", 2, [cell_condition("主表.状态", "活跃")], {"备注": "空"})
    rs = make_rs(
        [r1, r2],
        ["备注"],
        tests=[
            TestCase("输入写 空 = 空值", {"主表.状态": "空", "主表.备注": "旧"}, {"备注": "没状态"}),
            TestCase("期望写 空 = 期望为空", {"主表.状态": "活跃", "主表.备注": "旧"}, {"备注": "空"}),
            TestCase("没清就失败", {"主表.状态": "休眠", "主表.备注": "旧"}, {"备注": "空"}),
        ],
        variables=["主表.状态"],
    )
    outcomes = run_tests(rs)
    assert [o.passed for o in outcomes] == [True, True, False]
    assert outcomes[2].problems == ["备注: 期望 空, 实际 旧"]

    ns = dict(ns_two_rows)
    add_derived(make_rs([], [], derived=[("占位", "空")]), ns, ns["主表.金额"].index, lenient=False)
    assert ns["占位"].isna().all()


def test_clear_takes_part_in_priority_merge(ns_two_rows):
    # 清空 和普通输出一样参与逐列合并: 优先级数字小的 清空 赢过数字大的写值, 反过来也一样
    r1 = Rule("R1", 1, [cell_condition("本月.存在", "是")], {"x": "清空"})
    r2 = Rule("R2", 2, [cell_condition("本月.存在", "是")], {"x": "写了"})
    res, _ = apply_rules(make_rs([r1, r2], ["x"]), ns_two_rows, {"x": pd.Series(["old"] * 3)})
    assert res["x"].tolist() == [None, "old", None]
    swapped = [Rule("R1", 2, r1.conditions, r1.outputs), Rule("R2", 1, r2.conditions, r2.outputs)]
    res, _ = apply_rules(make_rs(swapped, ["x"]), ns_two_rows, {"x": pd.Series(["old"] * 3)})
    assert res["x"].tolist() == ["写了", "old", "写了"]


def test_first_match_policy_stops_at_first_hit(ns_two_rows):
    """首次命中 (issue #3): 每行只听优先级数字最小的那条; 没有输出的规则也能拦住一行。"""
    guard = Rule("G", 1, [cell_condition("主表.状态", "活跃")], {})  # 只拦截, 什么都不改
    r2 = Rule("R2", 2, [cell_condition("本月.存在", "是")], {"x": "来自R2", "y": "来自R2"})
    r3 = Rule("R3", 3, [], {"x": "兜底"})  # 无条件
    originals = {"x": pd.Series(["old"] * 3), "y": pd.Series(["old"] * 3)}
    res, hit = apply_rules(make_rs([r3, r2, guard], ["x", "y"], hit_policy="首次命中"), ns_two_rows, originals)
    # 行 0: 活跃 -> G 拦住, 不改; 行 1: 休眠, 本月不存在 -> R3 兜底; 行 2: 状态空, 本月存在 -> R2
    assert res["x"].tolist() == ["old", "兜底", "来自R2"]
    assert res["y"].tolist() == ["old", "old", "来自R2"]
    assert hit.tolist() == ["G", "R3", "R2"]  # 只记录赢的那条
    # 同一组规则换成 优先级 策略: G 不占任何列, 行 0 的 x 被 R2 写、y 也被 R2 写
    res, hit = apply_rules(make_rs([r3, r2, guard], ["x", "y"]), ns_two_rows, originals)
    assert res["x"].tolist() == ["来自R2", "兜底", "来自R2"]
    assert hit.tolist() == ["G; R2; R3", "R3", "R2; R3"]


def test_first_match_ties_follow_definition_order(ns_two_rows):
    a = Rule("A", 5, [], {"x": "A"})
    b = Rule("B", 5, [], {"x": "B"})
    res, hit = apply_rules(make_rs([b, a], ["x"], hit_policy="首次命中"), ns_two_rows, {})
    assert res["x"].tolist() == ["B"] * 3 and hit.tolist() == ["B"] * 3


def test_join_policy_concatenates_all_hits_in_priority_order(ns_two_rows):
    """多条拼接 (issue #16, DMN 的 Collect): 每个输出列把所有命中规则的值按优先级顺序用分隔符接起来并替换原值;
    没写该列的规则不参与, 算出空值的规则不贡献, 没命中的行不动; 审计列记全部命中。"""
    r1 = Rule("R1", 1, [cell_condition("主表.状态", "活跃")], {"标签": "高管", "备注": "a"})
    r2 = Rule("R2", 3, [cell_condition("本月.存在", "是")], {"标签": "直管"})
    big = '=如果(主表.金额 > 6000, "大客户", 空)'
    r3 = Rule("R3", 2, [cell_condition("本月.存在", "是")], {"标签": big, "备注": "-"})
    originals = {"标签": pd.Series(["旧"] * 3), "备注": pd.Series(["旧"] * 3)}
    rs = make_rs([r1, r2, r3], ["标签", "备注"], hit_policy="多条拼接")
    res, hit = apply_rules(rs, ns_two_rows, originals)
    assert res["标签"].tolist() == ["高管/大客户/直管", "旧", "直管"]  # 默认分隔符 /; 行 2 的 R3 算出空, 不贡献
    assert res["备注"].tolist() == ["a", "旧", "旧"]
    assert hit.tolist() == ["R1; R3; R2", "", "R3; R2"]
    rs = make_rs([r1, r2, r3], ["标签", "备注"], hit_policy="多条拼接", settings={"拼接分隔符": "; "})
    res, _ = apply_rules(rs, ns_two_rows, originals)
    assert res["标签"].tolist() == ["高管; 大客户; 直管", "旧", "直管"]
    r4 = Rule("R4", 0, [], {"标签": 7.0, "备注": "本月.电话"})  # 数值 / 复制列 一样进拼接, 7.0 写成 7, 空的列值不贡献
    res, _ = apply_rules(make_rs([r1, r4], ["标签", "备注"], hit_policy="多条拼接"), ns_two_rows, originals)
    assert res["标签"].tolist() == ["7/高管", "7", "7"] and res["备注"].tolist() == ["139/a", "138", "旧"]


def _branch_rule_set(policy, second=None):
    from rule_sheet.model import Branch

    r4 = Rule(
        "R4",
        1,
        [cell_condition("主表.状态", "非空")],  # 各分支共有的条件
        {},
        branches=[
            Branch([cell_condition("主表.状态", "活跃")], {"类别": "A-活跃"}),
            Branch([cell_condition("主表.金额", "> 0")], {"类别": "B-有钱", "备注": "有金额"}),
            Branch([], {"类别": "C-其他"}),  # 无条件 = 规则内的兜底
        ],
    )
    r9 = second or Rule("R9", 2, [], {"备注": "兜底"})
    settings = {"拼接分隔符": "/"} if policy == "多条拼接" else {}
    return make_rs([r4, r9], ["类别", "备注"], hit_policy=policy, settings=settings)


def _branch_originals():
    idx = pd.RangeIndex(3)
    return {
        "类别": pd.Series([None] * 3, index=idx, dtype=object),
        "备注": pd.Series(["o1", "o2", "o3"], index=idx, dtype=object),
    }


def test_branches_follow_hit_policy(ns_two_rows):
    """issue #20: 规则内有序分支。行 0 (活跃, 金额 10000) 三个分支都成立, 行 1 (休眠, 金额 0) 只有兜底分支,
    行 2 状态为空不命中。首次命中 / 优先级: 取书写顺序第一个成立的分支; 多条拼接: 所有成立分支的值接起来;
    规则整体计一次, 审计记 R4。"""
    res, hits = apply_rules(_branch_rule_set("首次命中"), ns_two_rows, _branch_originals())
    assert res["类别"].tolist() == ["A-活跃", "C-其他", None] and res["备注"].tolist() == ["o1", "o2", "兜底"]
    assert hits.tolist() == ["R4", "R4", "R9"]
    res, hits = apply_rules(_branch_rule_set("优先级"), ns_two_rows, _branch_originals())
    assert res["类别"].tolist() == ["A-活跃", "C-其他", None]  # 第一个成立的分支不写 备注, 这一列由 R9 补
    assert res["备注"].tolist() == ["兜底", "兜底", "兜底"] and hits.tolist() == ["R4; R9", "R4; R9", "R9"]
    res, hits = apply_rules(_branch_rule_set("多条拼接"), ns_two_rows, _branch_originals())
    assert res["类别"].tolist() == ["A-活跃/B-有钱/C-其他", "C-其他", "o3" if False else None]
    assert res["备注"].tolist() == ["有金额/兜底", "兜底", "兜底"] and hits.tolist() == ["R4; R9", "R4; R9", "R9"]
    only_null = Rule("R9", 2, [cell_condition("主表.状态", "空")], {"备注": "兜底"})
    res, hits = apply_rules(
        _branch_rule_set("唯一", only_null), ns_two_rows, _branch_originals()
    )  # 多个分支成立不算重叠
    assert hits.tolist() == ["R4", "R4", "R9"] and res["类别"].tolist() == ["A-活跃", "C-其他", None]


def test_branches_with_same_value_join_once(ns_two_rows):
    from rule_sheet.model import Branch

    r = Rule(
        "R1",
        1,
        [],
        {},
        branches=[
            Branch([cell_condition("主表.状态", "活跃")], {"类别": "X"}),
            Branch([cell_condition("主表.金额", "> 0")], {"类别": "X"}),
        ],
    )
    rs = make_rs([r], ["类别"], hit_policy="多条拼接")
    res, hits = apply_rules(rs, ns_two_rows, _branch_originals())
    assert res["类别"].tolist() == ["X", None, "X"] and hits.tolist() == [
        "R1",
        "",
        "R1",
    ]  # 行 0 两支都成立, 相同值只记一次


def test_unique_policy_reports_label(ns_two_rows):
    r1 = Rule("R1", 1, [cell_condition("本月.存在", "是")], {"x": "1"})
    r2 = Rule("R2", 2, [expr_condition("主表.金额 > 100")], {"x": "2"})
    rs = make_rs([r1, r2], ["x"], hit_policy="唯一")
    with pytest.raises(ValueError, match="C001"):
        apply_rules(rs, ns_two_rows, {}, labels=pd.Series(["C001", "C002", "C003"]))


def test_duplicate_rule_ids_rejected(ns_two_rows):
    # masks 按 ID 建字典, 重复 ID 会静默丢掉一条规则; 两个加载器都已拦截, 引擎再兜一层
    rules = [Rule("R", 1, [], {"x": "a"}), Rule("R", 2, [], {"x": "b"})]
    with pytest.raises(ValueError, match="规则 ID 重复: R"):
        apply_rules(make_rs(rules, ["x"]), ns_two_rows, {})


def test_unknown_variable_names_rule(ns_two_rows):
    r = Rule("R9", 1, [cell_condition("不存在的变量", "x")], {"x": "1"})
    with pytest.raises(KeyError, match="R9"):
        apply_rules(make_rs([r], ["x"]), ns_two_rows, {})


# ---------------------------------------------------------------------------
# add_derived / run_tests
# ---------------------------------------------------------------------------


def test_derived_respects_given_value_and_lenient(ns_two_rows):
    rs = make_rs([], [], derived=[("偏差率", "abs(本月.金额 - 主表.金额) / 主表.金额"), ("坏的", "没有这个变量 + 1")])
    ns = dict(ns_two_rows)
    ns["偏差率"] = pd.Series([9.9, 9.9, 9.9])
    add_derived(rs, ns, ns["主表.金额"].index, lenient=True)
    assert ns["偏差率"].tolist() == [9.9, 9.9, 9.9]
    assert ns["坏的"].isna().all()
    with pytest.raises(ValueError):
        add_derived(rs, dict(ns_two_rows), ns_two_rows["主表.金额"].index, lenient=False)


def test_run_tests_blank_expected_means_unchanged_and_missing_vars_are_null():
    r1 = Rule("R1", 1, [cell_condition("主表.状态", "活跃"), cell_condition("本月.电话", "非空")], {"备注": "有电话"})
    rs = make_rs(
        [r1],
        ["备注", "状态"],
        tests=[
            TestCase("命中", {"主表.状态": "活跃", "本月.电话": "1"}, {"备注": "有电话"}),
            TestCase("电话没给=空, 不命中", {"主表.状态": "活跃"}, {}),
            TestCase("故意写错", {"主表.状态": "活跃", "本月.电话": "1"}, {"备注": "别的"}),
            TestCase("状态保持原值", {"主表.状态": "休眠"}, {"状态": "休眠"}),
        ],
        variables=["主表.状态", "本月.电话"],
    )
    outcomes = run_tests(rs)
    assert [o.passed for o in outcomes] == [True, True, False, True]
    assert outcomes[0].hit == "R1" and outcomes[1].hit == ""
    assert "备注" in outcomes[2].problems[0]


def test_run_tests_expected_cleared():
    """期望格写 清空 / 为空 = 期望该列为空值; 留空仍是「不变」, 所以能区分「被清空」和「没动」。"""
    r1 = Rule("R1", 1, [cell_condition("主表.状态", "活跃")], {"备注": "清空"})
    rs = make_rs(
        [r1],
        ["备注"],
        tests=[
            TestCase("清空成功", {"主表.状态": "活跃", "主表.备注": "旧"}, {"备注": "清空"}),
            TestCase("也认 为空", {"主表.状态": "活跃", "主表.备注": "旧"}, {"备注": "为空"}),
            TestCase("留空=不变, 抓到被清空", {"主表.状态": "活跃", "主表.备注": "旧"}, {"备注": None}),
            TestCase("没命中, 期望空则失败", {"主表.状态": "休眠", "主表.备注": "旧"}, {"备注": "清空"}),
        ],
        variables=["主表.状态"],
    )
    outcomes = run_tests(rs)
    assert [o.passed for o in outcomes] == [True, True, False, False]
    assert outcomes[2].problems == ["备注: 期望 旧, 实际 空"]
    assert outcomes[3].problems == ["备注: 期望 空, 实际 旧"]


def test_date_columns_written_with_date_format(tmp_path):
    """规则写过之后变成 object 的日期列也按 yyyy-mm-dd 写出; 混着文本的列和全空列走原来的路径 (不当日期列)。"""
    idx = pd.RangeIndex(3)
    df = pd.DataFrame(
        {
            "全日期": pd.Series(
                [pd.Timestamp("2026-10-20"), None, pd.Timestamp("2026-09-20")], index=idx, dtype=object
            ),
            "混合": pd.Series([pd.Timestamp("2026-10-20"), "待定", None], index=idx, dtype=object),
            "全空": pd.Series([None, None, None], index=idx, dtype=object),
            "原生": pd.to_datetime(pd.Series(["2026-10-20", None, "2026-09-20"], index=idx)),
        }
    )
    assert [_is_date_column(df[c]) for c in df.columns] == [True, False, False, True]
    out = write_xlsx_safely(df, tmp_path / "o.xlsx")
    ws = load_workbook(out)["Sheet1"]
    formats = {ws.cell(1, c).value: ws.cell(2, c).number_format for c in range(1, 5)}
    assert formats["全日期"] == "yyyy-mm-dd" and formats["原生"] == "yyyy-mm-dd"
    assert formats["混合"] != "yyyy-mm-dd" and formats["全空"] != "yyyy-mm-dd"


def test_output_column_widths_fit_all_rows(tmp_path):
    """输出文件列宽和 Excel 里双击列边一样: 看全部行 (以前只看前 50 行、上限 60, 后面的长值被右边的列挡住)。"""
    long_text = "华东区域-上海-浦东新区-陆家嘴金融贸易区-某某大厦-第三十八层-综合管理部-行政与后勤保障组"
    df = pd.DataFrame({"岗位": ["短"] * 60 + [long_text], "编码": ["P1"] * 61})
    ws = load_workbook(write_xlsx_safely(df, tmp_path / "o.xlsx"))["Sheet1"]
    assert ws.column_dimensions["A"].width >= display_width(long_text) > 60
    assert ws.column_dimensions["B"].width == 8  # 短列不低于 8
