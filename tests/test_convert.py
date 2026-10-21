"""Excel -> YAML 导出的无损性 (逐形态行为等价 + 三个样例端到端), 以及两份规则文件的语义比对。"""

import shutil
from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from rule_sheet import cli
from rule_sheet.conditions import to_mask, unary_test
from rule_sheet.convert import CellTranslator, ExprToCell, diff_rules, excel_to_yaml, export_workbook
from rule_sheet.engine import run, run_tests
from rule_sheet.expressions import eval_expr
from rule_sheet.loaders import load_rules
from rule_sheet.model import InputSpec, RuleSet

from .conftest import (
    CAND_DIR,
    CAND_EXCEL,
    CAND_YAML,
    CARRY_DIR,
    CARRY_EXCEL,
    CARRY_YAML,
    DIFF_DIR,
    DIFF_EXCEL,
    DIFF_YAML,
    EXCEL_RULES,
    STATUS_DIR,
    STATUS_EXCEL,
    STATUS_YAML,
    TAGS_DIR,
    TAGS_EXCEL,
    TAGS_YAML,
    YAML_RULES,
)

SAMPLES = pytest.mark.parametrize(
    ("excel", "yaml_file", "data_dir"),
    [
        (EXCEL_RULES, YAML_RULES, EXCEL_RULES.parent),
        (STATUS_EXCEL, STATUS_YAML, STATUS_DIR),
        (TAGS_EXCEL, TAGS_YAML, TAGS_DIR),
        (CARRY_EXCEL, CARRY_YAML, CARRY_DIR),
        (CAND_EXCEL, CAND_YAML, CAND_DIR),
        (DIFF_EXCEL, DIFF_YAML, DIFF_DIR),
    ],
    ids=["customer", "status", "tags", "carry", "cand", "diff"],
)


@pytest.fixture
def translator():
    rs = RuleSet(
        inputs=[InputSpec("客户", primary=True), InputSpec("本月", keys=[("客户编号", "客户编号")])],
        settings={},
        derived=[("拟任人有变化", "客户.电话 != 本月.电话")],
        hit_policy="优先级",
        rules=[],
        outputs=[],
    )
    return CellTranslator(rs)


@pytest.fixture
def ns():
    idx = pd.RangeIndex(5)
    return {
        "客户.状态": pd.Series(["活跃", "休眠", None, "VIP", "是"], index=idx, dtype=object),
        "偏差率": pd.Series([0.08, 0.03, None, 11.0, 0.05], index=idx),
        "本月.存在": pd.Series([True, False, True, False, True], index=idx),
        "客户.电话": pd.Series(["138", None, "139", "007", "x"], index=idx, dtype=object),
        "本月.电话": pd.Series(["139", "138", None, "007", "x"], index=idx, dtype=object),
        "客户.上次活跃": pd.Series(
            pd.to_datetime(["2026-10-09", "2026-10-08", None, "2026-01-01", "2026-10-08"]), index=idx
        ),
        "拟任人有变化": pd.Series([True, False, True, False, True], index=idx),
        "基准日期": pd.Timestamp("2026-10-19"),
    }


@pytest.mark.parametrize(
    ("var", "cell", "expr"),
    [
        ("客户.状态", "活跃", '客户.状态 == "活跃"'),
        ("客户.状态", "!= 休眠", '客户.状态 != "休眠"'),
        ("客户.状态", "<> 休眠", '客户.状态 != "休眠"'),
        ("客户.状态", "= 活跃", '客户.状态 == "活跃"'),
        ("客户.状态", "VIP, 活跃", '属于(客户.状态, "VIP", "活跃")'),
        ("客户.状态", "活跃, 空", '属于(客户.状态, "活跃", 空)'),
        ("客户.状态", '"VIP 小幅偏差, 关注"', '客户.状态 == "VIP 小幅偏差, 关注"'),
        ("客户.状态", "= 空", "客户.状态 == 空"),
        ("客户.状态", "!= 空", "客户.状态 != 空"),
        ("客户.状态", "空", "客户.状态 == 空"),
        ("客户.状态", "为空", "客户.状态 == 空"),
        ("客户.状态", "非空", "客户.状态 != 空"),
        ("客户.状态", "not(活跃)", '客户.状态 != "活跃"'),
        ("客户.状态", "not(VIP, 活跃)", '客户.状态 not in ("VIP", "活跃")'),
        ("客户.状态", "not(空)", "客户.状态 != 空"),
        ("客户.状态", "not(为空)", "客户.状态 != 空"),
        ("客户.状态", "not(非空)", "客户.状态 == 空"),
        ("客户.状态", "not(包含 活)", '客户.状态 != 空 且 非 (包含(客户.状态, "活"))'),
        ("客户.状态", "包含 活, 休", '包含(客户.状态, "活", "休")'),
        ("客户.状态", "开头是(活)", '开头是(客户.状态, "活")'),
        ("客户.状态", "不包含 眠", '不包含(客户.状态, "眠")'),
        ("客户.状态", "是", '客户.状态 == "是"'),  # 文本列上的 是 就是文字
        ("偏差率", "> 0.05", "偏差率 > 0.05"),
        ("偏差率", ">= 0.05", "偏差率 >= 0.05"),
        ("偏差率", "[0.02..0.05]", "0.02 <= 偏差率 <= 0.05"),
        ("偏差率", "(0.02..0.05]", "0.02 < 偏差率 <= 0.05"),
        ("偏差率", "[0.02..0.05)", "0.02 <= 偏差率 < 0.05"),
        ("偏差率", "> 10, 为空", "偏差率 > 10 或 偏差率 == 空"),
        ("偏差率", "not(> 0.05)", "偏差率 != 空 且 非 (偏差率 > 0.05)"),
        ("偏差率", 0.05, "偏差率 == 0.05"),
        ("偏差率", 11, "偏差率 == 11"),
        ("偏差率", "11", "偏差率 == 11"),
        ("本月.存在", "是", "本月.存在 == 是"),
        ("本月.存在", "否", "本月.存在 == 否"),
        ("本月.存在", True, "本月.存在 == 是"),
        ("拟任人有变化", "是", "拟任人有变化 == 是"),
        ("本月.电话", "!= 客户.电话", "本月.电话 != 客户.电话"),
        ("本月.电话", "= 客户.电话", "本月.电话 == 客户.电话"),
        ("本月.电话", "客户.电话", "本月.电话 == 客户.电话"),
        ("客户.电话", "007", "客户.电话 == 7"),  # 和单元格一样: 裸的 007 经 literal 变成 7, 要保文本就加引号
        ("客户.电话", '"007"', '客户.电话 == "007"'),
        ("客户.电话", "v1.2", '客户.电话 == "v1.2"'),
        ("客户.上次活跃", "> 2026-10-08", '客户.上次活跃 > 日期("2026-10-08")'),
        ("客户.上次活跃", "2026-10-08", '客户.上次活跃 == 日期("2026-10-08")'),
        ("客户.上次活跃", datetime(2026, 10, 8), '客户.上次活跃 == 日期("2026-10-08")'),
        ("客户.状态", "-", None),
        ("客户.状态", None, None),
        ("客户.状态", "任意", None),
        ("客户.状态", "活跃, -", None),
    ],
)
def test_cell_translation_matches_unary_test(translator, ns, var, cell, expr):
    """每种单元格形态翻出来的表达式, 在含空值的数据上与 unary_test 的 mask 完全一致 (这是「无损」的定义)。"""
    assert translator.cell(var, cell) == expr
    want = unary_test(cell, ns[var], ns).tolist()
    got = to_mask(eval_expr(expr, ns), ns).tolist() if expr is not None else [True] * 5
    assert got == want, (cell, expr)


def test_names_with_special_chars_get_backticks(translator):
    assert translator.cell("客户.金额(万元)", "> 100") == "`客户.金额(万元)` > 100"
    assert translator.cell("客户.类别-2024", "x") == '`客户.类别-2024` == "x"'
    assert translator.cell("客户.类别-会议", "x") == '客户.类别-会议 == "x"'
    assert translator.cell("客户.岗位/职级", "x") == '客户.岗位/职级 == "x"'
    assert translator.cell("客户.电话", "!= 客户.金额(万元)") == "客户.电话 != `客户.金额(万元)`"
    assert translator.cell("客户.电话", '= "a\'b"') == '客户.电话 == "a\'b"'


@SAMPLES
def test_export_is_lossless_end_to_end(excel, yaml_file, data_dir, tmp_path):
    """导出件自己的用例全过; 导出件跑真实数据与 Excel 结果相同; 与手写 YAML 语义一致 (只有输出路径这种提示)。"""
    out = export_workbook(excel, tmp_path / "exported.yaml")
    text = out.read_text(encoding="utf-8")
    assert text.startswith("# 由 rule-sheet export 从") and "输入:" in text and "规则:" in text
    rs = load_rules(out)
    outcomes = run_tests(rs)
    failed = [(o.name, o.problems, o.error) for o in outcomes if not o.passed]
    assert outcomes and not failed, failed
    a = run(excel, dry_run=True).out
    b = run(out, base_dir=data_dir, dry_run=True).out
    cols = list(load_rules(excel).outputs) + ["命中规则"]
    pd.testing.assert_frame_equal(a[cols], b[cols], check_dtype=False)
    res = diff_rules(excel, out)
    assert res.same and res.cases == len(rs.tests), res.differences
    res = diff_rules(out, yaml_file)  # 手写的 YAML 样例带注释, 不由导出覆盖; 靠 diff 守住两者一致
    assert res.same, res.differences
    assert all("输出" in n or "默认文件" in n for n in res.notes)


@SAMPLES
def test_hand_written_yaml_matches_excel(excel, yaml_file, data_dir):
    """这就是主张的同步机制: Excel 是源, 手写 / 生成的 YAML 用 diff 对账。"""
    res = diff_rules(excel, yaml_file)
    assert res.same, res.differences


def test_diff_reports_each_kind_of_change(tmp_path):
    src = STATUS_YAML.read_text(encoding="utf-8")
    p = tmp_path / "changed.yaml"
    s3_cond = "- 本期.类别-会议 == 上期.类别-会议      # 上期没有"
    changed = (
        src.replace("优先级: 4\n", "优先级: 40\n", 1)
        .replace("状态-会议: 03-按计划进行\n", "状态-会议: 03-按计划\n", 1)
        .replace(s3_cond, s3_cond.replace("==", "!="), 1)
    )
    assert changed != src and s3_cond in src
    p.write_text(changed, encoding="utf-8")
    res = diff_rules(STATUS_EXCEL, p)
    text = "\n".join(res.differences)
    assert "规则 S4 优先级: A 4, B 40" in text
    assert "规则 S4 输出: A {'状态-会议': '03-按计划进行'}, B {'状态-会议': '03-按计划'}" in text
    assert "用例「T3」命中规则: A S3, B" in text  # S3 的条件改了, 结构比不出来, 交叉验证能抓到
    assert any("输出「状态-会议」" in d for d in res.differences)
    p.write_text(src.replace("  - 编号: S5\n", "  - 编号: S6\n", 1), encoding="utf-8")
    res = diff_rules(STATUS_EXCEL, p)
    assert "规则 S5 只在 A 里" in res.differences and "规则 S6 只在 B 里" in res.differences


def test_export_cli_both_directions(tmp_path, capsys):
    shutil.copy(YAML_RULES, tmp_path / "rules.yaml")
    assert cli.main(["export", str(tmp_path / "rules.yaml")]) == 0
    out = capsys.readouterr().out
    assert "已导出" in out and "rules.xlsx" in out and (tmp_path / "rules.xlsx").exists()
    src = tmp_path / "r.xlsx"
    shutil.copy(TAGS_EXCEL, src)
    assert cli.main(["export", str(src)]) == 0
    out = capsys.readouterr().out
    assert "已导出" in out and "(6 条规则, 8 条用例)" in out and (tmp_path / "r.yaml").exists()
    assert cli.main(["diff", str(src), str(tmp_path / "r.yaml")]) == 0
    assert "语义一致" in capsys.readouterr().out
    assert cli.main(["diff", str(src), str(STATUS_YAML)]) == 1
    assert "发现" in capsys.readouterr().out
    assert cli.main(["export", str(tmp_path / "nope.txt")]) == 2


# ---------------------------------------------------------------------------
# YAML -> Excel: 表格子集
# ---------------------------------------------------------------------------


@pytest.fixture
def to_cell():
    rs = RuleSet(
        inputs=[InputSpec("客户", primary=True), InputSpec("本月", keys=[("客户编号", "客户编号")])],
        settings={},
        derived=[("拟任人有变化", "客户.电话 != 本月.电话"), ("偏差率", "客户.金额 * 2")],
        hit_policy="优先级",
        rules=[],
        outputs=[],
        variables=["客户.状态", "客户.电话", "本月.电话", "客户.上次活跃", "客户.等级", "本月.存在", "偏差率"],
    )
    return ExprToCell(rs)


@pytest.mark.parametrize(
    ("expr", "cells"),
    [
        ('客户.状态 == "活跃"', [("客户.状态", "活跃")]),
        ('客户.状态 != "休眠"', [("客户.状态", "!= 休眠")]),
        ('"活跃" == 客户.状态', [("客户.状态", "活跃")]),
        ("偏差率 > 0.05", [("偏差率", "> 0.05")]),
        ("0.05 < 偏差率", [("偏差率", "> 0.05")]),
        ("偏差率 >= 11", [("偏差率", ">= 11")]),
        ("0.02 <= 偏差率 <= 0.05", [("偏差率", "[0.02..0.05]")]),
        ("0.02 < 偏差率 <= 0.05", [("偏差率", "(0.02..0.05]")]),
        ("0.05 >= 偏差率 > 0.02", [("偏差率", "(0.02..0.05]")]),
        ('属于(客户.状态, "VIP", "活跃")', [("客户.状态", "VIP, 活跃")]),
        ('属于(客户.状态, "活跃", 空)', [("客户.状态", "活跃, 空")]),
        ('客户.状态 in ("VIP", "活跃")', [("客户.状态", "VIP, 活跃")]),
        ('客户.状态 not in ("VIP", "活跃")', [("客户.状态", "not(VIP, 活跃)")]),
        ("本月.电话 == 空", [("本月.电话", "= 空")]),
        ("本月.电话 != 空", [("本月.电话", "!= 空")]),
        ("为空(本月.电话)", [("本月.电话", "为空")]),
        ("非空(本月.电话)", [("本月.电话", "非空")]),
        ("非 (本月.电话 == 空)", [("本月.电话", "非空")]),
        ("本月.存在 == 是", [("本月.存在", "是")]),
        ("本月.存在 == 否", [("本月.存在", "否")]),
        ("本月.存在 != 是", [("本月.存在", "否")]),
        ("本月.存在", [("本月.存在", "是")]),
        ("非 本月.存在", [("本月.存在", "否")]),
        ("拟任人有变化", [("拟任人有变化", "是")]),
        ("本月.电话 != 客户.电话", [("本月.电话", "!= 客户.电话")]),
        ("本月.电话 == 客户.电话", [("本月.电话", "= 客户.电话")]),
        ('包含(客户.状态, "活", "休")', [("客户.状态", "包含 活, 休")]),
        ('不包含(客户.状态, "眠")', [("客户.状态", "不包含 眠")]),
        ('开头是(客户.状态, "活")', [("客户.状态", "开头是 活")]),
        ('客户.上次活跃 > 日期("2026-10-08")', [("客户.上次活跃", "> 2026-10-08")]),
        ('客户.上次活跃 == 日期("2026-10-08")', [("客户.上次活跃", "2026-10-08")]),
        ('客户.状态 == "007"', [("客户.状态", '"007"')]),  # 裸 007 会被单元格当成数字, 必须加引号
        ('客户.状态 == "2026-10-08"', [("客户.状态", '"2026-10-08"')]),
        ('客户.状态 == "VIP 小幅偏差, 关注"', [("客户.状态", '"VIP 小幅偏差, 关注"')]),
        ('客户.状态 == "是"', [("客户.状态", '"是"')]),
        ('客户.状态 == "空"', [("客户.状态", '"空"')]),
        ('客户.状态 == "-"', [("客户.状态", '"-"')]),
        ('客户.状态 == "> 1"', [("客户.状态", '"> 1"')]),
        ('客户.状态 != 空 且 非 (包含(客户.状态, "活"))', [("客户.状态", "not(包含 活)")]),
        ("偏差率 != 空 且 非 (偏差率 > 0.05)", [("偏差率", "not(> 0.05)")]),
        ("偏差率 > 10 或 偏差率 == 空", [("偏差率", "> 10, = 空")]),
        ('客户.状态 == "活跃" 且 偏差率 > 0.05', [("客户.状态", "活跃"), ("偏差率", "> 0.05")]),
    ],
)
def test_expr_to_cells(to_cell, ns, expr, cells):
    """表格子集里的每种表达式都有对应的单元格, 且翻回去的 mask 与原表达式在含空值的数据上一致。"""
    assert to_cell.cells(expr, "规则 R") == cells
    want = to_mask(eval_expr(expr, ns), ns).tolist()
    got = [True] * 5
    for var, cell in cells:
        got = [g and h for g, h in zip(got, unary_test(cell, ns[var], ns).tolist(), strict=True)]
    assert got == want, (expr, cells)


@pytest.mark.parametrize(
    ("expr", "why"),
    [
        ("abs(客户.金额 - 本月.金额) > 1", "比较的一侧必须是变量"),
        ("偏差率 * 2 > 1", "比较的一侧必须是变量"),
        ("偏差率 * 2", "不是一元测试"),
        ('非 (客户.状态 == "活跃")', "纯取反"),
        ("偏差率 > 0.05 或 本月.存在", "跨变量的或"),
        ("偏差率 > 空", "空 只能和 == / != 搭配"),
        ("客户.状态 == 是", "是 / 否 比较只对"),
        ("客户.状态", "裸变量名"),
        ("天数差(基准日期, 客户.上次活跃) > 1", "比较的一侧必须是变量"),
        ("天数差(客户.上次活跃, 基准日期)", "不能直接当单元格条件"),
        ("0.02 <= 偏差率 >= 0.01", "区间"),
        ("偏差率 >", "表达式语法错误"),
    ],
)
def test_expr_outside_subset_is_rejected(to_cell, expr, why):
    with pytest.raises(ValueError, match=why):
        to_cell.cells(expr, "规则 R")


def test_branches_merge_ranges_and_split_alternatives(to_cell):
    assert to_cell.branches(["偏差率 >= 0.02", "偏差率 <= 0.05"], "R") == [{"偏差率": "[0.02..0.05]"}]
    assert to_cell.branches(["偏差率 > 0.02", "偏差率 < 0.05"], "R") == [{"偏差率": "(0.02..0.05)"}]
    cond = ["本月.存在 == 是", {"任一满足": ["本月.电话 == 空", ["客户.电话 == 空", "偏差率 > 1"]]}]
    rows = to_cell.branches(cond, "R")
    assert rows == [{"本月.存在": "是", "本月.电话": "= 空"}, {"本月.存在": "是", "客户.电话": "= 空", "偏差率": "> 1"}]
    assert to_cell.branches({"不满足": "本月.电话 == 空"}, "R") == [{"本月.电话": "非空"}]
    assert to_cell.branches("本月.存在", "R") == [{"本月.存在": "是"}]
    with pytest.raises(ValueError, match="两个条件"):
        to_cell.branches(["偏差率 > 0.02", "偏差率 > 0.05"], "R")
    with pytest.raises(ValueError, match="只能有一个顶层 任一满足"):
        to_cell.branches([{"任一满足": ["本月.存在"]}, {"任一满足": ["拟任人有变化"]}], "R")
    with pytest.raises(ValueError, match="嵌套在条件里面的 任一满足"):
        to_cell.branches([{"任一满足": [{"任一满足": ["本月.存在"]}]}], "R")
    with pytest.raises(ValueError, match="不满足: 是纯取反"):
        to_cell.branches({"不满足": '客户.状态 == "活跃"'}, "R")


@SAMPLES
def test_yaml_to_excel_round_trips(excel, yaml_file, data_dir, tmp_path):
    """手写 YAML -> Excel: 结构能加载、内嵌用例全过、与手写 Excel 和 YAML 都语义一致、跑真实数据结果相同;
    再 Excel -> YAML -> Excel 一圈仍然一致。"""
    out = export_workbook(yaml_file, tmp_path / "from_yaml.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames[0] == "输入" and "决策表" in wb.sheetnames and "测试" in wb.sheetnames
    assert wb["决策表"].freeze_panes and wb["决策表"].cell(1, 1).value == "命中策略"  # 版式已套
    rs = load_rules(out)
    failed = [(o.name, o.problems, o.error) for o in run_tests(rs) if not o.passed]
    assert rs.tests and not failed, failed
    assert diff_rules(out, excel).same, diff_rules(out, excel).differences
    assert diff_rules(out, yaml_file).same, diff_rules(out, yaml_file).differences
    a = run(yaml_file, dry_run=True).out
    b = run(out, base_dir=data_dir, dry_run=True).out
    cols = list(rs.outputs) + ["命中规则"]
    pd.testing.assert_frame_equal(a[cols], b[cols], check_dtype=False)
    again = export_workbook(export_workbook(excel, tmp_path / "a.yaml"), tmp_path / "b.xlsx")
    assert diff_rules(again, excel).same


def test_yaml_to_excel_keeps_continuation_rows_merged(tmp_path):
    out = export_workbook(STATUS_YAML, tmp_path / "s.xlsx")
    ws = load_workbook(out)["决策表"]
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert merged == {"A6:A7", "B6:B7"}  # S3 的 任一满足 两支 -> 两行, 规则ID / 优先级 合并
    assert [ws.cell(r, 1).value for r in range(4, 10)] == ["S1", "S2", "S3", None, "S4", "S5"]
    assert ws.cell(6, 3).value == "= 上期.类别-会议" or ws.cell(6, 4).value == "= 上期.类别-会议"


def test_field_groups_survive_both_directions_and_diff(tmp_path):
    """导出保留字段组的形态 (字段组: 键 + 则: 里的 组名: 输入名), YAML -> Excel 生成「字段组」sheet, diff 比字段组。"""
    text = excel_to_yaml(CARRY_EXCEL)
    assert "字段组:\n  会议记录: [例会时间, 类别-会议, 状态-会议, 例会结论, 备注, 待跟进事宜]\n" in text
    assert (
        "会议记录: 上期" in text and "例会结论: 上期.例会结论" not in text
    )  # 则: 里不展开 (用例输入里当然有 上期.例会结论)
    out = export_workbook(CARRY_YAML, tmp_path / "c.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["输入", "设置", "字段组", "决策表", "测试"]
    assert [wb["字段组"].cell(2, c).value for c in (1, 2)] == [
        "会议记录",
        "例会时间, 类别-会议, 状态-会议, 例会结论, 备注, 待跟进事宜",
    ]
    ws = wb["决策表"]
    names = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    assert names[names.index("会议记录") - 1] == "上期.状态-会议" and names[-4:] == [
        "会议记录",
        "例会时间",
        "状态-会议",
        "说明",
    ]
    assert (
        ws.cell(4, names.index("会议记录") + 1).value == "上期"
        and ws.cell(5, names.index("会议记录") + 1).value is None
    )
    tests_ws = wb["测试"]
    expects = [
        tests_ws.cell(2, c).value for c in range(1, tests_ws.max_column + 1) if tests_ws.cell(1, c).value == "期望"
    ]
    assert expects == ["例会时间", "类别-会议", "状态-会议", "例会结论", "备注", "待跟进事宜"]  # 期望列是展开后的逐列
    p = tmp_path / "g.yaml"
    p.write_text(
        CARRY_YAML.read_text(encoding="utf-8").replace(", 待跟进事宜]", ", 待跟进事宜, 岗位名称]"), encoding="utf-8"
    )
    res = diff_rules(CARRY_EXCEL, p)
    assert any(d.startswith("字段组不同") for d in res.differences)
    assert any("规则 R1 输出" in d and "上期.岗位名称" in d for d in res.differences)  # 展开后的逐列输出也跟着不同


def test_concat_input_options_survive_both_directions(tmp_path):
    text = excel_to_yaml(CAND_EXCEL)
    assert "推荐: {文件: 推荐记录.xlsx, 匹配: 岗位编码, 重复键: 拼接, 排序: 推荐时间 降序, 空值占位: NA}" in text
    assert (
        "冻结: {文件: 冻结记录.xlsx, 匹配: 工号 = 岗位.现任工号, 重复键: 拼接, 排序: 冻结开始 降序, 带序号: 是}" in text
    )
    assert "推荐.现岗: 拼接(推荐.现部门" in text  # 带前缀的派生变量照常导出
    out = export_workbook(CAND_YAML, tmp_path / "c.xlsx")
    ws = load_workbook(out)["输入"]
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    assert header == ["名称", "匹配", "重复键", "排序", "分隔符", "空值占位", "带序号", "文件", "说明"]
    rows = {ws.cell(r, 1).value: [ws.cell(r, c).value for c in range(2, 9)] for r in (2, 3, 4)}
    assert rows["推荐"] == ["岗位编码", "拼接", "推荐时间 降序", None, "NA", None, "推荐记录.xlsx"]
    assert rows["冻结"] == ["工号 = 岗位.现任工号", "拼接", "冻结开始 降序", None, None, "是", "冻结记录.xlsx"]
    p = tmp_path / "g.yaml"
    p.write_text(CAND_YAML.read_text(encoding="utf-8").replace("带序号: 是", "带序号: 否"), encoding="utf-8")
    res = diff_rules(CAND_EXCEL, p)
    assert any(d.startswith("输入「冻结」的选项") for d in res.differences)


def test_append_option_survives_both_directions(tmp_path):
    src = tmp_path / "a.yaml"
    src.write_text(
        "输入:\n  本期: {文件: a.xlsx, 匹配: 主表}\n  上期: {文件: b.xlsx, 匹配: 编码, 独有行: 追加}\n"
        "规则:\n  - 编号: X1\n    当: [本期.存在 == 否]\n    则: {差异类型: REMOVED}\n"
        "测试:\n  - 用例: T1\n    输入: {本期.存在: 否}\n    期望: {差异类型: REMOVED}\n",
        encoding="utf-8",
    )
    out = export_workbook(src, tmp_path / "a.xlsx")
    ws = load_workbook(out)["输入"]
    assert [ws.cell(1, c).value for c in range(1, 7)] == ["名称", "匹配", "重复键", "独有行", "文件", "说明"]
    assert [ws.cell(3, c).value for c in range(1, 6)] == ["上期", "编码", "取第一条", "追加", "b.xlsx"]
    rs = load_rules(out)
    assert rs.lookups[0].append is True and rs.variables == ["本期.存在"]
    assert [o.passed for o in run_tests(rs)] == [True]
    assert "独有行: 追加}" in excel_to_yaml(out)
    assert diff_rules(src, out).same
    src.write_text(src.read_text(encoding="utf-8").replace(", 独有行: 追加", ""), encoding="utf-8")
    assert any(d.startswith("输入「上期」的选项") for d in diff_rules(src, out).differences)


def test_yaml_outside_subset_reports_rule(tmp_path):
    p = tmp_path / "r.yaml"
    p.write_text(
        "输入:\n  客户: {匹配: 主表}\n  本月: {匹配: k}\n规则:\n"
        "  - 编号: X\n    当: [abs(本月.金额 - 客户.金额) > 1]\n    则: {y: 1}\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="规则 X: 条件 .* 转不成决策表单元格 .*变量: 里起个名字"):
        export_workbook(p, tmp_path / "r.xlsx")
    with pytest.raises(ValueError, match="export 只认"):
        export_workbook(tmp_path / "r.txt")


def test_exported_yaml_round_trips_through_loader(tmp_path):
    """导出的 YAML 再被加载时, 规则 ID / 优先级 / 输出 / 用例与 Excel 一致 (条件靠上面的端到端测试)。"""
    out = export_workbook(STATUS_EXCEL, tmp_path / "s.yaml")
    a, b = load_rules(STATUS_EXCEL), load_rules(out)
    assert [(r.rule_id, r.priority, r.desc) for r in a.rules] == [(r.rule_id, r.priority, r.desc) for r in b.rules]
    assert [t.name for t in a.tests] == [t.name for t in b.tests]
    assert b.hit_policy == "首次命中" and b.outputs == ["状态-会议"]
    text = out.read_text(encoding="utf-8")
    assert "任一满足:" in text  # S3 的两行续行
    assert excel_to_yaml(STATUS_EXCEL) == excel_to_yaml(STATUS_EXCEL)  # 确定性
