"""生成「跨批次带入」样例: 两份数据 + Excel 决策表, 全部写到本脚本所在目录。

业务: 每个批次一张全量岗位表 (本期), 冒泡的岗位要上会跟踪; 上一批次的会议记录 (上期) 要按岗位编码带入本期。
会议记录有六列 (例会时间 / 类别-会议 / 状态-会议 / 例会结论 / 备注 / 待跟进事宜), 三条规则里两条都要整组带入。

这个样例展示 0.7 加的「字段组」:
    - 「字段组」sheet 给这六列起个名字 会议记录; 决策表里以 会议记录 为列名的输出格一格就是整组:
      写 上期 = 组内每列从上期的同名列复制, 写 清空 = 整组清空, - = 不改
    - 同一条规则显式写了的列以显式为准: R1 整组带入但 例会时间 改成本批例会日, R3 整组带入但 状态-会议 标成 06
    - 带入是整组同步: 上期某列为空, 本期该列也会被清成空 (K005 的备注), 变更明细里能看到
    - 「本批冒泡」是文本列, 单元格里的 是 / 否 就是文字 (只有 上期.存在 这类布尔列上的 是 / 否 才是布尔)

运行:
    uv run python examples/carry-forward/make_samples.py

这里只写内容; 颜色 / 列宽 / 换行 / 冻结等版式统一由 rule_sheet.style 套用 (与 `rule-sheet format` 同一套规范)。

产物:
    本期全量.xlsx      被更新的主表 (规则里叫 本期)
    上期全量.xlsx      参照表 (规则里叫 上期)
    规则_决策表.xlsx   四张 sheet: 设置 (含输入表) / 字段组 / 决策表 / 测试 (没有变量定义, 它是可选的)
"""
from __future__ import annotations

from collections.abc import Iterable
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from rule_sheet.style import style_data_sheet, style_rules_workbook

ROOT = Path(__file__).resolve().parent


def new_sheet(wb: Workbook, title: str, rows: Iterable[Iterable[Any]], *, first: bool = False) -> Worksheet:
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    for row in rows:
        ws.append(list(row))
    return ws


# ---------------------------------------------------------------------------
# 样例数据: 每一行对应一种情形, 便于对照决策表看命中结果
# ---------------------------------------------------------------------------
NOTICE = "变动通报"
PLAN = "近2个月酝酿及上会计划"
LAST = datetime(2026, 9, 20)   # 上批例会日
BASE = datetime(2026, 10, 20)  # 本批例会日 = 设置里的基准日期

CURRENT_HEADER = ["岗位编码", "岗位名称", "本批冒泡", "例会时间", "类别-会议", "状态-会议", "例会结论", "备注", "待跟进事宜"]
CURRENT_ROWS = [
    # 编号    名称            冒泡  例会时间 类别     状态          结论        备注       待跟进   上期情况                      预期
    ["K001", "财务部-部长",   "是", None,  None,    None,         None,      None,      None],  # 上期有 02-进展更新           -> R1 整组带入, 例会时间改成本批
    ["K002", "人事部-副部长", "是", None,  None,    None,         None,      None,      None],  # 上期无                       -> R2 首次冒泡: 例会时间 + 01-新增
    ["K003", "研发中心-主任", "否", None,  None,    None,         None,      None,      None],  # 上期有 03-按计划进行           -> R3 整组带入 (例会时间仍是上批的), 状态标 06
    ["K004", "市场部-经理",   "否", None,  None,    None,         None,      None,      None],  # 上期无                       -> 不命中
    ["K005", "法务部-主管",   "是", LAST,  NOTICE,  "02-进展更新", "候选人面谈", "临时备注", None],  # 上期有, 上期备注为空          -> R1 整组覆盖: 备注被清成空
    ["K006", "采购部-经理",   "否", None,  None,    None,         None,      None,      None],  # 上期有但 05-已关闭            -> R3 不命中 (条件 != 05-已关闭)
    ["K007", "质量部-部长",   "是", None,  None,    None,         None,      None,      None],  # 上期有 01-新增, 其余列为空     -> R1 带入, 空列仍是空
]

PREVIOUS_HEADER = ["岗位编码", "例会时间", "类别-会议", "状态-会议", "例会结论", "备注", "待跟进事宜"]
PREVIOUS_ROWS = [
    ["K001", LAST,                    NOTICE, "02-进展更新",  "同意推进", "候选人待定", "确认任命时间"],
    ["K003", LAST,                    NOTICE, "03-按计划进行", "按计划",   None,        "跟进到位情况"],
    ["K005", LAST,                    NOTICE, "03-按计划进行", "面谈通过", None,        "走审批"],
    ["K006", datetime(2026, 8, 20),   NOTICE, "05-已关闭",    "已任命",   None,        None],
    ["K007", LAST,                    PLAN,   "01-新增",      None,      None,        None],
]


def build_data() -> None:
    wb = Workbook()
    style_data_sheet(new_sheet(wb, "本期全量", [CURRENT_HEADER, *CURRENT_ROWS], first=True))
    wb.save(ROOT / "本期全量.xlsx")

    wb = Workbook()
    style_data_sheet(new_sheet(wb, "上期全量", [PREVIOUS_HEADER, *PREVIOUS_ROWS], first=True))
    wb.save(ROOT / "上期全量.xlsx")


# ---------------------------------------------------------------------------
# 规则工作簿
# ---------------------------------------------------------------------------
def build_rules() -> None:
    wb = Workbook()

    # ---- sheet 1: 设置 (上方设置项; 空一行; 下方输入表) ---------------------------
    new_sheet(wb, "设置", [
        ("基准日期", BASE, "本批例会日: 决策表里 =基准日期 就是它, 也是内嵌测试用的日期。每批跑的时候用 --var 基准日期=... 给真实的例会日"),
        ("输出", "output/本期全量_更新后.xlsx", "默认输出路径, 相对本规则文件所在目录, 命令行 -o 可覆盖"),
        (None, None, None),
        ("名称", "匹配", "重复键", "文件", "说明"),
        ("本期", "主表", None, "本期全量.xlsx", "被更新的表: 输出的行就是它的行。规则里引用它的列写 本期.本批冒泡"),
        ("上期", "岗位编码", "取第一条", "上期全量.xlsx", "上一批次的全量表, 按 岗位编码 对齐到主表; 上期.存在 = 上批有没有这个岗位"),
    ], first=True)

    # ---- sheet 2: 字段组 ---------------------------------------------------
    new_sheet(wb, "字段组", [
        ("名称", "列名", "说明"),
        ("会议记录", "例会时间, 类别-会议, 状态-会议, 例会结论, 备注, 待跟进事宜",
         "给这六列起个名字。决策表里以 会议记录 为列名的输出格: 写 上期 = 组内每列从上期的同名列复制; 写 清空 = 整组清空; - = 不改。"
         "同一条规则显式写了的列以显式为准 (R1 的 例会时间, R3 的 状态-会议)"),
    ])

    # ---- sheet 3: 决策表 (第 1 行元信息, 第 2 行列类型, 第 3 行列名, 第 4 行起每行一条规则) ----
    kinds = ["规则ID", "优先级", "输入", "输入", "输入", "输出", "输出", "输出", "说明"]
    names = ["规则ID", "优先级", "本期.本批冒泡", "上期.存在", "上期.状态-会议", "会议记录", "例会时间", "状态-会议", "说明"]
    rules = [
        ["R1", 1, "是", "是", "-",            "上期", "=基准日期", "-",            "本批冒泡且上批有记录: 会议记录整组从上期带入, 例会时间改成本批例会日 (显式写的列优先于字段组里的同名列)"],
        ["R2", 2, "是", "否", "-",            "-",    "=基准日期", "01-新增",      "首次冒泡: 记例会时间, 状态 01-新增"],
        ["R3", 3, "否", "是", "!= 05-已关闭", "上期", "-",        "06-本批未冒泡", "上批冒泡、本批没冒泡且上批还没关闭: 历史记录整组带入 (例会时间也是上批的), 状态标成 06。已关闭的不动"],
    ]
    new_sheet(wb, "决策表", [
        ["命中策略", "首次命中", "三条规则互斥, 每行最多命中一条; 不冒泡且上批没有记录的岗位什么都不动。本批冒泡 是文本列, 格里的 是 / 否 是文字"],
        kinds, names, *rules,
    ])

    # ---- sheet 4: 测试 --------------------------------------------------------
    kinds = ["用例", *["输入"] * 9, *["期望"] * 6, "说明"]
    names = ["用例", "本期.本批冒泡", "上期.存在", "上期.例会时间", "上期.类别-会议", "上期.状态-会议", "上期.例会结论", "上期.备注", "上期.待跟进事宜", "本期.备注",
             "例会时间", "类别-会议", "状态-会议", "例会结论", "备注", "待跟进事宜", "说明"]
    tests = [
        ["T1", "是", "是", LAST, NOTICE, "02-进展更新",  "同意推进", "候选人待定", "确认任命时间", None,  BASE, NOTICE, "02-进展更新",  "同意推进", "候选人待定", "确认任命时间", "冒泡且上批有: 六列整组带入, 例会时间是本批的 (基准日期) 而不是上批的"],
        ["T2", "是", "否", None, None,   None,          None,      None,        None,          None,  BASE, None,   "01-新增",      None,      None,        None,          "冒泡且上批没有: 首次冒泡"],
        ["T3", "否", "是", LAST, NOTICE, "03-按计划进行", "按计划",   None,        "跟进到位情况",  None,  LAST, NOTICE, "06-本批未冒泡", "按计划",   None,        "跟进到位情况",  "不冒泡且上批有未关闭记录: 整组带入, 例会时间是上批的, 状态标 06"],
        ["T4", "否", "否", None, None,   None,          None,      None,        None,          None,  None, None,   None,          None,      None,        None,          "不冒泡且上批没有: 不动"],
        ["T5", "是", "是", LAST, NOTICE, "03-按计划进行", "面谈通过", None,        "走审批",       "临时备注", BASE, NOTICE, "03-按计划进行", "面谈通过", "空",        "走审批",       "本期备注原来有值、上期为空: 整组带入会把它清成空 (期望格写 空)"],
        ["T6", "否", "是", None, NOTICE, "05-已关闭",    "已任命",   None,        None,          None,  None, None,   None,          None,      None,        None,          "不冒泡且上批已关闭: R3 的 != 05-已关闭 不成立, 不动"],
    ]
    new_sheet(wb, "测试", [kinds, names, *tests])

    style_rules_workbook(wb)
    wb.save(ROOT / "规则_决策表.xlsx")


if __name__ == "__main__":
    build_data()
    build_rules()
    print(f"已生成到 {ROOT}: 本期全量.xlsx  上期全量.xlsx  规则_决策表.xlsx")
