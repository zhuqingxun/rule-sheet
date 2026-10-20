# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""生成「岗位上会状态」样例: 两份数据 + Excel 决策表, 全部写到本脚本所在目录。

这个样例展示 customer-sync 里没有的几件事:
    - 字段名含 `-` (类别-会议 / 状态-会议), 直接来自业务方表头
    - 输出格 `清空`: 命中后把目标列写成空值 (区别于 `-` 的不改)
    - 命中策略 `首次命中`: 每行只听优先级数字最小的那条, 命中即停, 后面的规则不再看这一行;
      S2 没有任何输出, 只是把「酝酿中」的岗位拦下来不让兜底规则改它
    - 决策表里的跨变量「或」: S3 占两行, 规则 ID 相同 (合并单元格), 每行各格 AND、行间 OR
    - 变量定义里用 且 / 或 组合条件, `!=` 两侧有空值就不成立 (与单元格一致), 要把「上期没填、本期填了」算作变化得显式写
    - 判空统一写 `= 空` / `!= 空` (单元格) 与 `x == 空` / `x != 空` (表达式); `为空` / `非空` 是同义写法
    - 「测试」sheet 的表头直接从第 1 行开始 (表头位置按内容识别, 不固定行号)

运行 (任选其一):
    uv run --script examples/status-machine/make_samples.py
    uv run python examples/status-machine/make_samples.py

产物:
    本期岗位.xlsx      被更新的左表
    上期岗位.xlsx      参照用的右表
    规则_决策表.xlsx   四张 sheet: 匹配设置 / 变量定义 / 决策表 / 测试
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent

FILL_BY_KIND = {
    "输入": "DDEBF7",   # 浅蓝: 业务人员改的条件格
    "输出": "E2EFDA",   # 浅绿: 命中后写回的值
    "期望": "E2EFDA",
    "说明": "FFF2CC",   # 浅黄: 纯注释
    "规则ID": "EDEDED",
    "优先级": "EDEDED",
    "用例": "EDEDED",
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)


def put(ws: Worksheet, row: int, col: int, value: Any, *, bold: bool = False, fill: str | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # 以 = 开头的 rule-sheet 表达式按文本写, 否则 openpyxl / Excel 会当成公式
        cell.number_format = "@"
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center")
    if bold:
        cell.font = BOLD
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def set_widths(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_data_sheet(ws: Worksheet, header: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    for c, name in enumerate(header, start=1):
        put(ws, 1, c, name, bold=True, fill="D9D9D9")
    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            put(ws, r, c, val)
    ws.freeze_panes = "A2"
    set_widths(ws, widths)


# ---------------------------------------------------------------------------
# 样例数据: 每一行对应一种情形, 便于对照决策表看命中结果
# ---------------------------------------------------------------------------
PLAN = "近2个月酝酿及上会计划"
NOTICE = "变动通报"

CURRENT_HEADER = ["岗位编码", "岗位名称", "类别-会议", "状态-会议", "拟任人", "流程节点"]
CURRENT_ROWS = [
    # 编号    名称            类别-会议   状态-会议       拟任人  流程节点     上期情况                      预期
    ["P001", "财务部-部长",   None,      "01-新增",     None,   None],     # 上期 变动通报                  -> S1 清空状态
    ["P002", "人事部-副部长", PLAN,      None,          None,   None],     # 上期无                        -> S2 拦截不动 (优先级 策略下会被 S5 写成 01-新增)
    ["P003", "研发中心-主任", NOTICE,    "01-新增",     "王五", "部门评审"], # 上期 变动通报, 拟任人 张三      -> S3 拟任人变了, 02-进展更新
    ["P004", "市场部-经理",   NOTICE,    None,          None,   None],     # 上期无                        -> S5 01-新增
    ["P005", "法务部-主管",   None,      None,          None,   None],     # 上期无                        -> S1 清空空值, 无变更
    ["P006", "采购部-经理",   NOTICE,    "02-进展更新",  None,   None],     # 上期 酝酿 (类别不同)           -> S5 01-新增 (重新计为新增)
    ["P007", "运营部-总监",   NOTICE,    "01-新增",     "李四", "公司审批"], # 上期 变动通报, 节点 部门评审    -> S3 流程节点变了 (第二个分支), 02-进展更新
    ["P008", "行政部-经理",   NOTICE,    "02-进展更新",  "周七", "公司审批"], # 上期 变动通报, 都没变          -> S4 03-按计划进行
    ["P009", "质量部-部长",   NOTICE,    "01-新增",     "赵六", None],     # 上期 变动通报, 拟任人为空       -> S3 新任命算变化, 02-进展更新
]

PREVIOUS_HEADER = ["岗位编码", "类别-会议", "状态-会议", "拟任人", "流程节点"]
PREVIOUS_ROWS = [
    ["P001", NOTICE, "01-新增",     None,   None],
    ["P003", NOTICE, "01-新增",     "张三", "部门评审"],
    ["P006", PLAN,   None,          None,   None],
    ["P007", NOTICE, "01-新增",     "李四", "部门评审"],
    ["P008", NOTICE, "02-进展更新",  "周七", "公司审批"],
    ["P009", NOTICE, "01-新增",     None,   None],
]


def build_data() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "本期岗位"
    write_data_sheet(ws, CURRENT_HEADER, CURRENT_ROWS, widths=[10, 16, 24, 14, 10, 12])
    wb.save(ROOT / "本期岗位.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "上期岗位"
    write_data_sheet(ws, PREVIOUS_HEADER, PREVIOUS_ROWS, widths=[10, 24, 14, 10, 12])
    wb.save(ROOT / "上期岗位.xlsx")


# ---------------------------------------------------------------------------
# 规则工作簿
# ---------------------------------------------------------------------------
def write_kv_sheet(ws: Worksheet, title: tuple[str, str, str], rows: list[tuple[Any, Any, Any]],
                   widths: list[int]) -> None:
    for c, name in enumerate(title, start=1):
        put(ws, 1, c, name, bold=True, fill="D9D9D9")
    for r, (k, v, note) in enumerate(rows, start=2):
        put(ws, r, 1, k, bold=True)
        put(ws, r, 2, v)
        put(ws, r, 3, note, fill=FILL_BY_KIND["说明"])
    set_widths(ws, widths)


def write_rule_table(ws: Worksheet, top_note: list[Any] | None, kinds: list[str], names: list[str],
                     rows: list[list[Any]], widths: list[int]) -> int:
    """表头两行 (列类型 / 列名), 之前可以有一行元信息 (命中策略 / 说明); 表头之后每行一条规则。返回正文起始行号。"""
    start = 1
    if top_note:
        for c, val in enumerate(top_note, start=1):
            put(ws, 1, c, val, bold=(c == 1), fill=FILL_BY_KIND["说明"] if c == 3 else None)
        start = 2
    for c, (kind, name) in enumerate(zip(kinds, names, strict=True), start=1):
        fill = FILL_BY_KIND.get(kind)
        put(ws, start, c, kind, bold=True, fill=fill)
        put(ws, start + 1, c, name, bold=True, fill=fill)
    body = start + 2
    for r, row in enumerate(rows, start=body):
        for c, val in enumerate(row, start=1):
            kind = kinds[c - 1]
            fill = FILL_BY_KIND.get(kind) if kind in ("输入", "输出", "期望") else None
            put(ws, r, c, val, fill=fill)
            horizontal = "center" if kind in ("输入", "规则ID", "优先级", "用例") else "left"
            ws.cell(row=r, column=c).alignment = Alignment(horizontal=horizontal, vertical="center")
    ws.freeze_panes = f"C{body}"
    set_widths(ws, widths)
    return body


def build_rules() -> None:
    wb = Workbook()

    # ---- sheet 1: 匹配设置 --------------------------------------------------
    ws = wb.active
    ws.title = "匹配设置"
    write_kv_sheet(ws, ("设置项", "值", "说明"), [
        ("左表", "本期岗位.xlsx", "被更新的表, 相对路径以本规则文件所在目录为基准"),
        ("右表", "上期岗位.xlsx", "参照表, 只读"),
        ("左表别名", "本期", "规则里引用左表字段的前缀, 如 本期.类别-会议"),
        ("右表别名", "上期", "规则里引用右表字段的前缀, 如 上期.类别-会议"),
        ("匹配键", "岗位编码", "两表按此列对行"),
        ("输出", "output/本期岗位_更新后.xlsx", "结果文件, 原表不动"),
    ], widths=[14, 32, 52])

    # ---- sheet 2: 变量定义 --------------------------------------------------
    ws = wb.create_sheet("变量定义")
    write_kv_sheet(ws, ("变量", "定义", "说明"), [
        ("拟任人有变化", "本期.拟任人 != 上期.拟任人 或 (上期.拟任人 == 空 且 本期.拟任人 != 空)",
         "!= 两侧有空值就不成立 (和单元格一样), 所以「上期没填、本期填了」要单独写一支; == 空 / != 空 就是判空"),
        ("流程节点有变化", "本期.流程节点 != 上期.流程节点 或 (上期.流程节点 == 空 且 本期.流程节点 != 空)",
         "同上。两期都空不算变化"),
        (None, None, None),
        ("连接词", None, "且 / 或 / 非 (也认 and / or / not), 两侧留空格; 可用括号分组"),
        ("空值", None, "x == 空 / x != 空 判空 (也认 为空(x) / 非空(x)); 空 只能和 == / != 搭配, 不能参与大小比较和运算"),
        ("可用函数", None, "abs  round  天数差(晚, 早)  非空(x)  为空(x)  属于(x, 值1, 值2, ...)"),
    ], widths=[16, 70, 60])

    # ---- sheet 3: 决策表 ----------------------------------------------------
    ws = wb.create_sheet("决策表")
    kinds = ["规则ID", "优先级", "输入", "输入", "输入", "输入", "输出", "说明"]
    names = ["规则ID", "优先级", "本期.类别-会议", "上期.类别-会议", "拟任人有变化", "流程节点有变化", "状态-会议", "说明"]
    rules = [
        ["S1", 1,    "= 空", "-",              "-", "-", "清空",        "本期不再上会 (= 空 即判空, 也可写 为空), 清掉状态 (清空 = 写成空值, 也可写 空; - 才是不改)"],
        ["S2", 2,    PLAN,   "-",              "-", "-", "-",           "酝酿中的岗位由人工维护, 不处理。首次命中 下命中即停, 后面的规则不再看这一行 (优先级 策略下会被 S5 改成 01-新增)"],
        ["S3", 3,    "-",    "= 本期.类别-会议", "是", "-", "02-进展更新", "类别与上期相同, 且拟任人或流程节点有变化。两行同一个规则ID = 同一条规则的两个「或」分支: 每行各格 AND, 行间 OR"],
        ["S3", None, "-",    "= 本期.类别-会议", "-", "是", None,         "S3 的第二个分支。续行的优先级和输出留空 (或与首行相同), 规则ID 可以合并单元格"],
        ["S4", 4,    "-",    "= 本期.类别-会议", "-", "-", "03-按计划进行", "类别与上期相同, 没有变化"],
        ["S5", 5,    "-",    "-",              "-", "-", "01-新增",     "兜底: 前面都没命中的行 (有类别、不是酝酿、与上期不同) 都是新增"],
    ]
    body = write_rule_table(
        ws,
        top_note=["命中策略", "首次命中", "每行只听优先级数字最小的那条命中规则, 命中即停; 同优先级按行序。输出格 清空 = 写成空值, - = 不改"],
        kinds=kinds, names=names, rows=rules,
        widths=[8, 8, 22, 18, 14, 16, 14, 70],
    )
    s3 = body + 2  # S3 的两行: 合并 规则ID 与 优先级, 视觉上就是一条规则
    ws.merge_cells(start_row=s3, start_column=1, end_row=s3 + 1, end_column=1)
    ws.merge_cells(start_row=s3, start_column=2, end_row=s3 + 1, end_column=2)

    # ---- sheet 4: 测试 (表头直接从第 1 行开始, 表头位置按内容识别) --------------
    ws = wb.create_sheet("测试")
    kinds = ["用例", "输入", "输入", "输入", "输入", "输入", "输入", "输入", "期望", "说明"]
    names = ["用例", "本期.类别-会议", "本期.状态-会议", "上期.类别-会议",
             "本期.拟任人", "上期.拟任人", "本期.流程节点", "上期.流程节点", "状态-会议", "说明"]
    tests = [
        ["T1",  None,   "01-新增",    NOTICE, None,   None,   None,      None,      "清空",         "类别为空, 清掉状态。期望格写 清空 = 期望该列为空 (留空是期望不变)"],
        ["T2",  PLAN,   None,         None,   None,   None,   None,      None,      None,           "S2 拦截, 状态保持空; 优先级 策略下会被 S5 改成 01-新增, 这条用例会失败"],
        ["T3",  NOTICE, "01-新增",    NOTICE, "王五", "张三", "部门评审", "部门评审", "02-进展更新",   "类别相同, 拟任人变了 -> 进展更新"],
        ["T4",  NOTICE, None,         None,   None,   None,   None,      None,      "01-新增",      "上期没有这个岗位 -> 兜底新增"],
        ["T5",  None,   None,         None,   None,   None,   None,      None,      None,           "清空本来就空的格, 无变化"],
        ["T6",  NOTICE, "02-进展更新", PLAN,   None,   None,   None,      None,      "01-新增",      "类别变了 -> 重新计为新增"],
        ["T7",  NOTICE, "01-新增",    NOTICE, "李四", "李四", "公司审批", "部门评审", "02-进展更新",   "类别相同, 流程节点变了 -> S3 的第二个分支"],
        ["T8",  NOTICE, "02-进展更新", NOTICE, "周七", "周七", "公司审批", "公司审批", "03-按计划进行", "类别相同, 都没变 -> 按计划进行"],
        ["T9",  NOTICE, "01-新增",    NOTICE, "赵六", None,   None,      None,      "02-进展更新",   "上期拟任人为空、本期有 = 新任命, 算变化 (变量定义里显式写的那一支)"],
        ["T10", NOTICE, "01-新增",    NOTICE, None,   None,   "部门评审", "部门评审", "03-按计划进行", "两期拟任人都空不算变化: != 两侧有空值就不成立, 决策表与 YAML 一致"],
    ]
    write_rule_table(
        ws,
        top_note=None,
        kinds=kinds, names=names, rows=tests,
        widths=[6, 22, 14, 22, 10, 10, 12, 12, 14, 64],
    )

    wb.save(ROOT / "规则_决策表.xlsx")


if __name__ == "__main__":
    build_data()
    build_rules()
    print(f"已生成到 {ROOT}: 本期岗位.xlsx  上期岗位.xlsx  规则_决策表.xlsx")
