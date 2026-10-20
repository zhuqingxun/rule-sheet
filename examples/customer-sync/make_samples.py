# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""生成样例数据（两份业务 Excel）和 Excel 版决策表, 全部写到本脚本所在目录。

运行 (任选其一):
    uv run --script examples/customer-sync/make_samples.py
    uv run python examples/customer-sync/make_samples.py

产物:
    客户主表.xlsx      被更新的主表 (规则里叫 客户)
    本月流水.xlsx      参照表 (规则里叫 本月)
    规则_决策表.xlsx   五张 sheet: 输入 / 设置 / 变量定义 / 决策表 / 测试
"""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent

FILL_BY_KIND = {
    "输入": "DDEBF7",   # 浅蓝: 业务人员改的条件格
    "输出": "E2EFDA",   # 浅绿: 命中后写回的值
    "期望": "E2EFDA",
    "说明": "FFF2CC",   # 浅黄: 纯注释
    "规则ID": "EDEDED",
    "优先级": "EDEDED",
    "用例": "EDEDED",
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)


def put(ws: Worksheet, row: int, col: int, value: Any, *, bold: bool = False,
        fill: str | None = None, fmt: str | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center")
    if bold:
        cell.font = BOLD
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if fmt and value is not None:
        cell.number_format = fmt


def set_widths(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_data_sheet(ws: Worksheet, header: list[str], rows: list[list[Any]],
                     text_cols: set[int], date_cols: set[int], widths: list[int]) -> None:
    for c, name in enumerate(header, start=1):
        put(ws, 1, c, name, bold=True, fill="D9D9D9")
    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            fmt = "@" if c in text_cols else ("yyyy-mm-dd" if c in date_cols else None)
            put(ws, r, c, val, fmt=fmt)
    ws.freeze_panes = "A2"
    set_widths(ws, widths)


# ---------------------------------------------------------------------------
# 样例数据: 每一行对应一种情形, 便于对照决策表看命中结果
# ---------------------------------------------------------------------------
MASTER_HEADER = ["客户编号", "客户名称", "等级", "状态", "金额", "电话", "上次活跃"]
MASTER_ROWS = [
    # 编号    名称        等级    状态    金额    电话            上次活跃
    ["C001", "华信贸易", "普通", "活跃", 10000, "13800000001", date(2026, 10, 8)],   # 本月 10800, 偏差 8%  -> R1
    ["C002", "远大机械", "普通", "活跃", 10000, "13800000002", date(2026, 10, 2)],   # 本月 10300, 偏差 3%  -> 无
    ["C003", "恒通物流", "普通", "活跃", 8000,  "13800000003", date(2026, 3, 1)],    # 本月缺失, 232 天      -> R2
    ["C004", "星辰科技", "普通", "活跃", 5000,  "13800000004", date(2026, 9, 15)],   # 本月缺失, 34 天       -> 无
    ["C005", "蓝海餐饮", "普通", "试用", 2000,  "13800000005", date(2026, 10, 10)],  # 偏差 30% 但状态试用, 电话变 -> R3
    ["C006", "鼎盛集团", "VIP",  "活跃", 12000, "13800000006", date(2026, 10, 11)],  # 偏差 8.3% 且电话变     -> R1 + R3
    ["C007", "旧友商行", "普通", "休眠", 0,     "13800000007", date(2025, 12, 1)],   # 已休眠, 本月缺失       -> 无 (R2 的 != 休眠 排除)
    ["C008", "金桥投资", "金卡", "活跃", 9000,  "13800000008", date(2026, 10, 1)],   # 偏差 4%, 金卡, 本月电话空 -> R4 + R5
]

MONTHLY_HEADER = ["客户编号", "金额", "电话", "交易日期"]
MONTHLY_ROWS = [
    ["C001", 10800, "13800000001", date(2026, 10, 15)],
    ["C002", 10300, "13800000002", date(2026, 10, 12)],
    ["C005", 2600,  "13900000005", date(2026, 10, 9)],
    ["C006", 13000, "13900000006", date(2026, 10, 16)],
    ["C008", 9360,  None,          date(2026, 10, 14)],
]


def build_data() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "客户主表"
    write_data_sheet(ws, MASTER_HEADER, MASTER_ROWS, text_cols={6}, date_cols={7},
                     widths=[10, 12, 8, 8, 10, 16, 12])
    wb.save(ROOT / "客户主表.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "本月流水"
    write_data_sheet(ws, MONTHLY_HEADER, MONTHLY_ROWS, text_cols={3}, date_cols={4},
                     widths=[10, 10, 16, 12])
    wb.save(ROOT / "本月流水.xlsx")


# ---------------------------------------------------------------------------
# 规则工作簿
# ---------------------------------------------------------------------------
def write_kv_sheet(ws: Worksheet, title: tuple[str, str, str], rows: list[tuple[Any, Any, Any]],
                   widths: list[int]) -> None:
    for c, name in enumerate(title, start=1):
        put(ws, 1, c, name, bold=True, fill="D9D9D9")
    for r, (k, v, note) in enumerate(rows, start=2):
        put(ws, r, 1, k, bold=True)
        put(ws, r, 2, v, fmt="yyyy-mm-dd" if isinstance(v, date) else None)
        put(ws, r, 3, note, fill=FILL_BY_KIND["说明"])
    set_widths(ws, widths)


def write_table_sheet(ws: Worksheet, header: tuple[str, ...], rows: list[tuple[Any, ...]], widths: list[int]) -> None:
    """第 1 行列名, 之后一行一条记录 (「输入」sheet 用)。"""
    for c, name in enumerate(header, start=1):
        put(ws, 1, c, name, bold=True, fill="D9D9D9")
    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            put(ws, r, c, val, bold=(c == 1), fill=FILL_BY_KIND["说明"] if header[c - 1] == "说明" else None)
    set_widths(ws, widths)


def write_rule_table(ws: Worksheet, top_note: list[Any], kinds: list[str], names: list[str],
                     rows: list[list[Any]], widths: list[int]) -> None:
    """第 1 行元信息 (命中策略 / 说明), 第 2 行列类型 (输入/输出), 第 3 行变量名, 第 4 行起每行一条规则。
    表头按内容识别 (写着 规则ID / 用例 的那一行), 元信息行可以没有。"""
    for c, val in enumerate(top_note, start=1):
        put(ws, 1, c, val, bold=(c == 1), fill=FILL_BY_KIND["说明"] if c == 3 else None)
    for c, (kind, name) in enumerate(zip(kinds, names, strict=True), start=1):
        fill = FILL_BY_KIND.get(kind)
        put(ws, 2, c, kind, bold=True, fill=fill)
        put(ws, 3, c, name, bold=True, fill=fill)
    for r, row in enumerate(rows, start=4):
        for c, val in enumerate(row, start=1):
            kind = kinds[c - 1]
            fill = FILL_BY_KIND.get(kind) if kind in ("输入", "输出", "期望") else None
            put(ws, r, c, val, fill=fill)
            horizontal = "center" if kind in ("输入", "规则ID", "优先级", "用例") else "left"
            ws.cell(row=r, column=c).alignment = Alignment(horizontal=horizontal, vertical="center")
    ws.freeze_panes = "C4"
    set_widths(ws, widths)


def build_rules() -> None:
    wb = Workbook()

    # ---- sheet 1: 输入 (形参: 名称 + 怎么对齐到主表; 文件只是默认绑定, 命令行 -i 可覆盖) ----
    ws = wb.active
    ws.title = "输入"
    write_table_sheet(ws, ("名称", "匹配", "重复键", "文件", "说明"), [
        ("客户", "主表", None, "客户主表.xlsx", "被更新的表: 输出的行就是它的行。规则里引用它的列写 客户.状态; 相对路径以本规则文件所在目录为基准"),
        ("本月", "客户编号", "取第一条", "本月流水.xlsx", "参照表, 按 客户编号 对齐到主表 (多个键用逗号分隔); 引用写 本月.金额, 本月.存在 = 是否匹配到"),
    ], widths=[10, 12, 10, 20, 80])

    # ---- sheet 2: 设置 (可选) ------------------------------------------------
    ws = wb.create_sheet("设置")
    write_kv_sheet(ws, ("设置项", "值", "说明"), [
        ("基准日期", date(2026, 10, 19), "天数差计算用的「今天」, 固定下来保证每次运行结果一致; 不写取运行当天"),
        ("输出", "output/客户主表_更新后.xlsx", "默认输出路径, 命令行 -o 可覆盖; 不写则是 output/<主表文件名>_更新后.xlsx"),
    ], widths=[14, 32, 60])

    # ---- sheet 3: 变量定义 --------------------------------------------------
    ws = wb.create_sheet("变量定义")
    write_kv_sheet(ws, ("变量", "定义", "说明"), [
        ("偏差率", "abs(本月.金额 - 客户.金额) / 客户.金额", "本月金额相对主表金额的偏差比例, 本月缺失时为空"),
        ("距上次活跃天数", "天数差(基准日期, 客户.上次活跃)", "基准日期见「设置」"),
        (None, None, None),
        ("内置变量", None, "客户.<列名>  本月.<列名>  本月.存在(参照表是否匹配到)  基准日期"),
        ("连接词", None, "且 / 或 / 非 (也认 and / or / not), 两侧留空格; 可用括号分组"),
        ("空值", None, "x == 空 / x != 空 判空 (也认 为空(x) / 非空(x)); 空 只能和 == / != 搭配"),
        ("可用函数", None, "abs  round  天数差(晚, 早)  非空(x)  为空(x)  属于(x, 值1, 值2, ...)"),
    ], widths=[16, 44, 52])

    # ---- sheet 4: 决策表 ----------------------------------------------------
    ws = wb.create_sheet("决策表")
    kinds = ["规则ID", "优先级", "输入", "输入", "输入", "输入", "输入", "输入", "输入",
             "输出", "输出", "输出", "输出", "输出", "说明"]
    names = ["规则ID", "优先级", "客户.状态", "客户.等级", "偏差率", "本月.存在", "距上次活跃天数", "本月.电话", "客户.电话",
             "状态", "电话", "复核标记", "备注", "提醒", "说明"]
    rules = [
        ["R1",  1, "活跃",    "-",         "> 0.05",       "是", "-",     "-",            "-",    "-",    "-",        "待复核", "金额偏差超 5%",      "-",         "活跃客户金额偏差超 5% 需复核"],
        ["R2",  2, "!= 休眠", "-",         "-",            "否", "> 180", "-",            "-",    "休眠", "-",        "-",      "连续 6 个月无流水",  "-",         "本月无流水且长期不活跃, 转休眠"],
        ["R3",  3, "-",       "-",         "-",            "是", "-",     "!= 客户.电话", "-",    "-",    "本月.电话", "-",      "电话变更",          "-",         "本月电话与主表不同则覆盖 (本月电话为空不算不同)"],
        ["R4",  4, "活跃",    "VIP, 金卡", "[0.02..0.05]", "是", "-",     "-",            "-",    "-",    "-",        "-",      "VIP 小幅偏差, 关注", "-",         "高等级客户小幅偏差也要留意"],
        ["R5",  5, "-",       "-",         "-",            "是", "-",     "为空",         "-",    "-",    "-",        "-",      "-",                  "电话有缺失", "本月电话为空, 提醒补录。下一行同 ID = 同一条规则的另一个「或」分支 (YAML 里 R5 的 任一满足)"],
        ["R5",  None, "-",    "-",         "-",            "是", "-",     "-",            "为空", "-",    "-",        "-",      "-",                  None,        "主表电话为空, 提醒补录。续行的优先级和输出留空 (或与首行相同)"],
    ]
    write_rule_table(
        ws,
        top_note=["命中策略", "优先级", "多条规则同时命中时, 每个输出列取优先级数字最小的那条; 输出格填 - 表示不改; 同一行各条件之间是 AND, 相邻同 ID 行是同一规则的 OR 分支"],
        kinds=kinds, names=names, rows=rules,
        widths=[8, 8, 10, 11, 13, 10, 15, 14, 12, 8, 12, 10, 20, 12, 44],
    )

    # ---- sheet 5: 测试 ------------------------------------------------------
    ws = wb.create_sheet("测试")
    kinds = ["用例", "输入", "输入", "输入", "输入", "输入", "输入", "输入",
             "期望", "期望", "期望", "期望", "期望", "说明"]
    names = ["用例", "客户.状态", "客户.等级", "偏差率", "本月.存在", "距上次活跃天数", "客户.电话", "本月.电话",
             "状态", "电话", "复核标记", "备注", "提醒", "说明"]
    tests = [
        ["T1", "活跃", "普通", 0.08, "是", 5,   "13800000001", "13800000001", None,   None,          "待复核", "金额偏差超 5%",      None,         "偏差 8% 触发 R1"],
        ["T2", "活跃", "普通", 0.03, "是", 5,   "13800000002", "13800000002", None,   None,          None,     None,                 None,         "偏差 3% 普通客户, 什么都不发生"],
        ["T3", "活跃", "普通", None, "否", 200, "13800000003", None,          "休眠", None,          None,     "连续 6 个月无流水",  None,         "本月缺失且 200 天, 触发 R2"],
        ["T4", "休眠", "普通", None, "否", 400, "13800000007", None,          None,   None,          None,     None,                 None,         "已休眠不重复命中 R2"],
        ["T5", "试用", "普通", 0.30, "是", 3,   "13800000005", "13900000005", None,   "13900000005", None,     "电话变更",          None,         "试用客户不触发 R1, 只换电话"],
        ["T6", "活跃", "金卡", 0.04, "是", 10,  "13800000008", None,          None,   None,          None,     "VIP 小幅偏差, 关注", "电话有缺失", "金卡 4% 落在 R4 区间; 本月电话空不触发 R3 但触发 R5"],
        ["T7", "活跃", "VIP",  0.08, "是", 2,   "13800000006", "13900000006", None,   "13900000006", "待复核", "金额偏差超 5%",      None,         "R1 与 R3 同时命中, 备注听优先级高的 R1"],
        ["T8", "活跃", "普通", 0.01, "是", 3,   None,          "13800000009", None,   None,          None,     None,                 "电话有缺失", "主表电话空, 触发 R5 的第二个分支"],
    ]
    write_rule_table(
        ws,
        top_note=["说明", None, "每行一个用例: 填输入变量的值, 期望列留空表示「该列不变」。运行时先跑本表, 全部通过才处理真实数据"],
        kinds=kinds, names=names, rows=tests,
        widths=[6, 10, 11, 8, 10, 15, 14, 14, 8, 14, 10, 20, 12, 40],
    )

    wb.save(ROOT / "规则_决策表.xlsx")


if __name__ == "__main__":
    build_data()
    build_rules()
    print(f"已生成到 {ROOT}: 客户主表.xlsx  本月流水.xlsx  规则_决策表.xlsx")
