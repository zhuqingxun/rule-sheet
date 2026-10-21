# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""生成「岗位打标签」样例: 两份数据 + Excel 决策表, 全部写到本脚本所在目录。

这个样例展示 0.6 加的几件事 (issue #9 #10 #11 #14 #15 #16):
    - 命中策略 `多条拼接`: 一个岗位可以同时命中多条规则, 「标签」列把它们按优先级顺序用 / 接起来 (多标签分类)
    - 单元格里的文本关键字: `包含 CEO, 总裁` / `不包含 助理` / `开头是 2-, 3-` (按字面匹配, 任一命中即真, 空值不成立)
    - 清单文件当参照表: 关键岗位清单.xlsx 按 岗位编码 对齐到主表, `关键岗位清单.存在` 就是「在不在清单里」
    - 变量定义里的 `数字(岗位.职级)` (22级 / 23 / P23 -> 22 / 23 / 23) 和时长字面量 `基准日期 - 365天`
    - 输出格里的 `如果(条件, 真值, 假值)` 与 `拼接(...)`: 一层组织为空就只用岗位名称当显示名

运行 (任选其一):
    uv run --script examples/position-tags/make_samples.py
    uv run python examples/position-tags/make_samples.py

产物:
    岗位清单.xlsx        被更新的主表 (规则里叫 岗位)
    关键岗位清单.xlsx    参照表 / 清单 (规则里叫 关键岗位清单)
    规则_决策表.xlsx     五张 sheet: 输入 / 设置 / 变量定义 / 决策表 / 测试
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent

FILL_BY_KIND = {
    "输入": "DDEBF7",   # 浅蓝: 业务人员改的条件格
    "输出": "E2EFDA",   # 浅绿: 命中后写回的值
    "期望": "E2EFDA",
    "说明": "FFF2CC",   # 浅黄: 纯注释
    "规则ID": "EDEDED",
    "优先级": "EDEDED",
    "用例": "EDEDED",
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)


def put(ws: Worksheet, row: int, col: int, value: Any, *, bold: bool = False, fill: str | None = None) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # 以 = 开头的 rule-sheet 表达式按文本写, 否则 openpyxl / Excel 会当成公式
        cell.number_format = "@"
    elif isinstance(value, datetime):
        cell.number_format = "yyyy-mm-dd"
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center")
    if bold:
        cell.font = BOLD
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def set_widths(ws: Worksheet, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_data_sheet(ws: Worksheet, header: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    for c, name in enumerate(header, start=1):
        put(ws, 1, c, name, bold=True, fill="D9D9D9")
    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            put(ws, r, c, val)
    ws.freeze_panes = "A2"
    set_widths(ws, widths)


# ---------------------------------------------------------------------------
# 样例数据: 每一行对应一种情形, 便于对照决策表看命中结果 (基准日期 2026-10-01)
# ---------------------------------------------------------------------------
D = datetime

POSITION_HEADER = ["岗位编码", "岗位名称", "一层组织", "职级", "推荐时间", "标签", "显示名"]
POSITION_ROWS = [
    # 编码    名称                一层组织        职级     推荐时间          标签    显示名   预期
    ["J001", "集团CEO",          "1-集团总部",   "22级",  D(2026, 6, 1),   None,  None],  # G1 高管 (包含 CEO) + G5 新推荐; 1- 不是直管
    ["J002", "华东代表处-行政助理", "3-华东代表处", "20级",  D(2025, 1, 1),   "普通", None],  # 3- 开头但名称含 助理, G2 不命中; 什么都没命中, 标签保持 普通
    ["J003", "华东代表处-代表",    "3-华东代表处", "23级",  None,            "待定", None],  # G2 直管 + G3 在清单里 + G4 职级 23; 原标签 待定 被整体替换
    ["J004", "终端BG-产品总裁",    "2-终端BG",     24,      D(2026, 9, 15),  None,  None],  # G1 + G2 + G4 (职级格是数字 24) + G5, 四个标签
    ["J005", "法务经理",          None,          "19级",  D(2026, 2, 1),   None,  None],  # G5; 一层组织为空, 显示名走 如果() 的假分支
    ["J006", "研发-SPDT主管",     "4-研发体系",   "P23",   D(2024, 12, 31), None,  None],  # G3 在清单里 + G4 (P23 -> 23); 推荐超过一年, 不是新推荐
    ["J007", "财务部-会计",       "5-职能平台",   None,    D(2026, 10, 1),  None,  None],  # G5; 职级为空, 数字() 取不到, >= 23 不成立
    ["J008", "供应链-采购经理",    "6-供应链",     "22级",  D(2025, 10, 1),  None,  None],  # G5: 正好 365 天, >= 成立 (边界)
]

KEY_HEADER = ["岗位编码", "说明"]
KEY_ROWS = [
    ["J003", "代表处一把手"],
    ["J006", "核心研发岗"],
    ["J999", "已撤销的岗位, 主表里没有 (清单里多出来的行不影响)"],
]


def build_data() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "岗位"
    write_data_sheet(ws, POSITION_HEADER, POSITION_ROWS, widths=[10, 22, 16, 8, 12, 26, 32])
    wb.save(ROOT / "岗位清单.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "关键岗位"
    write_data_sheet(ws, KEY_HEADER, KEY_ROWS, widths=[10, 50])
    wb.save(ROOT / "关键岗位清单.xlsx")


# ---------------------------------------------------------------------------
# 规则工作簿
# ---------------------------------------------------------------------------
def write_kv_sheet(ws: Worksheet, title: tuple[str, str, str], rows: list[tuple[Any, Any, Any]],
                   widths: list[int]) -> None:
    for c, name in enumerate(title, start=1):
        put(ws, 1, c, name, bold=True, fill="D9D9D9")
    for r, (k, v, note) in enumerate(rows, start=2):
        put(ws, r, 1, k, bold=True)
        put(ws, r, 2, v)
        put(ws, r, 3, note, fill=FILL_BY_KIND["说明"])
    set_widths(ws, widths)


def write_table_sheet(ws: Worksheet, header: tuple[str, ...], rows: list[tuple[Any, ...]], widths: list[int]) -> None:
    """第 1 行列名, 之后一行一条记录 (「输入」sheet 用)。"""
    for c, name in enumerate(header, start=1):
        put(ws, 1, c, name, bold=True, fill="D9D9D9")
    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            put(ws, r, c, val, bold=(c == 1), fill=FILL_BY_KIND["说明"] if header[c - 1] == "说明" else None)
    set_widths(ws, widths)


def write_rule_table(ws: Worksheet, top_note: list[Any] | None, kinds: list[str], names: list[str],
                     rows: list[list[Any]], widths: list[int]) -> int:
    """表头两行 (列类型 / 列名), 之前可以有一行元信息 (命中策略 / 说明); 表头之后每行一条规则。返回正文起始行号。"""
    start = 1
    if top_note:
        for c, val in enumerate(top_note, start=1):
            put(ws, 1, c, val, bold=(c == 1), fill=FILL_BY_KIND["说明"] if c == 3 else None)
        start = 2
    for c, (kind, name) in enumerate(zip(kinds, names, strict=True), start=1):
        fill = FILL_BY_KIND.get(kind)
        put(ws, start, c, kind, bold=True, fill=fill)
        put(ws, start + 1, c, name, bold=True, fill=fill)
    body = start + 2
    for r, row in enumerate(rows, start=body):
        for c, val in enumerate(row, start=1):
            kind = kinds[c - 1]
            fill = FILL_BY_KIND.get(kind) if kind in ("输入", "输出", "期望") else None
            put(ws, r, c, val, fill=fill)
            horizontal = "center" if kind in ("输入", "规则ID", "优先级", "用例") else "left"
            ws.cell(row=r, column=c).alignment = Alignment(horizontal=horizontal, vertical="center")
    ws.freeze_panes = f"C{body}"
    set_widths(ws, widths)
    return body


DISPLAY = '=如果(岗位.一层组织 != 空, 拼接(岗位.一层组织, "/", 岗位.岗位名称), 岗位.岗位名称)'


def build_rules() -> None:
    wb = Workbook()

    # ---- sheet 1: 输入 ------------------------------------------------------
    ws = wb.active
    ws.title = "输入"
    write_table_sheet(ws, ("名称", "匹配", "重复键", "文件", "说明"), [
        ("岗位", "主表", None, "岗位清单.xlsx", "被更新的表: 输出的行就是它的行。规则里引用它的列写 岗位.岗位名称"),
        ("关键岗位清单", "岗位编码", "取第一条", "关键岗位清单.xlsx",
         "清单文件就是一个参照表: 按 岗位编码 对齐到主表, 关键岗位清单.存在 = 这个岗位在不在清单里 (issue #15 的「属于文件」)"),
    ], widths=[14, 12, 10, 20, 90])

    # ---- sheet 2: 设置 ------------------------------------------------------
    ws = wb.create_sheet("设置")
    write_kv_sheet(ws, ("设置项", "值", "说明"), [
        ("基准日期", datetime(2026, 10, 1), "「近一年推荐」以它为准; 不填取运行当天"),
        ("拼接分隔符", "/", "命中策略 多条拼接 用的分隔符, 默认就是 /"),
        ("输出", "output/岗位清单_更新后.xlsx", "默认输出路径, 命令行 -o 可覆盖"),
    ], widths=[14, 32, 60])

    # ---- sheet 3: 变量定义 --------------------------------------------------
    ws = wb.create_sheet("变量定义")
    write_kv_sheet(ws, ("变量", "定义", "说明"), [
        ("职级数", "数字(岗位.职级)",
         "文本里的第一个数: 22级 / 23 / P23 -> 22 / 23 / 23; 取不到 (空格或没有数字) 为空, 空值不参与 >= 比较"),
        ("近一年推荐", "岗位.推荐时间 >= 基准日期 - 365天",
         "N天 是时长字面量, 可以加减日期; 推荐时间为空不成立。等价写法: 天数差(基准日期, 岗位.推荐时间) <= 365"),
        (None, None, None),
        ("连接词", None, "且 / 或 / 非 (也认 and / or / not), 两侧留空格; 可用括号分组"),
        ("空值", None, "x == 空 / x != 空 判空 (也认 为空(x) / 非空(x)); 空 只能和 == / != 搭配, 不能参与大小比较和运算"),
        ("可用函数", None,
         "abs  round  天数差(晚, 早)  非空(x)  为空(x)  属于(x, 值1, ...)  "
         "包含 / 不包含 / 开头是 / 开头不是 / 结尾是 / 结尾不是 (x, 文本1, ...)  拼接(a, b, ...)  数字(x)  如果(条件, 真值, 假值)"),
    ], widths=[14, 44, 100])

    # ---- sheet 4: 决策表 ----------------------------------------------------
    ws = wb.create_sheet("决策表")
    kinds = ["规则ID", "优先级", "输入", "输入", "输入", "输入", "输入", "输出", "输出", "说明"]
    names = ["规则ID", "优先级", "岗位.岗位名称", "岗位.一层组织", "关键岗位清单.存在", "职级数", "近一年推荐",
             "标签", "显示名", "说明"]
    rules = [
        ["G1", 1, "包含 CEO, 总裁", "-",            "-", "-",     "-", "高管",    "-",     "子串匹配, 任一命中即真。按字面匹配不是正则; 注意「总裁助理」也含「总裁」, 要精确就用 开头是 / 等于"],
        ["G2", 2, "不包含 助理",    "开头是 2-, 3-", "-", "-",     "-", "直管",    "-",     "同一行各格 AND: 一层组织以 2- 或 3- 开头, 且名称不含 助理。不包含 对空值也不成立 (空值参与任何比较都不成立)"],
        ["G3", 3, "-",             "-",            "是", "-",     "-", "关键岗位", "-",     "清单文件当参照表, 存在 = 在清单里"],
        ["G4", 4, "-",             "-",            "-", ">= 23",  "-", "高职级",  "-",     "职级数 是变量定义里用 数字() 剥掉「级」后缀得到的; 单元格里只写比较"],
        ["G5", 5, "-",             "-",            "-", "-",      "是", "新推荐",  "-",     "近一年推荐 是布尔派生变量, 单元格写 是"],
        ["G9", 9, "-",             "-",            "-", "-",      "-", "-",      DISPLAY, "无条件; 多条拼接 下只有它写 显示名, 所以不发生拼接。如果() 是向量化三元, 条件为空走假分支"],
    ]
    write_rule_table(
        ws,
        top_note=["命中策略", "多条拼接", "每个输出列把所有命中规则写的值按优先级顺序用「拼接分隔符」接起来 (DMN 的 Collect); 没命中的行不动, 原值被整体替换"],
        kinds=kinds, names=names, rows=rules,
        widths=[8, 8, 18, 16, 18, 10, 12, 12, 64, 90],
    )

    # ---- sheet 5: 测试 ------------------------------------------------------
    ws = wb.create_sheet("测试")
    kinds = ["用例", "输入", "输入", "输入", "输入", "输入", "输入", "期望", "期望", "说明"]
    names = ["用例", "岗位.岗位名称", "岗位.一层组织", "关键岗位清单.存在", "岗位.职级", "岗位.推荐时间", "岗位.标签",
             "标签", "显示名", "说明"]
    tests = [
        ["T1", "集团CEO",          "1-集团总部",   "否", "22级", D(2026, 6, 1),   None,  "高管/新推荐",        "1-集团总部/集团CEO",          "两条命中, 按优先级顺序拼接"],
        ["T2", "华东代表处-行政助理", "3-华东代表处", "否", "20级", D(2025, 1, 1),   "普通", None,                "3-华东代表处/华东代表处-行政助理", "名称含 助理, G2 不命中; 标签期望留空 = 保持 普通"],
        ["T3", "华东代表处-代表",    "3-华东代表处", "是", "23级", None,            "待定", "直管/关键岗位/高职级", "3-华东代表处/华东代表处-代表",   "原标签 待定 被整体替换; 职级数 / 近一年推荐 两列不在用例里, 按变量定义推导"],
        ["T4", "法务经理",          None,          "否", "19级", D(2026, 2, 1),   None,  "新推荐",             "法务经理",                    "一层组织为空: 如果() 走假分支, 显示名只有岗位名称"],
        ["T5", "研发-SPDT主管",     "4-研发体系",   "是", "P23",  D(2024, 12, 31), None,  "关键岗位/高职级",     "4-研发体系/研发-SPDT主管",    "P23 -> 23; 推荐超过一年不算新推荐"],
        ["T6", "财务部-会计",       "5-职能平台",   "否", None,   D(2026, 10, 1),  None,  "新推荐",             "5-职能平台/财务部-会计",      "职级为空, 数字() 取不到, >= 23 不成立"],
        ["T7", "供应链-采购经理",    "6-供应链",     "否", "22级", D(2025, 10, 1),  None,  "新推荐",             "6-供应链/供应链-采购经理",     "正好 365 天, >= 成立"],
        ["T8", "供应链-采购经理",    "6-供应链",     "否", "22级", D(2025, 9, 30),  None,  None,                "6-供应链/供应链-采购经理",     "366 天, 不算; 标签期望留空 = 不变 (原本就空)"],
    ]
    write_rule_table(
        ws,
        top_note=None,
        kinds=kinds, names=names, rows=tests,
        widths=[6, 20, 14, 16, 10, 12, 8, 22, 30, 70],
    )

    wb.save(ROOT / "规则_决策表.xlsx")


if __name__ == "__main__":
    build_data()
    build_rules()
    print(f"已生成到 {ROOT}: 岗位清单.xlsx  关键岗位清单.xlsx  规则_决策表.xlsx")
