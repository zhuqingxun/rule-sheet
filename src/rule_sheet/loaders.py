"""规则文件加载: Excel 决策表 与 YAML 两种前端, 都产出 RuleSet。

规则文件声明的是「形参」: 有哪些输入 (名字、怎么对齐到主表)、变量、规则、测试; 输入的默认文件只是可选的绑定,
真正的文件可以在命令行给。加载阶段尽量把「写错了但不报错、测试假通过」的情况拦下来: 表头位置、列类型拼写、
列名缺失、测试用例的输入列 / 期望列对不上输出列或变量、重复的规则 ID、空的测试 sheet, 都在这里报错。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .conditions import all_of, any_of, build_condition, cell_condition, strings_in
from .expressions import FUNCS, find_names, split_top_level
from .model import (
    APPEND_KEY,
    APPEND_WORDS,
    BASE_DATE_VAR,
    BLANK_WORDS,
    CLEAR_WORDS,
    CONCAT_OPTS,
    CONCAT_POLICY,
    DEFAULT_CONCAT_SEP,
    DUP_POLICIES,
    EXIST_VAR,
    FALSE_WORDS,
    FIELD_GROUP_KEY,
    HIT_POLICIES,
    JOIN_POLICY,
    JOIN_SEP_KEY,
    LEGACY_INPUTS_SHEET,
    NULL_WORD,
    PRIMARY_WORD,
    SETTING_KEYS,
    SETTINGS_SHEET,
    TRUE_WORDS,
    Branch,
    Condition,
    InputSpec,
    Rule,
    RuleSet,
    TestCase,
)

RULE_KINDS = ("规则ID", "优先级", "输入", "输出", "说明")
TEST_KINDS = ("用例", "输入", "期望", "说明")
INPUT_COLS = ("名称", "匹配", "重复键", APPEND_KEY, *CONCAT_OPTS, "文件", "说明")
INPUT_KEYS = ("文件", "匹配", "重复键", APPEND_KEY, *CONCAT_OPTS, "说明")  # YAML 输入声明认的键
# YAML 各层认的键: 拼错的键 (命中测略 / 期待 / ID) 静默不生效会让规则悄悄换策略、用例假通过, 一律报错
YAML_TOP_KEYS = ("输入", "匹配", "设置", "变量", FIELD_GROUP_KEY, "命中策略", JOIN_SEP_KEY, "规则", "测试")
YAML_RULE_KEYS = ("编号", "名称", "优先级", "当", "则", "分支")
YAML_TEST_KEYS = ("用例", "输入", "期望")
INPUT_OPTS = (APPEND_KEY, *CONCAT_OPTS)  # 「输入」里除 名称 / 匹配 / 重复键 / 文件 之外的可选项, 都只对参照表有意义
FIELD_GROUP_COLS = ("名称", "列名", "说明")
POLICY_LABEL = "命中策略"
LEGACY_DEFAULT_NAMES = ("主表", "本月")  # 0.4 以前「匹配设置」里左右表别名的默认值


def _check_policy(policy: Any) -> str:
    name = str(policy or "优先级").strip()
    if name not in HIT_POLICIES:
        raise ValueError(f"未知的命中策略「{name}」, 只认: {' / '.join(HIT_POLICIES)}")
    return name


def _check_join_policy(rs: RuleSet) -> RuleSet:
    """多条拼接 下输出格写 清空 没有意义 (拼接只接有值的); 没用这个策略却设了分隔符, 多半是策略名写错了。"""
    if rs.hit_policy != JOIN_POLICY:
        if JOIN_SEP_KEY in rs.settings:
            raise ValueError(f"设置了「{JOIN_SEP_KEY}」, 但命中策略是「{rs.hit_policy}」而不是 {JOIN_POLICY}")
        return rs
    for rule in rs.rules:
        for outputs in (rule.outputs, *[b.outputs for b in rule.branches]):
            for col, v in outputs.items():
                if _text(v).lstrip("=").strip() in CLEAR_WORDS | {NULL_WORD}:
                    raise ValueError(
                        f"规则 {rule.rule_id} 输出「{col}」= {_text(v)}: "
                        f"命中策略 {JOIN_POLICY} 只拼接有值的输出, 不能写 清空"
                    )
    return rs


def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _blank(v: Any) -> bool:
    return _text(v) in BLANK_WORDS


def _same_cell(a: Any, b: Any) -> bool:
    if isinstance(a, str) or isinstance(b, str):
        return _text(a) == _text(b)
    return a == b


def _at(row: list[Any], i: int | None) -> Any:
    return row[i] if i is not None and i < len(row) else None


# ---------------------------------------------------------------------------
# 输入声明 (两种前端共用)
# ---------------------------------------------------------------------------


def _parse_keys(text: str, name: str, primary: str, names: list[str], where: str) -> list[tuple[str, str]]:
    """「匹配」格: `岗位编码` (两边同名) / `工号 = 本期.拟任人工号` (本输入的列 = 主表的列), 多个键逗号分隔。"""
    if not text:
        raise ValueError(
            f"{where}: 输入「{name}」没有写「匹配」: 主表写 {PRIMARY_WORD}, "
            f"其它输入写对齐到主表的键 (如 岗位编码, 或 工号 = {primary}.拟任人工号)"
        )
    keys: list[tuple[str, str]] = []
    for part in split_top_level(text):
        mine, _, theirs = part.partition("=")
        mine, theirs = mine.strip(), theirs.strip() or mine.strip()
        prefix, dot, rest = theirs.partition(".")
        if dot and prefix == primary:
            theirs = rest
        elif dot and prefix in names:
            raise ValueError(
                f"{where}: 输入「{name}」的匹配键「{part.strip()}」只能对齐到主表「{primary}」, 不能对齐到「{prefix}」"
            )
        keys.append((mine, theirs))
    return keys


def _parse_sort(text: str, name: str, where: str) -> list[tuple[str, bool]]:
    """「排序」格: `推荐时间 降序, 工号 升序`; 不写方向默认升序。"""
    out: list[tuple[str, bool]] = []
    for part in split_top_level(text):
        words = part.split()
        if len(words) == 1:
            out.append((words[0], True))
        elif len(words) == 2 and words[1] in ("升序", "降序"):
            out.append((words[0], words[1] == "升序"))
        else:
            raise ValueError(
                f"{where}: 输入「{name}」的「排序」写法不对 ({part.strip()!r}); "
                f"写 列名 或 列名 升序 / 列名 降序, 多个用逗号分隔"
            )
    return out


def _concat_options(opts: dict[str, Any], name: str, where: str) -> dict[str, Any]:
    """重复键: 拼接 的四个可选项 -> InputSpec 字段。"""
    numbered_text = _text(opts.get("带序号"))
    if numbered_text and numbered_text not in TRUE_WORDS | FALSE_WORDS:
        raise ValueError(f"{where}: 输入「{name}」的「带序号」只认 是 / 否, 不认识「{numbered_text}」")
    sep = opts.get("分隔符")
    return {
        "sort": _parse_sort(_text(opts.get("排序")), name, where),
        "sep": DEFAULT_CONCAT_SEP if sep in (None, "") else str(sep),
        "placeholder": _text(opts.get("空值占位")),
        "numbered": numbered_text in TRUE_WORDS,
    }


def _build_inputs(rows: list[tuple[str, Any, Any, Any, dict[str, Any]]], where: str) -> list[InputSpec]:
    """rows: (名称, 匹配, 重复键, 文件, {排序 / 分隔符 / 空值占位 / 带序号})。
    恰有一个输入的「匹配」写 主表 (只有一个输入时可省), 其它写对齐到主表的键。"""
    if not rows:
        raise ValueError(f"{where}: 没有声明任何输入")
    names = [r[0] for r in rows]
    dup = [n for n in names if names.count(n) > 1]
    if dup:
        raise ValueError(f"{where}: 输入名「{dup[0]}」重复")
    primaries = [r[0] for r in rows if _text(r[1]) == PRIMARY_WORD]
    if len(rows) == 1 and not primaries and not _text(rows[0][1]):
        primaries = [rows[0][0]]
    if len(primaries) != 1:
        raise ValueError(
            f"{where}: 必须恰有一个输入的「匹配」写 {PRIMARY_WORD} (输出的行就是它的行), "
            f"现在有 {len(primaries)} 个 ({', '.join(primaries) or '无'}); 其它输入写对齐到主表的键"
        )
    primary = primaries[0]
    specs: list[InputSpec] = []
    for name, match, dup_policy, path, opts in rows:
        file = _text(path) or None
        policy = _text(dup_policy) or "取第一条"
        given = [k for k in INPUT_OPTS if _text(opts.get(k))]
        if name == primary:
            if given or _text(dup_policy):
                raise ValueError(f"{where}: 主表「{name}」不能写「{(given or ['重复键'])[0]}」, 它们只对参照表有意义")
            specs.append(InputSpec(name, primary=True, path=file))
            continue
        if policy not in DUP_POLICIES:
            raise ValueError(f"{where}: 输入「{name}」的「重复键」只认 {' / '.join(DUP_POLICIES)}, 不认识「{policy}」")
        concat_given = [k for k in given if k in CONCAT_OPTS]
        if concat_given and policy != CONCAT_POLICY:
            raise ValueError(
                f"{where}: 输入「{name}」写了「{concat_given[0]}」, 但「重复键」是 {policy} 而不是 {CONCAT_POLICY}; "
                f"排序 / 分隔符 / 空值占位 / 带序号 只在 重复键: {CONCAT_POLICY} 时有意义"
            )
        append_text = _text(opts.get(APPEND_KEY))
        if append_text and append_text not in APPEND_WORDS:
            raise ValueError(
                f"{where}: 输入「{name}」的「{APPEND_KEY}」只认 {' / '.join(APPEND_WORDS)}, 不认识「{append_text}」"
            )
        extra = _concat_options(opts, name, where) if policy == CONCAT_POLICY else {}
        specs.append(
            InputSpec(
                name,
                keys=_parse_keys(_text(match), name, primary, names, where),
                dup_policy=policy,
                path=file,
                append=APPEND_WORDS.get(append_text, False),
                **extra,
            )
        )
    return specs


def _check_settings(settings: dict[str, Any], where: str, *, legacy_sep: bool) -> dict[str, Any]:
    """legacy_sep: 旧写法 (单独的「输入」sheet / YAML 设置:) 里的 拼接分隔符 仍认; 新版「设置」sheet 里写它报错指路。"""
    if JOIN_SEP_KEY in settings and not legacy_sep:
        raise ValueError(f"{where}: 「{JOIN_SEP_KEY}」写在决策表左上角「{POLICY_LABEL}」下面一行 (只给 多条拼接 用)")
    bad = [k for k in settings if k not in (*SETTING_KEYS, JOIN_SEP_KEY)]
    if bad:
        raise ValueError(f"{where}: 设置项「{bad[0]}」不认识, 只认: {' / '.join(SETTING_KEYS)}")
    return {k: v for k, v in settings.items() if v not in (None, "")}


def _merge_join_sep(settings: dict[str, Any], sep: Any, where: str) -> dict[str, Any]:
    """决策表左上角 / YAML 顶层的 拼接分隔符 并进 settings (引擎从那里读); 与旧写法两处都写了报错, 不猜听谁的。"""
    if sep in (None, ""):
        return settings
    if JOIN_SEP_KEY in settings:
        raise ValueError(f"「{JOIN_SEP_KEY}」写了两处 ({where} 和「设置」), 删掉「设置」里的那个")
    return {**settings, JOIN_SEP_KEY: sep}


def _legacy_inputs(match: dict[str, Any]) -> tuple[list[InputSpec], dict[str, Any]]:
    """0.4 以前的「匹配设置」/ YAML 匹配: 左表 / 右表 / 左表别名 / 右表别名 / 匹配键 / 右表重复键 / 基准日期 / 输出。"""
    la = _text(match.get("左表别名")) or LEGACY_DEFAULT_NAMES[0]
    ra = _text(match.get("右表别名")) or LEGACY_DEFAULT_NAMES[1]
    keys = [k.strip() for k in str(match.get("匹配键") or "").split(",") if k.strip()]
    policy = _text(match.get("右表重复键")) or "取第一条"
    inputs = [
        InputSpec(la, primary=True, path=_text(match.get("左表")) or None),
        InputSpec(
            ra,
            keys=[(k, k) for k in keys],
            dup_policy="取第一条" if policy == "取第一条" else "报错",
            path=_text(match.get("右表")) or None,
        ),
    ]
    settings = {k: match[k] for k in SETTING_KEYS if match.get(k) not in (None, "")}
    return inputs, settings


# ---------------------------------------------------------------------------
# 字段组 (两种前端共用): 给一组列起个名字, 决策表里以组名为列名的输出格一格表达「整组从某个输入复制 / 整组清空」
# ---------------------------------------------------------------------------


def _check_field_groups(
    groups: dict[str, list[str]], names: list[str], where: str, derived: list[str] = ()
) -> dict[str, list[str]]:
    all_cols = {c for cols in groups.values() for c in cols}
    for name, cols in groups.items():
        if not cols:
            raise ValueError(f"{where}: 字段组「{name}」没有列; 写上它包含的列名 (逗号分隔)")
        dup = [c for c in cols if cols.count(c) > 1]
        if dup:
            raise ValueError(f"{where}: 字段组「{name}」里的列「{dup[0]}」重复")
        if name in names:
            raise ValueError(f"{where}: 字段组「{name}」与输入名重复; 组名要能和输入名区分开")
        if name in derived:
            raise ValueError(f"{where}: 字段组「{name}」与变量定义里的派生变量重名")
        if name in all_cols:
            raise ValueError(f"{where}: 字段组名「{name}」同时是某个字段组里的列名; 组名不能是列名")
    return groups


_DIFF_CALL_RE = re.compile(r"差异\(\s*([^\s,()]+)\s*,\s*([^\s,()]+)\s*,\s*([^\s,()]+)\s*[,)]")


def diff_call_variables(texts: list[str], groups: dict[str, list[str]], names: list[str]) -> set[str]:
    """`差异(本期, 上期, 对比字段)` 引用到的列 (两个输入 × 组里的每一列) 也是变量: 用例能给值, 没给的视为空。"""
    out: set[str] = set()
    for t in texts:
        for new, old, group in _DIFF_CALL_RE.findall(t):
            if group in groups:
                out |= {f"{n}.{c}" for n in (new, old) if n in names for c in groups[group]}
    return out


def _read_field_groups_sheet(ws: Any) -> dict[str, list[str]]:
    """「字段组」sheet: 第 1 行列名 (名称 / 列名 / 说明), 之后每行一个组, 「列名」格里逗号分隔。"""
    grid = _rows(ws)
    header = [_text(v) for v in grid[0]] if grid else []
    col = {h: i for i, h in enumerate(header) if h}
    for need in ("名称", "列名"):
        if need not in col:
            raise ValueError(
                f"「{FIELD_GROUP_KEY}」sheet 第 1 行缺少「{need}」列 (列名: {' / '.join(FIELD_GROUP_COLS)})"
            )
    groups: dict[str, list[str]] = {}
    for row in grid[1:]:
        name = _text(_at(row, col["名称"]))
        if not name:
            continue
        if name in groups:
            raise ValueError(f"「{FIELD_GROUP_KEY}」sheet: 字段组「{name}」重复")
        groups[name] = split_top_level(_text(_at(row, col["列名"])))
    return groups


def _yaml_field_groups(doc: dict[str, Any]) -> dict[str, list[str]]:
    """YAML: 字段组: {组名: [列, 列, ...]} (也认逗号分隔的一个字符串)。"""
    groups: dict[str, list[str]] = {}
    for name, cols in (doc.get(FIELD_GROUP_KEY) or {}).items():
        if isinstance(cols, str):
            cols = split_top_level(cols)
        elif not isinstance(cols, list):
            raise ValueError(f"YAML {FIELD_GROUP_KEY}: 「{name}」的值要是列名清单 ([a, b] 或 'a, b')")
        groups[str(name)] = [str(c).strip() for c in cols]
    return groups


def expand_group_outputs(
    outputs: dict[str, Any], groups: dict[str, list[str]], names: list[str], where: str
) -> dict[str, Any]:
    """把写在字段组名下的输出格展开成逐列: 格里写输入名 = 组内每列从该输入的同名列复制 (`上期.列`),
    写 清空 = 整组清空, 留空 / `-` = 整组不改; 同一条规则显式写了 (非空) 的列以显式为准。引擎只见到展开后的逐列输出。

    Excel 里每条规则的每个输出列都在 outputs 里 (没写的是 - / None), 所以「显式」按值判, 不按键判。"""
    explicit = {c for c, v in outputs.items() if c not in groups and _text(v) not in BLANK_WORDS}
    out: dict[str, Any] = {}
    for col, v in outputs.items():
        if col not in groups:
            if col in explicit or col not in out:  # 留空的列不覆盖字段组已经展开出来的值
                out[col] = v
            continue
        text = _text(v)
        if text in BLANK_WORDS:
            continue
        if text in CLEAR_WORDS or text == NULL_WORD:
            expanded = {c: "清空" for c in groups[col]}
        elif text in names:
            expanded = {c: f"{text}.{c}" for c in groups[col]}
        else:
            raise ValueError(
                f"{where} 输出「{col}」= {text}: 字段组格只能写一个输入名 ({' / '.join(names)}, "
                f"表示组内每列从它的同名列复制)、清空 或 - (不改)"
            )
        for c, val in expanded.items():
            if c not in explicit:
                out[c] = val
    return out


def all_output_values(rules: list[Rule]) -> list[Any]:
    """规则级与各分支的输出格值 (变量扫描用)。"""
    return [v for r in rules for outs in (r.outputs, *[b.outputs for b in r.branches]) for v in outs.values()]


def expand_output_names(headers: list[str], groups: dict[str, list[str]]) -> list[str]:
    """输出列清单: 字段组名换成它的列 (在原位置), 重复的只留第一次出现。"""
    out: list[str] = []
    for h in headers:
        for c in groups.get(h, [h]):
            if c not in out:
                out.append(c)
    return out


def _allowed_test_inputs(
    variables: list[str], derived: list[tuple[str, str]], inputs: list[InputSpec], outputs: list[str]
) -> set[str]:
    """测试用例里可以给值的名字: 规则引用的变量、派生变量、内置变量 (含主表的 .存在, 追加行上为 否)、输出列的原值。"""
    primary = next(s.name for s in inputs if s.primary)
    names = set(variables) | {k for k, _ in derived} | {BASE_DATE_VAR}
    names |= {f"{s.name}.{EXIST_VAR}" for s in inputs}
    return names | {f"{primary}.{o}" for o in outputs}


def _check_test_columns(
    where: str, inputs: list[str], expects: list[str], allowed: set[str], outputs: list[str]
) -> None:
    """期望列名不是输出列时整列会被忽略、输入列名不是变量时填了也没用: 两种情况都会让用例假通过, 直接报错。"""
    bad = [n for n in expects if n not in outputs]
    if bad:
        raise ValueError(
            f"{where}: 期望列「{bad[0]}」不是输出列 (输出列: {', '.join(outputs)}); "
            f"期望列对不上时会被忽略, 用例就会假通过"
        )
    bad = [n for n in inputs if n not in allowed]
    if bad:
        raise ValueError(
            f"{where}: 输入列「{bad[0]}」不是规则引用到的变量, 填了也不起作用 (可用: {', '.join(sorted(allowed))})"
        )


# ---------------------------------------------------------------------------
# Excel 决策表
# ---------------------------------------------------------------------------


def _rows(ws: Any) -> list[list[Any]]:
    """读整张 sheet。合并单元格只有左上格有值, 这里把它展开到区域内每一格 (续行的规则 ID / 输出格常被合并)。"""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        val = _at(grid[rng.min_row - 1], rng.min_col - 1) if rng.min_row - 1 < len(grid) else None
        for r in range(rng.min_row - 1, min(rng.max_row, len(grid))):
            for c in range(rng.min_col - 1, min(rng.max_col, len(grid[r]))):
                grid[r][c] = val
    return grid


def _read_kv(ws: Any) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for row in _rows(ws)[1:]:
        if row and row[0] not in (None, "") and len(row) > 1 and row[1] not in (None, ""):
            out[str(row[0]).strip()] = row[1]
    return out


def _read_inputs_table(grid: list[list[Any]], header_row: int, where: str) -> list[InputSpec]:
    """输入表: header_row 行是列名 (名称 / 匹配 / 重复键 / 文件 / 说明, 顺序随意), 之后每行一个输入。"""
    header = [_text(v) for v in grid[header_row]] if header_row < len(grid) else []
    col = {h: i for i, h in enumerate(header) if h}
    for need in ("名称", "匹配"):
        if need not in col:
            raise ValueError(f"{where} 第 {header_row + 1} 行缺少「{need}」列 (列名: {' / '.join(INPUT_COLS)})")
    rows = []
    for row in grid[header_row + 1 :]:
        name = _text(_at(row, col["名称"]))
        if name:
            opts = {k: _at(row, col[k]) for k in INPUT_OPTS if k in col}
            rows.append((name, _at(row, col["匹配"]), _at(row, col.get("重复键")), _at(row, col.get("文件")), opts))
    return _build_inputs(rows, where)


def input_header_row(grid: list[list[Any]]) -> int | None:
    """「设置」sheet 里输入表的列名行: 第一个同时有「名称」和「匹配」两格的行 (0 起); 没有返回 None。"""
    for i, row in enumerate(grid):
        cells = {_text(v) for v in row}
        if "名称" in cells and "匹配" in cells:
            return i
    return None


def read_setting_rows(rows: list[list[Any]]) -> dict[str, Any]:
    """0.8「设置」的设置项行: 每行第一个有字的格是设置项, 紧挨着它右边一格是值, 再往右是说明; 值为空的行跳过
    (标题 / 说明行)。写着 设置项 / 值 的表头行也跳过 (从旧版剪过来的); 别的表头文字会被当成设置项报错。"""
    out: dict[str, Any] = {}
    for row in rows:
        j = next((k for k, v in enumerate(row) if _text(v)), None)
        if j is None:
            continue
        label = _text(row[j]).rstrip(":：")
        value = _at(row, j + 1)
        if label == "设置项" or value in (None, ""):
            continue
        out[label] = value
    return out


def _read_inputs_and_settings(wb: Any, path: Path) -> tuple[list[InputSpec], dict[str, Any]]:
    """0.8: 「设置」sheet 上方设置项、下方输入表。旧写法 (单独的「输入」sheet + 只有设置项的「设置」sheet) 保留一个版本;
    两处都有输入表报错, 不猜听谁的。"""
    grid = _rows(wb[SETTINGS_SHEET]) if SETTINGS_SHEET in wb.sheetnames else []
    header = input_header_row(grid)
    where = f"「{SETTINGS_SHEET}」sheet"
    if header is not None:
        if LEGACY_INPUTS_SHEET in wb.sheetnames:
            raise ValueError(
                f"{where} 里已经有输入表, 又有「{LEGACY_INPUTS_SHEET}」sheet; 输入表 0.8 起写在{where}里, "
                f"删掉「{LEGACY_INPUTS_SHEET}」sheet"
            )
        inputs = _read_inputs_table(grid, header, f"{where}的输入表")
        return inputs, _check_settings(read_setting_rows(grid[:header]), where, legacy_sep=False)
    if LEGACY_INPUTS_SHEET in wb.sheetnames:
        inputs = _read_inputs_table(_rows(wb[LEGACY_INPUTS_SHEET]), 0, f"「{LEGACY_INPUTS_SHEET}」sheet")
        # 旧「设置」照 0.7 的读法: 第 1 行无论写什么都是表头 (消费方的表头文字不一定是 设置项 / 值)
        settings = _read_kv(wb[SETTINGS_SHEET]) if SETTINGS_SHEET in wb.sheetnames else {}
        return inputs, _check_settings(settings, where, legacy_sep=True)
    if "匹配设置" in wb.sheetnames:  # 0.4 以前的写法, 保留一个版本的兼容
        return _legacy_inputs(_read_kv(wb["匹配设置"]))
    raise ValueError(
        f"规则文件缺少输入表: {where}里要有一行写列名 名称 / 匹配 (还可以有 重复键 / 文件 / 说明 等), "
        f"下面每行声明一个输入 (名字、怎么对齐到主表): {path.name}"
    )


def _parse_header(
    grid: list[list[Any]], sheet: str, key_kind: str, allowed: tuple[str, ...]
) -> tuple[list[tuple[int, str, str]], int, int]:
    """定位表头: 第一个含 key_kind (决策表「规则ID」/ 测试「用例」) 的行是列类型行, 下一行是列名行, 再下一行起是正文。

    之前的行是元信息 (命中策略 / 说明), 可有可无。
    返回 (列清单 (列索引, 类型, 名字), 列类型行号, 正文起始行号), 行号 0 起。
    """
    kind_row = next((i for i, row in enumerate(grid) if any(_text(v) == key_kind for v in row)), None)
    kinds_doc = " / ".join(allowed)
    if kind_row is None:
        raise ValueError(
            f"「{sheet}」sheet 找不到列类型行: 需要有一行标出每列的类型 ({kinds_doc}), "
            f"其中一格写「{key_kind}」, 下一行写列名"
        )
    if kind_row + 1 >= len(grid):
        raise ValueError(f"「{sheet}」sheet 第 {kind_row + 1} 行是列类型行, 但下面没有列名行")
    kinds, names = grid[kind_row], grid[kind_row + 1]
    cols: list[tuple[int, str, str]] = []
    for i, k in enumerate(kinds):
        kind = _text(k)
        if not kind:
            continue  # 没标类型的列不参与解析, 可以放备注
        at = f"{get_column_letter(i + 1)}{kind_row + 1}"
        if kind not in allowed:
            raise ValueError(f"「{sheet}」sheet {at} 格的列类型「{kind}」不认识, 只认: {kinds_doc}")
        name = _text(_at(names, i))
        if not name:
            below = f"{get_column_letter(i + 1)}{kind_row + 2}"
            raise ValueError(f"「{sheet}」sheet {at} 格标了「{kind}」列, 但 {below} 格没有写列名")
        cols.append((i, kind, name))
    keys = [i for i, k, _ in cols if k == key_kind]
    if len(keys) != 1:
        raise ValueError(f"「{sheet}」sheet 必须恰有一列类型为「{key_kind}」, 现在有 {len(keys)} 列")
    return cols, kind_row, kind_row + 2


def _find_policy(grid: list[list[Any]], kind_row: int) -> Any:
    """列类型行之上是元信息行: 找「命中策略」标签, 取它右边第一个非空格; 没有标签时认一个单独写着策略名的格。"""
    fallback = None
    for row in grid[:kind_row]:
        cells = [_text(v).rstrip(":：") for v in row]
        if POLICY_LABEL in cells:
            j = cells.index(POLICY_LABEL)
            return next((c for c in cells[j + 1 :] if c), "")
        fallback = fallback or next((c for c in cells if c in HIT_POLICIES), None)
    return fallback


def find_meta_value(grid: list[list[Any]], kind_row: int, label: str) -> Any:
    """列类型行之上的元信息行里, 写着 label 的格紧挨着的右边一格 (原值); 没有这个标签或右边为空返回 None。
    只取紧挨着的一格: 值留空时不能把再右边的说明文字当成值。"""
    for row in grid[:kind_row]:
        for j, v in enumerate(row):
            if _text(v).rstrip(":：") == label:
                value = _at(row, j + 1)
                return None if value in (None, "") else value  # 不 strip: 分隔符可以是空格
    return None


@dataclass
class RawRule:
    """决策表里一条规则的原始格值 (相邻同 ID 的续行已合并成分支); 加载器从它构造 Rule, 导出器从它翻译成表达式。"""

    rule_id: str
    priority: float
    branches: list[dict[str, Any]] = field(default_factory=list)  # 每个分支: 输入列名 -> 原始格值
    outputs: dict[str, Any] = field(default_factory=dict)  # 首行的 输出列名 -> 原始格值
    desc: str = ""
    row: int = 0  # 首行的 Excel 行号
    row_outputs: list[dict[str, Any]] = field(
        default_factory=list
    )  # 每一行 (含首行) 的输出格原文; 续行留空 = 与首行相同

    def branch_outputs(self) -> list[dict[str, Any]]:
        """每个分支生效的输出格: 续行留空的格继承首行。"""
        return [
            {col: (self.outputs.get(col) if _blank(v) else v) for col, v in outs.items()} for outs in self.row_outputs
        ]

    def has_distinct_outputs(self) -> bool:
        """续行的输出与首行不同 → 这是有各自输出的分支 (0.7, issue #20), 否则是单纯的「或」。"""
        first = {c: _text(v) for c, v in self.outputs.items()}
        return any({c: _text(v) for c, v in outs.items()} != first for outs in self.branch_outputs()[1:])


def read_decision_rows(grid: list[list[Any]], cols: list[tuple[int, str, str]], body_start: int) -> list[RawRule]:
    """走决策表正文: 相邻的同 ID 行是同一条规则的多个分支 (优先级 / 说明以首行为准, 续行优先级不一致报错;
    续行的输出留空 = 与首行相同, 写了不同的值 = 各分支有各自的输出), 不相邻的重复 ID 报错, 优先级留空取 1000 + 行序。"""
    id_col = next(i for i, k, _ in cols if k == "规则ID")
    pri_col = next((i for i, k, _ in cols if k == "优先级"), None)
    desc_col = next((i for i, k, _ in cols if k == "说明"), None)
    in_cols = [(i, n) for i, k, n in cols if k == "输入"]
    out_cols = [(i, n) for i, k, n in cols if k == "输出"]
    rules: list[RawRule] = []
    seen: dict[str, int] = {}
    for offset, row in enumerate(grid[body_start:]):
        excel_row = body_start + offset + 1
        rid = _text(_at(row, id_col))
        if not rid:
            continue
        pri_raw = _at(row, pri_col)
        outs = {name: _at(row, i) for i, name in out_cols}
        branch = {name: _at(row, i) for i, name in in_cols}
        if rules and rid == rules[-1].rule_id:
            rule = rules[-1]
            where = f"规则 {rid} 的续行 (第 {excel_row} 行)"
            head = f"首行 (第 {rule.row} 行)"
            if not _blank(pri_raw) and float(pri_raw) != rule.priority:
                raise ValueError(f"{where} 优先级与{head}不一致; 续行的优先级留空或与首行相同")
            rule.branches.append(branch)
            rule.row_outputs.append(outs)
            continue
        if rid in seen:
            raise ValueError(
                f"规则ID「{rid}」在第 {seen[rid]} 行和第 {excel_row} 行重复出现; "
                f"同一规则的多个条件分支 (或) 必须写在相邻的行"
            )
        seen[rid] = excel_row
        rules.append(
            RawRule(
                rule_id=rid,
                priority=1000.0 + offset if _blank(pri_raw) else float(pri_raw),
                branches=[branch],
                outputs=outs,
                desc=_text(_at(row, desc_col)),
                row=excel_row,
                row_outputs=[outs],
            )
        )
    return rules


def _branch_conditions(branches: list[dict[str, Any]]) -> list[Condition]:
    """每个分支内各输入格 AND, 分支之间 OR。"""
    per_branch = [[cell_condition(name, v) for name, v in b.items()] for b in branches]
    if len(per_branch) == 1:
        return per_branch[0]
    return [any_of([all_of(b) for b in per_branch])]


def _excel_rule(r: RawRule, groups: dict[str, list[str]], names: list[str]) -> Rule:
    """续行输出都与首行相同 → 一条「或」规则 (0.6); 有续行写了不同的输出 → 各分支有各自输出的规则 (0.7)。"""
    where = f"规则 {r.rule_id} (第 {r.row} 行)"
    if not r.has_distinct_outputs():
        return Rule(
            r.rule_id,
            r.priority,
            _branch_conditions(r.branches),
            expand_group_outputs(r.outputs, groups, names, where),
            r.desc,
        )
    branches = []
    for cells, outs in zip(r.branches, r.branch_outputs(), strict=True):
        expanded = expand_group_outputs(outs, groups, names, where)
        outputs = {c: v for c, v in expanded.items() if not _blank(v)}  # 分支里 - 就是不写, 与 YAML 的 则: 一致
        branches.append(Branch([cell_condition(name, v) for name, v in cells.items()], outputs))
    return Rule(r.rule_id, r.priority, [], {}, r.desc, branches=branches)


def load_excel_rules(path: Path) -> RuleSet:
    return load_excel_raw(path)[0]


def load_excel_raw(path: Path) -> tuple[RuleSet, list[RawRule]]:
    """加载 Excel 规则工作簿, 同时返回决策表的原始格值 (导出器用)。"""
    # data_only=False: 以 = 开头的格 (= 活跃 / =主表.金额 * 1.1) 在 Excel 里会被当成公式存起来,
    # data_only=True 读到的是公式的缓存结果 (没算过是 None, 算过是 #NAME?), 条件就静默变成「不关心」。
    # 这里按原文读, 规则工作簿里的 = 一律是 rule-sheet 表达式, 不支持写 Excel 公式。
    wb = load_workbook(path, data_only=False)
    if "决策表" not in wb.sheetnames:
        raise ValueError(f"规则文件缺少「决策表」sheet: {path.name}")
    inputs, settings = _read_inputs_and_settings(wb, path)
    names = [s.name for s in inputs]
    derived: list[tuple[str, str]] = []
    if "变量定义" in wb.sheetnames:
        derived = [
            (str(r[0]).strip(), str(r[1]).strip())
            for r in _rows(wb["变量定义"])[1:]
            if r and r[0] not in (None, "") and len(r) > 1 and r[1] not in (None, "")
        ]
    groups = _read_field_groups_sheet(wb[FIELD_GROUP_KEY]) if FIELD_GROUP_KEY in wb.sheetnames else {}
    _check_field_groups(groups, names, f"「{FIELD_GROUP_KEY}」sheet", [k for k, _ in derived])

    grid = _rows(wb["决策表"])
    cols, kind_row, body_start = _parse_header(grid, "决策表", "规则ID", RULE_KINDS)
    hit_policy = _check_policy(_find_policy(grid, kind_row))
    settings = _merge_join_sep(settings, find_meta_value(grid, kind_row, JOIN_SEP_KEY), "决策表左上角")
    raw = read_decision_rows(grid, cols, body_start)
    rules = [_excel_rule(r, groups, names) for r in raw]
    # 输入格原文 + 展开后的输出格 (含分支的), 之后扫描里面引用的变量 (!= 上期.x / =本期.y * 2 / 字段组展开出的 上期.z);
    # 区间 [0..本期.上限] 里变量紧跟在 .. 后面, 扫描器的名字边界不认 . 前缀, 先把 .. 隔开
    cells = [v for r in raw for b in r.branches for v in b.values()] + all_output_values(rules)
    cell_texts = [v.replace("..", " .. ") for v in cells if isinstance(v, str)]

    known = [k for k, _ in derived] + list(FUNCS) + [BASE_DATE_VAR]
    input_names = [n for _, k, n in cols if k == "输入"]
    texts = cell_texts + [e for _, e in derived]
    # 带输入前缀的派生变量 (报表7.现岗) 在用例里直接给值不推导, 也进变量清单 (没给的视为空)
    input_derived = [k for k, _ in derived if "." in k and k.partition(".")[0] in names]
    found = {n for t in texts for n in find_names(t, names, known)} | set(input_derived)
    found |= diff_call_variables(texts, groups, names)
    extra = sorted(found - set(input_names))
    variables = input_names + extra  # 表头输入列 + 单元格与派生变量表达式里引用的列, 用例没给的都视为空
    output_names = expand_output_names([n for _, k, n in cols if k == "输出"], groups)

    tests: list[TestCase] = []
    if "测试" in wb.sheetnames:
        tgrid = _rows(wb["测试"])
        tcols, _, tstart = _parse_header(tgrid, "测试", "用例", TEST_KINDS)
        case_col = next(i for i, k, _ in tcols if k == "用例")
        t_inputs = [(i, n) for i, k, n in tcols if k == "输入"]
        t_expect = [(i, n) for i, k, n in tcols if k == "期望"]
        allowed = _allowed_test_inputs(variables, derived, inputs, output_names)
        _check_test_columns("「测试」sheet", [n for _, n in t_inputs], [n for _, n in t_expect], allowed, output_names)
        for row in tgrid[tstart:]:
            if not _text(_at(row, case_col)):
                continue
            tests.append(
                TestCase(
                    name=_text(_at(row, case_col)),
                    inputs={n: _at(row, i) for i, n in t_inputs},
                    expected={n: _at(row, i) for i, n in t_expect},
                )
            )
        if not tests:
            raise ValueError("「测试」sheet 里没有任何用例 (「用例」列为空的行会跳过); 写上用例, 或删掉这张 sheet")

    rs = RuleSet(
        inputs,
        settings,
        derived,
        hit_policy,
        rules,
        output_names,
        tests,
        source=path.name,
        variables=variables,
        field_groups=groups,
    )
    return _check_join_policy(rs), raw


# ---------------------------------------------------------------------------
# YAML
# ---------------------------------------------------------------------------


def _yaml_inputs(doc: dict[str, Any]) -> tuple[list[InputSpec], dict[str, Any]]:
    """YAML: 输入: {名字: {文件, 匹配, 重复键}} + 设置: {基准日期, 输出}; 没有 输入: 时按 0.4 以前的 匹配: 折算。"""
    if "输入" in doc and "匹配" in doc:
        raise ValueError("YAML 里 输入: 和旧写法 匹配: 不能同时出现")
    if "输入" not in doc:
        inputs, settings = _legacy_inputs(dict(doc.get("匹配") or {}))
        # 设置: 不能因为没写 输入: 就被忽略
        settings.update(_check_settings(dict(doc.get("设置") or {}), "YAML 设置:", legacy_sep=True))
        return inputs, settings
    rows = []
    for name, v in (doc.get("输入") or {}).items():
        if isinstance(v, dict):
            bad = [k for k in v if k not in INPUT_KEYS]
            if bad:
                raise ValueError(f"YAML 输入「{name}」里的「{bad[0]}」不认识, 只认: {' / '.join(INPUT_KEYS)}")
            opts = {k: v.get(k) for k in INPUT_OPTS if k in v}
            rows.append((str(name), v.get("匹配"), v.get("重复键"), v.get("文件"), opts))
        else:
            rows.append((str(name), None, None, v, {}))  # 只写一个字符串 = 文件名
    inputs = _build_inputs(rows, "YAML 输入:")
    # 设置: 里的 拼接分隔符 是 0.7 的写法, 仍认; 0.8 起写在顶层 (与 命中策略: 并列), 在 load_yaml_rules 里合并
    settings = _check_settings(dict(doc.get("设置") or {}), "YAML 设置:", legacy_sep=True)
    return inputs, settings


def _yaml_branches(
    spec: Any, rid: str, defaults: dict[str, Any], groups: dict[str, list[str]], names: list[str]
) -> list[Branch]:
    """规则的 分支: [{当: ..., 则: ...}, ...]; 规则级 则 是各分支的默认值, 分支的 则 按列覆盖。"""
    if spec is None:
        return []
    if not isinstance(spec, list) or not spec:
        raise ValueError(f"规则 {rid} 的 分支: 要是一个非空列表, 每项写 当: / 则:")
    branches: list[Branch] = []
    for i, b in enumerate(spec, 1):
        if not isinstance(b, dict) or any(k not in ("当", "则", "名称") for k in b):
            raise ValueError(f"规则 {rid} 的第 {i} 个分支只认 当: / 则: / 名称:")
        conds = b.get("当") or []
        then = expand_group_outputs(
            {str(k): v for k, v in (b.get("则") or {}).items()}, groups, names, f"规则 {rid} 分支 {i}"
        )
        branches.append(Branch([build_condition(conds)] if conds else [], {**defaults, **then}))
    return branches


def _check_yaml_keys(item: Any, allowed: tuple[str, ...], where: str) -> dict[str, Any]:
    if not isinstance(item, dict):
        raise ValueError(f"{where} 要写成 键: 值 的形式, 键只认: {' / '.join(allowed)}")
    bad = [str(k) for k in item if k not in allowed]
    if bad:
        raise ValueError(f"{where} 里的「{bad[0]}」不认识, 只认: {' / '.join(allowed)}")
    return item


def load_yaml_rules(path: Path) -> RuleSet:
    doc = _check_yaml_keys(yaml.safe_load(path.read_text(encoding="utf-8")) or {}, YAML_TOP_KEYS, "YAML 顶层")
    for i, r in enumerate(doc.get("规则") or [], 1):
        _check_yaml_keys(r, YAML_RULE_KEYS, f"YAML 第 {i} 条规则")
    for i, t in enumerate(doc.get("测试") or [], 1):
        _check_yaml_keys(t, YAML_TEST_KEYS, f"YAML 第 {i} 条测试用例")
    inputs, settings = _yaml_inputs(doc)
    settings = _merge_join_sep(settings, doc.get(JOIN_SEP_KEY), "YAML 顶层")
    names = [s.name for s in inputs]
    derived = [(str(k), str(v)) for k, v in (doc.get("变量") or {}).items()]
    groups = _check_field_groups(_yaml_field_groups(doc), names, f"YAML {FIELD_GROUP_KEY}:", [k for k, _ in derived])
    rules: list[Rule] = []
    headers: list[str] = []
    for order, r in enumerate(doc.get("规则") or []):
        conds = r.get("当") or []
        then = {str(k): v for k, v in (r.get("则") or {}).items()}  # 则: 省略或为空 = 命中但什么都不改
        rid = str(r.get("编号", f"R{order + 1}"))
        if any(x.rule_id == rid for x in rules):
            raise ValueError(f"规则编号「{rid}」重复; 一条规则的多个条件分支请写在它的 当: 里 (任一满足:)")
        expanded = expand_group_outputs(then, groups, names, f"规则 {rid}")
        branches = _yaml_branches(r.get("分支"), rid, expanded, groups, names)
        rules.append(
            Rule(
                rule_id=rid,
                priority=float(r.get("优先级", 1000 + order)),
                conditions=[build_condition(conds)] if conds else [],
                outputs={} if branches else expanded,
                desc=str(r.get("名称", "")),
                branches=branches,
            )
        )
        for col in [*then, *[str(c) for b in (r.get("分支") or []) for c in (b.get("则") or {})]]:
            if col not in headers:
                headers.append(col)
    outputs = expand_output_names(headers, groups)
    # 扫描 当 / 则 (字段组已展开, 含分支) / 变量 里引用到的 <输入名>.<列>; YAML 里新增写变量名的位置要同步进 texts
    texts = [s for r in doc.get("规则") or [] for s in strings_in(r.get("当"))]
    texts += [s for r in doc.get("规则") or [] for b in (r.get("分支") or []) for s in strings_in(b.get("当"))]
    texts += [str(v) for v in all_output_values(rules)]
    texts += [v for _, v in derived]
    known = [k for k, _ in derived] + list(FUNCS) + [BASE_DATE_VAR]
    found = {n for t in texts for n in find_names(t, names, known)} | {k for k, _ in derived}
    variables = sorted(found | diff_call_variables(texts, groups, names))

    tests: list[TestCase] = []
    allowed = _allowed_test_inputs(variables, derived, inputs, outputs)
    for i, t in enumerate(doc.get("测试") or []):
        name = str(t.get("用例", f"T{i + 1}"))
        given = {str(k): v for k, v in (t.get("输入") or {}).items()}
        expected = {str(k): v for k, v in (t.get("期望") or {}).items()}
        _check_test_columns(f"测试用例「{name}」", list(given), list(expected), allowed, outputs)
        tests.append(TestCase(name, given, expected))
    if "测试" in doc and not tests:
        raise ValueError("YAML 的 测试: 下没有任何用例; 写上用例, 或删掉这个键")

    return _check_join_policy(
        RuleSet(
            inputs,
            settings,
            derived,
            _check_policy(doc.get("命中策略")),
            rules,
            outputs,
            tests,
            source=path.name,
            variables=variables,
            field_groups=groups,
        )
    )


def load_rules(path: str | Path) -> RuleSet:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"规则文件不存在: {p}")
    if p.suffix.lower() in (".yaml", ".yml"):
        return load_yaml_rules(p)
    return load_excel_rules(p)
