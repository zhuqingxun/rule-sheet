"""执行核心: 读入各输入表 -> 以主表为锚逐个合并参照表 -> 派生变量 -> 逐条规则求 mask -> 按命中策略合并 -> 审计列。"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .conditions import ns_index, resolve_output, to_mask
from .expressions import eval_expr, is_na, text_of, values_equal
from .loaders import load_rules
from .model import (
    BASE_DATE_VAR,
    CLEAR,
    CONCAT_POLICY,
    DEFAULT_JOIN_SEP,
    EXIST_VAR,
    EXPECT_NULL_WORDS,
    JOIN_POLICY,
    JOIN_SEP_KEY,
    NULL,
    NULL_WORD,
    RUNTIME_VARS,
    TRUE_WORDS,
    Condition,
    FieldGroupRef,
    InputRef,
    InputSpec,
    Namespace,
    Rule,
    RuleSet,
    TestCase,
)
from .style import display_width

Mask = pd.Series

# ---------------------------------------------------------------------------
# 读表
# ---------------------------------------------------------------------------


def resolve_path(p: str | Path, base_dir: Path | None) -> Path:
    path = Path(str(p))
    return path if path.is_absolute() or base_dir is None else base_dir / path


def split_sheet(spec: str | Path, base_dir: Path | None = None) -> tuple[Path, str | None]:
    """`路径#sheet名` 拆成 (路径, sheet名)。文件名本身也可以含 #: 整个字符串就是一个存在的文件时不拆。"""
    text = str(spec)
    whole = resolve_path(text, base_dir)
    if whole.exists() or "#" not in text:
        return whole, None
    base, sheet = text.rsplit("#", 1)
    return resolve_path(base, base_dir), sheet or None


def read_table(path: str | Path, sheet: str | None = None) -> pd.DataFrame:
    """读 Excel, 保住文本格式的数字串。

    pandas 3 的 read_excel 默认会把文本单元格里的 "007" / "13800000001" 转成数字 (前导零丢失, 电话变浮点)。
    先按 object 原样读, 再让 pandas 只对整列同类型的数据做推断: 数值列、日期列正常, 文本列保持文本。
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"数据表不存在: {p}")
    if sheet is not None:
        wb = load_workbook(p, read_only=True)
        names = wb.sheetnames
        wb.close()
        if sheet not in names:
            raise ValueError(f"数据表 {p.name} 里没有 sheet「{sheet}」, 有: {', '.join(names)}")
    return pd.read_excel(p, sheet_name=sheet if sheet is not None else 0, dtype=object).infer_objects()


# ---------------------------------------------------------------------------
# 命名空间: 以主表为锚, 把每个参照表按各自的匹配键合并进来, 变量名 <输入名>.<列>
# ---------------------------------------------------------------------------


def _add_input_derived(rs: RuleSet, spec: InputSpec, df: pd.DataFrame, base_date: pd.Timestamp) -> pd.DataFrame:
    """带输入前缀的派生变量 (`报表7.现岗`): 在那张表自己的行上逐行算, 发生在匹配 / 聚合之前, 只能引用同一张表的列。"""
    mine = [k for k, _ in rs.input_derived if k.partition(".")[0] == spec.name]
    if not mine:
        return df
    df = df.copy()
    local: Namespace = {f"{spec.name}.{c}": df[c] for c in df.columns}
    local[BASE_DATE_VAR] = base_date
    for name, expr in rs.input_derived:
        if name not in mine:
            continue
        col = name.partition(".")[2]
        if col in df.columns:
            raise ValueError(f"派生变量「{name}」与输入「{spec.name}」的列「{col}」同名; 换个名字")
        try:
            val = eval_expr(expr, local)
        except ValueError as exc:
            raise ValueError(f"派生变量「{name}」在输入「{spec.name}」的行上求值失败: {exc}") from exc
        if val is NULL:
            val = None
        df[col] = val if isinstance(val, pd.Series) else pd.Series([val] * len(df), index=df.index)
        local[name] = df[col]
    return df


def _filter_rows(spec: InputSpec, df: pd.DataFrame, base_date: pd.Timestamp, notes: list[str] | None) -> pd.DataFrame:
    """筛选: 在这张表自己的行上算表达式, 只留成立的行 (空值不成立)。筛完一行不剩报错: 多半是常量写错了,
    不报错的话这个输入的 .存在 会整列静默变成 否。"""
    if not spec.filter or df.empty:
        return df
    local: Namespace = {f"{spec.name}.{c}": df[c] for c in df.columns}
    local[BASE_DATE_VAR] = base_date
    try:
        keep = to_mask(eval_expr(spec.filter, local), local)
    except (ValueError, KeyError) as exc:
        msg = exc.args[0] if exc.args else exc
        raise ValueError(f"输入「{spec.name}」的筛选「{spec.filter}」求值失败: {msg}") from exc
    kept = int(keep.sum())
    if kept == 0:
        raise ValueError(
            f"输入「{spec.name}」按筛选「{spec.filter}」一行都没留下 (原有 {len(df)} 行); "
            f"多半是值写错了 (多了空格、全角半角不一致), 对照数据表里那一列的实际取值"
        )
    if notes is not None:
        notes.append(f"输入「{spec.name}」按筛选保留 {kept} / {len(df)} 行")
    return df[keep.to_numpy()]


def _concat_group(values: list[Any], spec: InputSpec) -> str | None:
    """重复键: 拼接 下一个键的一列 -> 一个文本: 空格子换成占位符 (位置保留), 可带序号; 没设占位符且整组全空则为空。"""
    items = [None if is_na(v) else text_of(v) for v in values]
    if not spec.placeholder and all(i is None for i in items):
        return None
    texts = [spec.placeholder if i is None else i for i in items]
    if spec.numbered:
        texts = [f"{n}.{t}" for n, t in enumerate(texts, 1)]
    return spec.sep.join(texts)


def _concat_rows(spec: InputSpec, df: pd.DataFrame, mine: list[str]) -> pd.DataFrame:
    """同键的行按 spec.sort 排 (空值一律排末尾), 每个非键列各自接成一个文本, 一个键一行。"""
    if spec.sort:
        missing = [c for c, _ in spec.sort if c not in df.columns]
        if missing:
            raise ValueError(f"输入「{spec.name}」没有排序列「{missing[0]}」")
        df = df.sort_values(
            by=[c for c, _ in spec.sort], ascending=[a for _, a in spec.sort], na_position="last", kind="stable"
        )
    cols = [c for c in df.columns if c not in mine]
    if not cols:
        return df.drop_duplicates(subset=mine)
    grouped = df.groupby(mine, sort=False)[cols].agg(lambda s: _concat_group(s.tolist(), spec))
    return grouped.reset_index()


def _prepare_lookup(
    rs: RuleSet,
    spec: InputSpec,
    df: pd.DataFrame,
    base_cols: list[str],
    base_date: pd.Timestamp,
    notes: list[str] | None,
) -> tuple[pd.DataFrame, list[str], list[str]]:
    """一个参照表在合并前的准备: 校验键列、剔掉键为空的行、算带前缀的派生变量、按「重复键」去重或聚合。"""
    primary = rs.primary
    if not spec.keys:
        raise ValueError(f"输入「{spec.name}」没有匹配键, 不知道怎么对齐到主表「{primary.name}」")
    mine = [m for m, _ in spec.keys]
    theirs = [t for _, t in spec.keys]
    for c in theirs:
        if c not in base_cols:
            raise ValueError(f"主表「{primary.name}」没有列「{c}」, 而输入「{spec.name}」要按它对齐")
    for c in mine:
        if c not in df.columns:
            raise ValueError(f"输入「{spec.name}」没有匹配键列「{c}」")
    df = df[df[mine].notna().all(axis=1)]  # 键为空的行不参与匹配: pandas merge 会把两边的 NaN 配成一对
    df = _add_input_derived(rs, spec, df, base_date)
    df = _filter_rows(spec, df, base_date, notes)  # 在去重 / 聚合之前: 别的类别里的同名键不算重复
    dup = df.duplicated(subset=mine, keep="first")
    if dup.any():
        if spec.dup_policy == CONCAT_POLICY:
            if notes is not None:
                notes.append(f"输入「{spec.name}」有 {int(dup.sum())} 行重复键, 按「{CONCAT_POLICY}」接成一行")
        elif spec.dup_policy != "取第一条":
            raise ValueError(
                f"输入「{spec.name}」有 {int(dup.sum())} 行重复键, 而它的「重复键」设置为 {spec.dup_policy}"
            )
        else:
            if notes is not None:
                notes.append(f"输入「{spec.name}」有 {int(dup.sum())} 行重复键, 按「取第一条」处理")
            df = df[~dup]
    if spec.dup_policy == CONCAT_POLICY:
        df = _concat_rows(spec, df, mine)  # 没有重复键也走一遍: 单条记录同样要占位 / 带序号, 形态一致
    return df, mine, theirs


def _append_only_rows(
    spec: InputSpec,
    df: pd.DataFrame,
    mine: list[str],
    theirs: list[str],
    base: pd.DataFrame,
    exists: pd.Series,
    notes: list[str] | None,
) -> tuple[pd.DataFrame, pd.Series]:
    """独有行: 追加 — 参照表里键在主表 (含前面已追加的行) 找不到的行接到行集末尾:
    键列按匹配映射填, 同名列带入, 主表独有的列为空; 主表 .存在 在这些行上为 否。"""
    keyed = base[theirs].drop_duplicates()
    flags = df[mine].merge(keyed, left_on=mine, right_on=theirs, how="left", indicator=True)["_merge"]
    only = df.loc[flags.eq("left_only").to_numpy()]
    if only.empty:
        return base, exists
    new = pd.DataFrame(index=pd.RangeIndex(len(only)), columns=base.columns, dtype=object)
    for m_col, t_col in zip(mine, theirs, strict=True):
        new[t_col] = only[m_col].to_numpy()
    for col in base.columns:
        if col in only.columns and col not in theirs:
            new[col] = only[col].to_numpy()
    if notes is not None:
        notes.append(f"输入「{spec.name}」有 {len(only)} 行主表没有的键, 按「追加」接在末尾")
    grown = pd.concat([base, new], ignore_index=True).infer_objects()
    flag = pd.concat([exists, pd.Series(False, index=pd.RangeIndex(len(only)))], ignore_index=True)
    return grown, flag


def resolve_base_date(rs: RuleSet, variables: dict[str, Any] | None = None) -> tuple[pd.Timestamp, str]:
    """真实数据用的基准日期与来源: 运行时给的 (variables) 优先, 其次规则文件「设置」, 都没有取运行当天。

    variables 的键只认 RUNTIME_VARS: 拼错的名字静默不生效会让结果悄悄用回规则文件里的值, 所以报错。
    内嵌测试不走这里 (case_namespace 用规则文件的值), 用例里的日期是按它写的。"""
    variables = dict(variables or {})
    unknown = [k for k in variables if k not in RUNTIME_VARS]
    if unknown:
        raise ValueError(f"运行时变量「{unknown[0]}」不认识, 只认: {' / '.join(RUNTIME_VARS)}")
    if BASE_DATE_VAR in variables:
        raw, source = variables[BASE_DATE_VAR], "运行时给定"
    elif rs.settings.get(BASE_DATE_VAR):
        raw, source = rs.settings[BASE_DATE_VAR], "规则文件「设置」"
    else:
        return pd.Timestamp(date.today()), "规则文件没设, 取运行当天"
    try:
        value = pd.Timestamp(raw)
    except (ValueError, TypeError):
        value = pd.NaT
    if pd.isna(value):
        raise ValueError(f"{BASE_DATE_VAR}「{raw}」({source}) 不是日期, 写成 2026-10-23 这样")
    return value, source


def build_namespace(
    rs: RuleSet,
    tables: dict[str, pd.DataFrame],
    notes: list[str] | None = None,
    *,
    variables: dict[str, Any] | None = None,
) -> Namespace:
    """以主表为锚, 每个参照表按各自的键 left join 进来。

    匹配之前: 每张表自己的行上先算带该表前缀的派生变量 (`报表7.现岗`); `重复键: 拼接` 的参照表按键聚合成一键一行;
    `独有行: 追加` 的参照表把主表没有的键追加到行集末尾 (按声明顺序, 后面的对齐到已扩的行集),
    主表 `.存在` 在这些行上为 否。直接调 apply_rules 的编排器必须从这里拿命名空间, 聚合 / 追加都发生在这一步;
    扩展后的主表用 primary_frame 取。
    variables: 运行时给的值 (目前只有 基准日期), 覆盖规则文件「设置」里的同名项, 见 resolve_base_date。
    """
    primary = rs.primary
    base_date, source = resolve_base_date(rs, variables)
    if notes is not None:
        notes.append(f"{BASE_DATE_VAR} = {text_of(base_date)} ({source})")
    base = tables[primary.name]
    base_cols = list(base.columns)
    prepared = {s.name: _prepare_lookup(rs, s, tables[s.name], base_cols, base_date, notes) for s in rs.lookups}
    exists = pd.Series(True, index=base.index)
    for spec in rs.lookups:
        if spec.append:
            df, mine, theirs = prepared[spec.name]
            base, exists = _append_only_rows(spec, df, mine, theirs, base, exists, notes)
    base = _add_input_derived(rs, primary, base, base_date)
    ns: Namespace = {f"{primary.name}.{col}": base[col] for col in base.columns}
    ns[f"{primary.name}.{EXIST_VAR}"] = exists
    for spec in rs.lookups:
        df, mine, theirs = prepared[spec.name]
        # 左侧只带主表的键列, 参照表之间同名列互不干扰; 参照表里与主表键列同名的非键列保留参照表的 (主表的加后缀)
        merged = base[theirs].merge(df, left_on=theirs, right_on=mine, how="left", suffixes=("__p", ""), indicator=True)
        merged.index = base.index
        for col in df.columns:
            if col not in mine:
                ns[f"{spec.name}.{col}"] = merged[col]
        ns[f"{spec.name}.{EXIST_VAR}"] = merged["_merge"].eq("both")
    ns[BASE_DATE_VAR] = base_date
    add_refs(rs, ns)
    return ns


def add_refs(rs: RuleSet, ns: Namespace) -> None:
    """输入名与字段组名本身也进命名空间 (InputRef / FieldGroupRef), 给 差异(本期, 上期, 对比字段) 这类函数当参数用。"""
    for s in rs.inputs:
        ns[s.name] = InputRef(s.name)
    for g, cols in rs.field_groups.items():
        ns[g] = FieldGroupRef(g, tuple(cols))


def primary_frame(rs: RuleSet, ns: Namespace, columns: list[str]) -> pd.DataFrame:
    """从命名空间取回 (可能已追加了独有行的) 主表: columns 是主表原有的列。输出的行就是它的行。"""
    return pd.DataFrame({c: ns[f"{rs.primary.name}.{c}"] for c in columns})


def add_derived(rs: RuleSet, ns: Namespace, index: pd.Index, *, lenient: bool) -> None:
    """按定义顺序求派生变量。测试用例里直接给了值的不再推导; lenient 下求值失败置空。

    带输入前缀的派生变量不在这里算: 真实数据里 build_namespace 已在那张表的行上算过, 用例里直接给值。"""
    for name, expr in rs.plain_derived:
        if name in ns:
            continue
        try:
            val = eval_expr(expr, ns)
        except ValueError:
            if not lenient:
                raise
            val = float("nan")
        if val is NULL:
            val = None
        if not isinstance(val, pd.Series):
            val = pd.Series([val] * len(index), index=index)
        ns[name] = val


# ---------------------------------------------------------------------------
# 规则执行
# ---------------------------------------------------------------------------


def _conditions_mask(conds: list[Condition], rule_id: str, ns: Namespace, index: pd.Index) -> Mask:
    m = pd.Series(True, index=index)
    for cond in conds:
        try:
            res = cond(ns)
        except KeyError as exc:
            raise KeyError(f"规则 {rule_id}: {exc.args[0]}") from exc
        m = m & to_mask(res, ns)
    return m


def branch_masks(rule: Rule, ns: Namespace, index: pd.Index) -> list[Mask]:
    """每个分支的 mask (规则级条件 AND 分支条件); 没有分支时为空列表。"""
    common = _conditions_mask(rule.conditions, rule.rule_id, ns, index)
    return [common & _conditions_mask(b.conditions, rule.rule_id, ns, index) for b in rule.branches]


def rule_mask(rule: Rule, ns: Namespace, index: pd.Index) -> Mask:
    """规则整体的 mask: 规则级条件, 有分支时再 AND 任一分支成立。"""
    common = _conditions_mask(rule.conditions, rule.rule_id, ns, index)
    if not rule.branches:
        return common
    any_branch = pd.Series(False, index=index)
    for b in branch_masks(rule, ns, index):
        any_branch = any_branch | b
    return common & any_branch


def _as_column(v: Any, index: pd.Index) -> pd.Series:
    if isinstance(v, pd.Series):
        return v.astype(object)
    return pd.Series([v] * len(index), index=index, dtype=object)


def _texts(s: pd.Series) -> pd.Series:
    return pd.Series([None if is_na(v) else text_of(v) for v in s.tolist()], index=s.index, dtype=object)


def _join_hits(rs: RuleSet, col: str, ordered: list[Rule], rule_value: Any, cur: pd.Series) -> pd.Series:
    """多条拼接: 按优先级顺序把每条命中规则对该列写的值接到后面; 清空 / 算出空值的规则不贡献。"""
    sep = str(rs.settings.get(JOIN_SEP_KEY, DEFAULT_JOIN_SEP))
    acc = pd.Series([None] * len(cur), index=cur.index, dtype=object)
    for rule in ordered:
        val, write = rule_value(rule, col)
        if val is None:
            continue
        texts = _texts(val)
        add = write & texts.notna()
        prefix = acc.map(lambda s: "" if s is None else s + sep)
        acc = acc.mask(add, prefix + texts.fillna(""))
    return cur.mask(acc.notna(), acc)


def apply_rules(
    rs: RuleSet,
    ns: Namespace,
    originals: dict[str, pd.Series],
    labels: pd.Series | None = None,
) -> tuple[dict[str, pd.Series], pd.Series]:
    """返回 (每个输出列的新值, 每行命中的规则 ID 串)。

    命中策略 优先级: 每个输出列取优先级数字最小的命中规则的值, 未写该列的规则不参与。
    命中策略 唯一: 任何行命中多条即报错。
    命中策略 首次命中: 每行只算优先级数字最小的那条命中 (同优先级按书写顺序), 其余规则对该行视为未命中;
        之后走与 优先级 相同的逐列合并, 此时每行至多一条规则, 结果就是该规则的输出 (- 的列保持原值)。
    命中策略 多条拼接: 每个输出列把所有命中规则写的值按优先级顺序 (同优先级按书写顺序) 用分隔符拼成一个文本,
        替换原值; 没写该列的规则不参与, 没有任何命中的行不动。
    规则内的分支 (Rule.branches): 规则整体参与上面的策略 (计一次, 审计记规则 ID); 规则对一行写什么由分支决定:
        多条拼接 下所有成立分支的值按书写顺序接起来 (相同值只记一次), 其它策略下取书写顺序第一个成立的分支。
    """
    index = ns_index(ns)
    dup = sorted({r.rule_id for r in rs.rules if sum(x.rule_id == r.rule_id for x in rs.rules) > 1})
    if dup:  # 两个加载器都已拦截, 这里兜底: masks 按 ID 建字典, 重复 ID 会静默丢掉一条规则
        raise ValueError(f"规则 ID 重复: {', '.join(dup)}")
    per_branch = {r.rule_id: branch_masks(r, ns, index) for r in rs.rules}
    masks: dict[str, Mask] = {}
    for r in rs.rules:
        if r.branches:
            any_branch = pd.Series(False, index=index)
            for b in per_branch[r.rule_id]:
                any_branch = any_branch | b
            masks[r.rule_id] = any_branch
        else:
            masks[r.rule_id] = _conditions_mask(r.conditions, r.rule_id, ns, index)
    ordered = sorted(rs.rules, key=lambda r: r.priority)  # 稳定排序: 同优先级保持书写顺序

    if rs.hit_policy == "唯一":
        hits = sum((masks[r.rule_id].astype(int) for r in rs.rules), pd.Series(0, index=index))
        bad = hits[hits > 1]
        if not bad.empty:
            detail = {
                str(labels.loc[i]) if labels is not None else f"第{int(i) + 2}行": [
                    r.rule_id for r in ordered if masks[r.rule_id].loc[i]
                ]
                for i in bad.index
            }
            raise ValueError(f"命中策略为「唯一」, 但以下行命中多条规则: {detail}")
    elif rs.hit_policy == "首次命中":
        taken = pd.Series(False, index=index)
        for rule in ordered:
            effective = masks[rule.rule_id] & ~taken
            taken = taken | effective
            masks[rule.rule_id] = effective

    sep = str(rs.settings.get(JOIN_SEP_KEY, DEFAULT_JOIN_SEP))
    nothing = pd.Series(False, index=index)

    def output_of(rule: Rule, outputs: dict[str, Any], col: str) -> Any:
        try:
            return resolve_output(outputs.get(col), ns)
        except KeyError as exc:
            raise KeyError(f"规则 {rule.rule_id} 输出「{col}」: {exc.args[0]}") from exc

    def rule_value(rule: Rule, col: str) -> tuple[pd.Series | None, Mask]:
        """规则对该列每行写的值 (object 列, 清空 已换成 None) 与写入 mask; 不写该列返回 (None, 全 False)。"""
        if not rule.branches:
            val = output_of(rule, rule.outputs, col)
            if val is None:
                return None, nothing
            return _as_column(None if val is CLEAR else val, index), masks[rule.rule_id]
        effective = [b & masks[rule.rule_id] for b in per_branch[rule.rule_id]]  # 首次命中 收窄过的规则 mask
        values = [output_of(rule, b.outputs, col) for b in rule.branches]
        if all(v is None for v in values):
            return None, nothing
        columns = [None if v is None else _as_column(None if v is CLEAR else v, index) for v in values]
        if rs.hit_policy == JOIN_POLICY:  # 所有成立分支的值按书写顺序接起来, 相同值只记一次
            texts = [None if c is None else _texts(c) for c in columns]
            rows: list[str | None] = []
            for i in range(len(index)):
                items: list[str] = []
                for m, t in zip(effective, texts, strict=True):
                    if t is None or not m.iloc[i]:
                        continue
                    v = t.iloc[i]
                    if v is not None and v not in items:
                        items.append(v)
                rows.append(sep.join(items) if items else None)
            joined = pd.Series(rows, index=index, dtype=object)
            return joined, joined.notna()
        val = pd.Series([None] * len(index), index=index, dtype=object)
        write = nothing.copy()
        seen = nothing.copy()
        for m, c in zip(effective, columns, strict=True):  # 书写顺序第一个成立的分支决定输出
            own = m & ~seen
            seen = seen | m
            if c is None:
                continue
            val = val.mask(own, c)
            write = write | own
        return val, write

    results: dict[str, pd.Series] = {}
    for col in rs.outputs:
        cur = originals.get(col)
        if cur is None:
            cur = pd.Series([None] * len(index), index=index, dtype=object)
        cur = cur.astype(object)
        if rs.hit_policy == JOIN_POLICY:
            results[col] = _join_hits(rs, col, ordered, rule_value, cur)
            continue
        for rule in reversed(ordered):  # 优先级数字大的先写, 小的最后写并覆盖
            val, write = rule_value(rule, col)
            if val is None:
                continue
            cur = cur.mask(write, val)
        results[col] = cur

    hit_ids = pd.Series(
        ["; ".join(r.rule_id for r in ordered if masks[r.rule_id].loc[i]) for i in index],
        index=index,
    )
    return results, hit_ids


# ---------------------------------------------------------------------------
# 比较与格式化
# ---------------------------------------------------------------------------


def same(a: Any, b: Any) -> bool:
    """变更判定 (审计列 / 用例期望 / diff 交叉验证) 的相等: 见 expressions.values_equal, 空串与空值视为相同 (0.7)。"""
    return values_equal(a, b)


def _tidy_nulls(s: pd.Series) -> pd.Series:
    """输出列 (apply_rules 之后总是 object) 里的空值统一成 None, 不再 NaN / None / NaT 混用 (Python API 的 res.out)。"""
    return pd.Series([None if is_na(v) else v for v in s.tolist()], index=s.index, dtype=object)


def fmt(x: Any) -> str:
    return "空" if is_na(x) else text_of(x)


# ---------------------------------------------------------------------------
# 规则级测试用例
# ---------------------------------------------------------------------------


@dataclass
class TestOutcome:
    name: str
    passed: bool
    hit: str = ""
    problems: list[str] = field(default_factory=list)
    error: str = ""


def case_namespace(rs: RuleSet, tc: TestCase) -> tuple[Namespace, dict[str, pd.Series]]:
    """把一条用例构造成 1 行的命名空间 (含派生变量) 和各输出列的原值; run_tests 与 diff 的交叉验证共用。"""
    exist_vars = {f"{s.name}.{EXIST_VAR}" for s in rs.inputs}
    derived_names = {d for d, _ in rs.plain_derived}  # 带输入前缀的派生变量按普通变量对待: 给了用, 没给为空
    index = pd.RangeIndex(1)
    ns: Namespace = {}
    for name, val in tc.inputs.items():
        if name in exist_vars and isinstance(val, str):
            val = val.strip() in TRUE_WORDS
        if isinstance(val, str) and val.strip() in ("", NULL_WORD):  # 留空或写 空 = 空值
            val = None
        if val is None and name in derived_names:
            continue  # 派生变量留空 = 没给, 按表达式推导 (Excel 测试 sheet 每列每行都在, 留空不能算「给了空值」)
        if isinstance(val, (date, datetime)):  # Excel 日期格 / YAML 日期 -> 与真实数据的日期列同型, 才能做大小比较
            val = pd.Timestamp(val)
        dtype = object if isinstance(val, str) or val is None else None
        ns[name] = pd.Series([val], index=index, dtype=dtype)
    ns.setdefault(BASE_DATE_VAR, resolve_base_date(rs)[0])  # 不传 variables: 用例永远按规则文件里的基准日期
    # 主表的 .存在 没给默认 是 (普通行); 要测追加来的行就写 否
    ns.setdefault(f"{rs.primary.name}.{EXIST_VAR}", pd.Series([True], index=index))
    for name in rs.variables:  # 用例没提到的变量 = 空值
        if name not in ns and name not in derived_names:
            ns[name] = pd.Series([None], index=index, dtype=object)
    add_refs(rs, ns)
    add_derived(rs, ns, index, lenient=True)
    originals = {
        col: ns.get(f"{rs.primary.name}.{col}", pd.Series([None], index=index, dtype=object)) for col in rs.outputs
    }
    return ns, originals


def run_tests(rs: RuleSet) -> list[TestOutcome]:
    outcomes: list[TestOutcome] = []
    for tc in rs.tests:
        ns, originals = case_namespace(rs, tc)
        try:
            results, hit_ids = apply_rules(rs, ns, originals)
        except Exception as exc:  # noqa: BLE001
            outcomes.append(TestOutcome(tc.name, False, error=str(exc)))
            continue
        problems = []
        for col in rs.outputs:
            exp = tc.expected.get(col)
            got = results[col].iloc[0]
            if exp is None or (isinstance(exp, str) and exp.strip() == ""):
                want = originals[col].iloc[0]  # 期望留空 = 该列不变
            elif isinstance(exp, str) and exp.strip() in EXPECT_NULL_WORDS:
                want = None  # 期望写 空 / 清空 / 为空 = 该列应为空值
            else:
                want = exp
            if not same(got, want):
                problems.append(f"{col}: 期望 {fmt(want)}, 实际 {fmt(got)}")
        outcomes.append(TestOutcome(tc.name, not problems, hit=hit_ids.iloc[0], problems=problems))
    return outcomes


# ---------------------------------------------------------------------------
# 输出与整体流程
# ---------------------------------------------------------------------------


def _is_date_column(s: pd.Series) -> bool:
    """日期列: datetime 类型, 或 (规则写过之后变成 object 的) 非空值全是日期的列。"""
    if pd.api.types.is_datetime64_any_dtype(s):
        return True
    values = [v for v in s.tolist() if not is_na(v)]
    return bool(values) and all(isinstance(v, (pd.Timestamp, datetime, date)) for v in values)


def write_xlsx_safely(df: pd.DataFrame, path: Path) -> Path:
    """目标被 Excel 占用时自动换 _v2/_v3 后缀, 不中断。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    candidates = [path] + [path.with_name(f"{path.stem}_v{i}{path.suffix}") for i in range(2, 10)]
    for p in candidates:
        try:
            with pd.ExcelWriter(p, engine="openpyxl") as xw:
                df.to_excel(xw, index=False)
                ws = xw.sheets[next(iter(xw.sheets))]
                for i, col in enumerate(df.columns, start=1):
                    # 和在 Excel 里双击列边自动调整一样: 看全部行, 不换行, 上限是 Excel 的 255
                    width = max([display_width(col)] + [display_width(v) for v in df[col].tolist()]) + 2
                    ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 255)
                    if _is_date_column(df[col]):
                        for r in range(2, len(df) + 2):
                            ws.cell(row=r, column=i).number_format = "yyyy-mm-dd"
                ws.freeze_panes = "B2"
            return p
        except PermissionError:
            continue
    raise PermissionError(f"无法写入 {path} 及其备选文件名, 请关闭占用的 Excel 窗口")


def bind_inputs(rs: RuleSet, base_dir: Path, overrides: dict[str, str] | None) -> dict[str, tuple[Path, str | None]]:
    """每个声明的输入对应一个文件: 命令行给的 (相对当前目录) 优先, 否则规则文件里的默认文件 (相对规则文件目录)。"""
    overrides = dict(overrides or {})
    unknown = [n for n in overrides if n not in rs.names]
    if unknown:
        raise ValueError(f"命令行给的输入「{unknown[0]}」不是规则声明的输入 (规则声明了: {', '.join(rs.names)})")
    bound: dict[str, tuple[Path, str | None]] = {}
    missing: list[str] = []
    for spec in rs.inputs:
        if spec.name in overrides:
            bound[spec.name] = split_sheet(overrides[spec.name])
        elif spec.path:
            bound[spec.name] = split_sheet(spec.path, base_dir)
        else:
            missing.append(spec.name)
    if missing:
        raise ValueError(
            f"输入「{'」「'.join(missing)}」没有绑定文件: 规则文件里没写默认文件, 命令行也没给 (-i 名称=路径)"
        )
    return bound


@dataclass
class RunResult:
    ruleset: RuleSet
    tests: list[TestOutcome]
    keys: list[str] = field(default_factory=list)  # 主表上用来标识一行的列 (第一个参照表的匹配键, 没有就是第一列)
    out: pd.DataFrame | None = None
    changed: pd.DataFrame | None = None
    written: Path | None = None
    stopped_by_tests: bool = False
    notes: list[str] = field(default_factory=list)
    inputs: dict[str, tuple[Path, int]] = field(default_factory=dict)  # 输入名 -> (实际读的文件, 行数)

    @property
    def tests_passed(self) -> int:
        return sum(1 for t in self.tests if t.passed)


def run(
    rules_path: str | Path,
    *,
    base_dir: str | Path | None = None,
    inputs: dict[str, str] | None = None,
    output: str | Path | None = None,
    variables: dict[str, Any] | None = None,
    ignore_tests: bool = False,
    dry_run: bool = False,
) -> RunResult:
    """完整流程: 加载规则 -> 跑内嵌测试 -> 绑定并读入各输入表 -> 执行 -> 写出。

    inputs: {输入名: 文件路径 (可带 #sheet名)}, 覆盖规则文件里的默认文件, 相对当前目录。
    variables: {名字: 值}, 目前只认 基准日期; 覆盖规则文件「设置」里的同名项, 只作用于真实数据
    (内嵌测试仍用规则文件的值)。
    规则文件里的相对路径 (默认文件 / 输出) 相对 base_dir, 默认为规则文件所在目录。
    """
    rules_path = Path(rules_path)
    base = Path(base_dir) if base_dir else rules_path.parent
    rs = load_rules(rules_path)
    resolve_base_date(rs, variables)  # 键名 / 日期写错在跑测试和读表之前就报
    outcomes = run_tests(rs) if rs.tests else []
    if any(not o.passed for o in outcomes) and not ignore_tests:
        return RunResult(rs, outcomes, stopped_by_tests=True)

    bound = bind_inputs(rs, base, inputs)
    tables = {name: read_table(p, sheet) for name, (p, sheet) in bound.items()}
    columns = list(tables[rs.primary.name].columns)
    clash = [g for g in rs.field_groups if g in columns]
    if clash:  # 组名在决策表里占一个输出列的位置, 加载时已展开成逐列; 主表真有同名列的话那一列永远写不到, 不能静默
        raise ValueError(
            f"字段组「{clash[0]}」与主表「{rs.primary.name}」的列同名; "
            f"组名只是决策表里的列名, 不能和真实的列撞名, 换个组名"
        )

    notes: list[str] = []
    ns = build_namespace(rs, tables, notes, variables=variables)
    primary_df = primary_frame(rs, ns, columns)  # 独有行: 追加 之后行集比读入的主表多, 输出的行就是它的行
    add_derived(rs, ns, primary_df.index, lenient=False)
    originals = {col: primary_df[col] for col in rs.outputs if col in primary_df.columns}
    first = rs.lookups[0].keys if rs.lookups else []
    keys = [t for _, t in first] or [str(primary_df.columns[0])]
    results, hit_ids = apply_rules(rs, ns, originals, labels=primary_df[keys[0]])

    out = primary_df.copy()
    details: list[str] = []
    for i in primary_df.index:
        parts = []
        for col in rs.outputs:
            old = primary_df[col].loc[i] if col in primary_df.columns else None
            new = results[col].loc[i]
            if not same(old, new):
                parts.append(f"{col}: {fmt(old)} -> {fmt(new)}")
        details.append("; ".join(parts))
    for col in rs.outputs:
        out[col] = _tidy_nulls(results[col])
    out["命中规则"] = hit_ids
    out["变更明细"] = details
    changed = out[out["变更明细"] != ""]

    written = None
    if not dry_run:
        if output:
            out_path = Path(output)
        else:
            default = rs.settings.get("输出") or f"output/{bound[rs.primary.name][0].stem}_更新后.xlsx"
            out_path = resolve_path(default, base)
        written = write_xlsx_safely(out, out_path)
    read_info = {name: (p, len(tables[name])) for name, (p, _) in bound.items()}
    return RunResult(rs, outcomes, keys, out, changed, written, notes=notes, inputs=read_info)
