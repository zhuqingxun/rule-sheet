"""两种写法之间的转换与比对。

- Excel 决策表 -> YAML (export_workbook): 无损。单元格文法的每一种形态都对应一条表达式, 翻译函数对照
  conditions.unary_test 逐分支写, 续行翻成 任一满足。导出的 YAML 是生成物: 源仍是 Excel, 改了 Excel 重新导出。
- YAML -> Excel (yaml_to_workbook): 只对「表格子集」成立, 每个条件必须是一个变量上的一元测试 (比较 / 区间 /
  属于 / 文本函数 / 判空 / 裸布尔名), 顶层 任一满足 翻成续行; 其它写法报错并指向变量定义。
- 两份规则文件的语义比对 (diff_rules): 结构逐项比 (输入 / 设置 / 字段组 / 变量 / 命中策略 / 规则 / 输出列 / 用例);
  条件是编译好的函数比不了文本, 改为把两边的测试用例合并, 每条在两套规则上各跑一次, 比实际输出和命中 ID
  (按用例交叉验证, 是抽样等价, 不是证明)。
"""

from __future__ import annotations

import ast
import numbers
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook

from .conditions import _CMP_RE, _NULL_TESTS, _RANGE_RE, _TEXT_TEST_RE
from .engine import apply_rules, case_namespace, fmt, same
from .expressions import FUNCS, TEXT_FUNCS, split_top_level, substitute_names
from .loaders import RawRule, load_excel_raw, load_rules, load_yaml_rules
from .model import (
    BASE_DATE_VAR,
    BLANK_WORDS,
    EXIST_VAR,
    FALSE_WORDS,
    FIELD_GROUP_KEY,
    NULL_WORD,
    PRIMARY_WORD,
    TRUE_WORDS,
    RuleSet,
    TestCase,
)
from .style import style_rules_workbook

_DATE_TEXT_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_CMP_OPS = {"=": "==", "==": "==", "<>": "!=", "!=": "!=", ">": ">", ">=": ">=", "<": "<", "<=": "<="}
# 名字里出现这些就得用反引号包起来: ASCII 标点 (除 . - /)、空白, 以及 - 或 / 后面紧跟数字 (会被当成减法 / 除法)
_NEEDS_BACKTICKS_RE = re.compile(r"[\s!-,:-@\[-^`{-~]|[-/]\d")


def _quote(s: str) -> str:
    if '"' not in s:
        return f'"{s}"'
    if "'" not in s:
        return f"'{s}'"
    return '"' + s.replace('"', '\\"') + '"'


def _number_text(x: Any) -> str:
    if isinstance(x, bool):
        return "是" if x else "否"
    if isinstance(x, numbers.Real) and float(x).is_integer():
        return str(int(x))
    return repr(x)


def _date_text(x: Any) -> str:
    ts = datetime(x.year, x.month, x.day) if not isinstance(x, datetime) else x
    text = ts.strftime("%Y-%m-%d") if (ts.hour, ts.minute, ts.second) == (0, 0, 0) else ts.strftime("%Y-%m-%d %H:%M:%S")
    return f'日期("{text}")'


class CellTranslator:
    """单元格条件 -> 表达式文本。知道的名字: 输入名前缀、派生变量、<参照表>.存在、基准日期。"""

    def __init__(self, rs: RuleSet) -> None:
        self.inputs = set(rs.names)
        self.derived = {k for k, _ in rs.derived}
        self.exists = {f"{s.name}.{EXIST_VAR}" for s in rs.lookups}
        self.known = list(self.derived) + list(FUNCS) + [BASE_DATE_VAR]

    def is_var(self, token: str) -> bool:
        if token in self.derived or token == BASE_DATE_VAR or token in self.exists:
            return True
        prefix, dot, rest = token.partition(".")
        return bool(dot and rest) and prefix in self.inputs

    def is_boolish(self, var: str) -> bool:
        """这些列上的 是 / 否 是布尔 (运行时 literal() 看的是列的 dtype, 导出时没有数据, 按这个约定走)。"""
        return var in self.exists or var in self.derived

    def name(self, var: str) -> str:
        if _NEEDS_BACKTICKS_RE.search(var) or any(
            re.search(rf"[-/]{re.escape(k)}(?![^\s!-/:-@\[-^`{{-~])", var) for k in self.known
        ):
            return f"`{var}`"
        return var

    def value(self, token: str, var: str) -> str:
        """比较右侧 / 枚举项 / 等于的字面值, 按 expressions.literal 的判定顺序翻译。"""
        t = token.strip()
        if len(t) >= 2 and t[0] == t[-1] and t[0] in "\"'":
            return _quote(t[1:-1])
        if t == NULL_WORD:
            return NULL_WORD
        if self.is_var(t):
            return self.name(t)
        if t in TRUE_WORDS or t in FALSE_WORDS:
            if self.is_boolish(var):
                return "是" if t in TRUE_WORDS else "否"
            return _quote(t)
        try:
            f = float(t)
            return _number_text(int(f) if f.is_integer() and "." not in t else f)
        except ValueError:
            pass
        if _DATE_TEXT_RE.match(t):
            return f'日期("{t}")'
        return _quote(t)

    def _is_plain(self, part: str) -> bool:
        """枚举项是不是单纯的字面值 (不是比较 / 区间 / 文本关键字 / 判空 / not)。"""
        p = part.strip()
        if p == NULL_WORD:
            return True  # 属于(x, "a", 空) / x not in ("a", 空) 都认 空
        return not (
            p in _NULL_TESTS or p.startswith("not(") or _RANGE_RE.match(p) or _CMP_RE.match(p) or _TEXT_TEST_RE.match(p)
        )

    def cell(self, var: str, cell: Any) -> str | None:
        """None = 不关心。分支顺序与 conditions.unary_test 一致。"""
        v = self.name(var)
        if cell is None:
            return None
        if isinstance(cell, bool) or isinstance(cell, numbers.Real):
            return f"{v} == {_number_text(cell)}"
        if isinstance(cell, (datetime, date)):
            return f"{v} == {_date_text(cell)}"
        s = str(cell).strip()
        if s in BLANK_WORDS:
            return None
        if s.startswith("not(") and s.endswith(")"):
            inner = s[4:-1].strip()
            if inner in (NULL_WORD, "为空"):
                return f"{v} != {NULL_WORD}"
            if inner == "非空":
                return f"{v} == {NULL_WORD}"
            parts = split_top_level(inner)
            if parts and all(self._is_plain(p) for p in parts):  # != / not in 本身就是空值不成立, 与 not(...) 一致
                if len(parts) == 1:
                    return f"{v} != {self.value(parts[0], var)}"
                return f"{v} not in ({', '.join(self.value(p, var) for p in parts)})"
            return f"{v} != {NULL_WORD} 且 非 ({self.cell(var, inner)})"
        if m := _TEXT_TEST_RE.match(s):
            args = [self.value(p, var) for p in split_top_level(m.group(2) or m.group(3))]
            return f"{m.group(1)}({v}, {', '.join(args)})"
        parts = split_top_level(s)
        if len(parts) > 1:
            if any(p in BLANK_WORDS for p in parts):
                return None  # 枚举里有「不关心」, 整格就是不关心
            if all(self._is_plain(p) for p in parts):
                return f"属于({v}, {', '.join(self.value(p, var) for p in parts)})"
            return " 或 ".join(str(self.cell(var, p)) for p in parts)
        if s == "非空":
            return f"{v} != {NULL_WORD}"
        if s in ("为空", NULL_WORD):
            return f"{v} == {NULL_WORD}"
        if m := _RANGE_RE.match(s):
            lo, hi = self.value(m.group(2), var), self.value(m.group(3), var)
            return f"{lo} {'<=' if m.group(1) == '[' else '<'} {v} {'<=' if m.group(4) == ']' else '<'} {hi}"
        if m := _CMP_RE.match(s):
            return f"{v} {_CMP_OPS[m.group(1)]} {self.value(m.group(2), var)}"
        return f"{v} == {self.value(s, var)}"


# ---------------------------------------------------------------------------
# Excel -> YAML
# ---------------------------------------------------------------------------


class FlowDict(dict):
    """dump 成 {a: 1, b: 2} 的一行写法 (输入声明、用例的输入 / 期望)。"""


class FlowList(list):
    """dump 成 [a, b, c] 的一行写法 (字段组的列名清单)。"""


class _Dumper(yaml.SafeDumper):
    pass


_Dumper.add_representer(FlowDict, lambda d, data: d.represent_mapping("tag:yaml.org,2002:map", data, flow_style=True))
_Dumper.add_representer(FlowList, lambda d, data: d.represent_sequence("tag:yaml.org,2002:seq", data, flow_style=True))


def _plain(v: Any) -> Any:
    """单元格值 -> YAML 值: 日期去掉 0 点的时间, 其余原样。"""
    if isinstance(v, datetime):
        return v.date() if (v.hour, v.minute, v.second) == (0, 0, 0) else v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _given(d: dict[str, Any]) -> FlowDict:
    """用例的输入 / 期望: 留空的键省略 (YAML 里不写 = 空值 / 不变, 与 Excel 留空同义)。"""
    return FlowDict({k: _plain(v) for k, v in d.items() if not _blank_value(v)})


def _blank_value(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _match_text(spec: Any, primary: str) -> str:
    if spec.primary:
        return PRIMARY_WORD
    return ", ".join(m if m == t else f"{m} = {primary}.{t}" for m, t in spec.keys)


def _rule_doc(tr: CellTranslator, r: RawRule) -> dict[str, Any]:
    doc: dict[str, Any] = {"编号": r.rule_id}
    if r.desc:
        doc["名称"] = r.desc
    doc["优先级"] = int(r.priority) if float(r.priority).is_integer() else r.priority
    branches = []
    for b in r.branches:
        exprs = [e for e in (tr.cell(var, cell) for var, cell in b.items()) if e is not None]
        branches.append(exprs)
    if len(branches) == 1:
        if branches[0]:
            doc["当"] = branches[0]
    else:
        doc["当"] = [{"任一满足": [b[0] if len(b) == 1 else b for b in branches]}]
    then = {col: _plain(v) for col, v in r.outputs.items() if v is not None and str(v).strip() not in BLANK_WORDS}
    if then:
        doc["则"] = then
    return doc


def excel_to_yaml_doc(path: str | Path) -> tuple[dict[str, Any], RuleSet]:
    rs, raw = load_excel_raw(Path(path))
    tr = CellTranslator(rs)
    primary = rs.primary.name
    doc: dict[str, Any] = {"输入": {}}
    for s in rs.inputs:
        entry: dict[str, Any] = {}
        if s.path:
            entry["文件"] = s.path
        entry["匹配"] = _match_text(s, primary)
        if not s.primary and s.dup_policy != "取第一条":
            entry["重复键"] = s.dup_policy
        doc["输入"][s.name] = FlowDict(entry)
    if rs.settings:
        doc["设置"] = {k: _plain(v) for k, v in rs.settings.items()}
    if rs.field_groups:
        doc[FIELD_GROUP_KEY] = {k: FlowList(v) for k, v in rs.field_groups.items()}
    if rs.derived:
        doc["变量"] = dict(rs.derived)
    doc["命中策略"] = rs.hit_policy
    doc["规则"] = [_rule_doc(tr, r) for r in raw]  # 输出格取原文: 字段组名下的格保持 组名: 输入名 的形态
    if rs.tests:
        doc["测试"] = [{"用例": t.name, "输入": _given(t.inputs), "期望": _given(t.expected)} for t in rs.tests]
    return doc, rs


def excel_to_yaml(path: str | Path) -> str:
    doc, rs = excel_to_yaml_doc(path)
    body = yaml.dump(doc, Dumper=_Dumper, allow_unicode=True, sort_keys=False, width=200)
    head = f"# 由 rule-sheet export 从 {Path(path).name} 生成, 请勿手改: 源是 Excel, 改了 Excel 后重新导出。\n\n"
    return head + body


def export_workbook(path: str | Path, out: str | Path | None = None) -> Path:
    """Excel 规则工作簿 -> YAML (默认同目录同名 .yaml); YAML -> Excel 规则工作簿 (默认同目录同名 .xlsx)。"""
    src = Path(path)
    suffix = src.suffix.lower()
    if suffix == ".xlsx":
        target = Path(out) if out else src.with_suffix(".yaml")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(excel_to_yaml(src), encoding="utf-8")
        return target
    if suffix in (".yaml", ".yml"):
        target = Path(out) if out else src.with_suffix(".xlsx")
        target.parent.mkdir(parents=True, exist_ok=True)
        wb = yaml_to_workbook(src)
        try:
            wb.save(target)
        except PermissionError as exc:
            raise PermissionError(f"无法写入 {target}: 文件可能正被 Excel 打开, 关闭后再试") from exc
        return target
    raise ValueError(f"export 只认 .xlsx (导出成 YAML) 或 .yaml (导出成 Excel): {src.name}")


# ---------------------------------------------------------------------------
# YAML -> Excel: 表格子集
# ---------------------------------------------------------------------------

UNARY_HINT = (
    "决策表单元格只放一元测试 (一个变量对字面量 / 变量的一次比较、区间、属于 / in、文本函数、判空、裸布尔名); "
    "其它写法先在 变量: 里起个名字, 条件里只引用它"
)
_FLIP = {ast.Lt: ast.Gt, ast.LtE: ast.GtE, ast.Gt: ast.Lt, ast.GtE: ast.LtE, ast.Eq: ast.Eq, ast.NotEq: ast.NotEq}
_OP_TEXT = {ast.Eq: "=", ast.NotEq: "!=", ast.Lt: "<", ast.LtE: "<=", ast.Gt: ">", ast.GtE: ">="}


def _is_number_text(t: str) -> bool:
    try:
        float(t)
        return True
    except ValueError:
        return False


def cell_literal(s: str) -> str:
    """字符串字面值写进单元格: 能裸写就裸写, 会被单元格文法认成别的东西的加引号。"""
    t = s.strip()
    needs = (
        t != s
        or not t
        or t in BLANK_WORDS
        or t in _NULL_TESTS
        or t == NULL_WORD
        or "," in t
        or t.startswith("not(")
        or t[0] in "\"'"
        or bool(_CMP_RE.match(t) or _RANGE_RE.match(t) or _TEXT_TEST_RE.match(t) or _DATE_TEXT_RE.match(t))
        or _is_number_text(t)
        or t in TRUE_WORDS
        or t in FALSE_WORDS
    )
    return _quote(s) if needs else s


class ExprToCell:
    """表达式 -> [(变量, 单元格文本)]; 不在表格子集里的写法报错。"""

    def __init__(self, rs: RuleSet) -> None:
        exists = {f"{s.name}.{EXIST_VAR}" for s in rs.lookups}
        derived = {k for k, _ in rs.derived}
        self.names = set(rs.variables) | derived | exists | {BASE_DATE_VAR}
        self.boolish = derived | exists
        self.used: dict[str, str] = {}
        self.where = ""
        self.expr = ""

    def _fail(self, why: str) -> ValueError:
        return ValueError(f"{self.where}: 条件 {self.expr!r} 转不成决策表单元格 ({why}); {UNARY_HINT}")

    def _var(self, node: ast.AST) -> str | None:
        return self.used.get(node.id) if isinstance(node, ast.Name) else None

    def _lit(self, node: ast.AST) -> str:
        """字面量 / 另一个变量名 / 空 / 日期("...") 在单元格里的写法。"""
        if isinstance(node, ast.Constant):
            v = node.value
            if isinstance(v, str):
                return cell_literal(v)
            if isinstance(v, bool) or not isinstance(v, numbers.Real):
                raise self._fail(f"不认识的字面量 {v!r}")
            return _number_text(v)
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.USub) and isinstance(node.operand, ast.Constant):
            return "-" + self._lit(node.operand)
        if isinstance(node, ast.Name):
            if node.id == NULL_WORD:
                return NULL_WORD
            if node.id in TRUE_WORDS or node.id in FALSE_WORDS:
                return "是" if node.id in TRUE_WORDS else "否"
            var = self._var(node)
            if var:
                return var
        if (
            isinstance(node, ast.Call)
            and isinstance(node.func, ast.Name)
            and node.func.id == "日期"
            and len(node.args) == 1
            and isinstance(node.args[0], ast.Constant)
            and isinstance(node.args[0].value, str)
        ):
            return node.args[0].value
        raise self._fail("右侧只能是字面量、另一个变量、空 或 日期(...)")

    def _compare(self, node: ast.Compare) -> tuple[str, str]:
        if len(node.ops) == 2:  # 链式区间 a <= x <= b / b >= x >= a
            lo, var_node, hi = node.left, node.comparators[0], node.comparators[1]
            var = self._var(var_node)
            ops = tuple(type(o) for o in node.ops)
            if var is None:
                raise self._fail("链式比较中间必须是变量")
            if all(o in (ast.Lt, ast.LtE) for o in ops):
                left, right = ops
            elif all(o in (ast.Gt, ast.GtE) for o in ops):
                lo, hi = hi, lo
                left, right = _FLIP[ops[1]], _FLIP[ops[0]]
            else:
                raise self._fail("链式比较只能是 a <= x <= b 这种区间")
            open_b = "[" if left is ast.LtE else "("
            close_b = "]" if right is ast.LtE else ")"
            return var, f"{open_b}{self._lit(lo)}..{self._lit(hi)}{close_b}"
        if len(node.ops) != 1:
            raise self._fail("链式比较只支持两段 (区间)")
        op, left, right = type(node.ops[0]), node.left, node.comparators[0]
        if op in (ast.In, ast.NotIn):
            var = self._var(left)
            if var is None or not isinstance(right, (ast.Tuple, ast.List)):
                raise self._fail("in 的左边必须是变量, 右边必须是字面量列表")
            items = ", ".join(self._lit(e) for e in right.elts)
            return var, items if op is ast.In else f"not({items})"
        var = self._var(left)
        other = right
        if var is None:
            var, other = self._var(right), left
            if var is None or op not in _FLIP:
                raise self._fail("比较的一侧必须是变量")
            op = _FLIP[op]
        if isinstance(other, ast.Name) and other.id == NULL_WORD:
            if op not in (ast.Eq, ast.NotEq):
                raise self._fail("空 只能和 == / != 搭配")
            return var, f"{_OP_TEXT[op]} {NULL_WORD}"
        if isinstance(other, ast.Name) and (other.id in TRUE_WORDS or other.id in FALSE_WORDS):
            if var not in self.boolish or op not in (ast.Eq, ast.NotEq):
                raise self._fail("和 是 / 否 比较只对 <参照表>.存在 与派生变量成立 (单元格里的 是 在文本列上是文字)")
            truthy = (other.id in TRUE_WORDS) == (op is ast.Eq)
            return var, "是" if truthy else "否"
        text = self._lit(other)
        if op is ast.Eq:
            return var, f"= {text}" if self._var(other) else text
        return var, f"{_OP_TEXT[op]} {text}"

    def _call(self, node: ast.Call) -> tuple[str, str]:
        fn = node.func.id if isinstance(node.func, ast.Name) else ""
        if not node.args or node.keywords:
            raise self._fail(f"函数 {fn} 的参数不对")
        var = self._var(node.args[0])
        if var is None:
            raise self._fail(f"函数 {fn} 的第一个参数必须是变量")
        rest = [self._lit(a) for a in node.args[1:]]
        if fn == "属于":
            return var, ", ".join(rest)
        if fn in ("为空", "非空") and not rest:
            return var, fn
        if fn in TEXT_FUNCS:
            return var, f"{fn} {', '.join(rest)}"
        raise self._fail(f"函数 {fn} 不能直接当单元格条件")

    def _is_null_test(self, node: ast.AST) -> tuple[str, bool] | None:
        """x == 空 / x != 空 / 为空(x) / 非空(x) -> (变量, 是否判空)。"""
        if isinstance(node, ast.Compare) and len(node.ops) == 1:
            left, right = node.left, node.comparators[0]
            if isinstance(right, ast.Name) and right.id == NULL_WORD and self._var(left):
                return self._var(left) or "", type(node.ops[0]) is ast.Eq
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in ("为空", "非空"):
            if len(node.args) == 1 and self._var(node.args[0]):
                return self._var(node.args[0]) or "", node.func.id == "为空"
        return None

    def _node(self, node: ast.AST) -> list[tuple[str, str]]:
        if isinstance(node, ast.BoolOp) and isinstance(node.op, ast.And):
            if len(node.values) == 2 and isinstance(node.values[1], ast.UnaryOp):
                nt = self._is_null_test(node.values[0])
                if nt and not nt[1] and isinstance(node.values[1].op, ast.Not):  # x != 空 且 非 (...) = 单元格 not(...)
                    inner = self._node(node.values[1].operand)
                    if len(inner) == 1 and inner[0][0] == nt[0]:
                        return [(nt[0], f"not({inner[0][1]})")]
            return [c for v in node.values for c in self._node(v)]
        if isinstance(node, ast.BoolOp) and isinstance(node.op, ast.Or):
            parts = [self._node(v) for v in node.values]
            if all(len(p) == 1 for p in parts) and len({p[0][0] for p in parts}) == 1:
                return [(parts[0][0][0], ", ".join(p[0][1] for p in parts))]  # 同一变量的 或 = 枚举
            raise self._fail("或 的两侧是不同变量; 跨变量的或写成规则顶层的 任一满足:, 转成决策表的续行")
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, ast.Not):
            var = self._var(node.operand)
            if var and var in self.boolish:
                return [(var, "否")]
            nt = self._is_null_test(node.operand)
            if nt:
                return [(nt[0], "非空" if nt[1] else "为空")]
            raise self._fail(
                "非 是纯取反 (空值行也为真), 单元格的 not(...) 只对有值的格成立, 两者不等价; "
                "写成 x != 值 / x not in (...), 或 x != 空 且 非 (...)"
            )
        if isinstance(node, ast.Name):
            var = self._var(node)
            if var and var in self.boolish:
                return [(var, "是")]
            raise self._fail("裸变量名当条件只对 <参照表>.存在 与派生变量成立")
        if isinstance(node, ast.Compare):
            return [self._compare(node)]
        if isinstance(node, ast.Call):
            return [self._call(node)]
        raise self._fail("不是一元测试")

    def cells(self, expr: str, where: str) -> list[tuple[str, str]]:
        """一条表达式 -> [(变量, 单元格文本)]; 且 连接的几个一元测试拆成几格。"""
        self.where, self.expr = where, expr
        text, self.used = substitute_names(expr, self.names)
        try:
            tree = ast.parse(text, mode="eval").body
        except SyntaxError as exc:
            raise ValueError(f"{where}: 表达式语法错误 {expr!r} ({exc.msg})") from exc
        return self._node(tree)

    def _flat_all(self, node: Any, where: str) -> list[tuple[str, str]]:
        if isinstance(node, str):
            return self.cells(node, where)
        if isinstance(node, list):
            return [c for n in node for c in self._flat_all(n, where)]
        if isinstance(node, dict) and len(node) == 1:
            key, val = next(iter(node.items()))
            if key == "全部满足":
                return self._flat_all(val, where)
            if key == "不满足":
                inner = self._flat_all(val, where)
                if len(inner) == 1 and inner[0][1] in ("为空", "非空", f"= {NULL_WORD}", f"!= {NULL_WORD}"):
                    var, cell = inner[0]
                    return [(var, "非空" if cell in ("为空", f"= {NULL_WORD}") else "为空")]
                raise ValueError(
                    f"{where}: 不满足: 是纯取反 (空值行也为真), 只有包着判空时才能转成决策表; "
                    f"其它情况写成 x != 值 / x not in (...), 或放进变量定义"
                )
            raise ValueError(
                f"{where}: 嵌套在条件里面的 {key}: 转不成决策表 (任一满足 只能放在规则条件的顶层, 对应续行)"
            )
        raise ValueError(f"{where}: 无法识别的条件写法: {node!r}")

    @staticmethod
    def _as_range(a: str, b: str) -> str | None:
        """同一变量的两个比较 (下界 + 上界) 合成区间格。"""
        ma, mb = _CMP_RE.match(a), _CMP_RE.match(b)
        if not ma or not mb:
            return None
        lo = next(((m.group(1), m.group(2)) for m in (ma, mb) if m.group(1) in (">", ">=")), None)
        hi = next(((m.group(1), m.group(2)) for m in (ma, mb) if m.group(1) in ("<", "<=")), None)
        if not lo or not hi:
            return None
        return f"{'[' if lo[0] == '>=' else '('}{lo[1].strip()}..{hi[1].strip()}{']' if hi[0] == '<=' else ')'}"

    def _merge(self, cells: list[tuple[str, str]], where: str) -> dict[str, str]:
        out: dict[str, str] = {}
        for var, cell in cells:
            if var in out:
                merged = self._as_range(out[var], cell)
                if merged is None:
                    raise ValueError(
                        f"{where}: 变量「{var}」在同一条规则里有两个条件 ({out[var]} 与 {cell}), "
                        f"决策表一列只放一格; 写成区间 / 枚举, 或放进变量定义"
                    )
                out[var] = merged
            else:
                out[var] = cell
        return out

    def branches(self, cond: Any, where: str) -> list[dict[str, str]]:
        """规则的 当: -> 若干分支 (每个分支: 变量 -> 单元格文本); 顶层 任一满足 的每一支与其余条件各组成一行。"""
        items = cond if isinstance(cond, list) else [cond]
        fixed: list[tuple[str, str]] = []
        alternatives: list[list[tuple[str, str]]] | None = None
        for item in items:
            if isinstance(item, dict) and len(item) == 1 and next(iter(item)) == "任一满足":
                if alternatives is not None:
                    raise ValueError(f"{where}: 一条规则只能有一个顶层 任一满足 (对应续行)")
                val = next(iter(item.values()))
                alternatives = [self._flat_all(alt, where) for alt in (val if isinstance(val, list) else [val])]
            else:
                fixed.extend(self._flat_all(item, where))
        rows = [fixed] if alternatives is None else [fixed + alt for alt in alternatives]
        return [self._merge(r, where) for r in rows]


def _excel_value(v: Any) -> Any:
    if isinstance(v, date) and not isinstance(v, datetime):
        return datetime(v.year, v.month, v.day)
    return v


def yaml_to_workbook(path: str | Path) -> Workbook:
    """YAML 规则文件 -> 规则工作簿 (五张 sheet, 版式按 style 规范套好)。条件不在表格子集里的报错。"""
    src = Path(path)
    doc = yaml.safe_load(src.read_text(encoding="utf-8")) or {}
    rs = load_yaml_rules(src)
    conv = ExprToCell(rs)
    primary = rs.primary.name
    columns: list[str] = []
    headers: list[str] = []  # 决策表的输出列按 则: 原文 (字段组名保留为一列); 测试 sheet 的期望列是展开后的 rs.outputs
    parsed: list[tuple[str, Any, list[dict[str, str]], dict[str, Any], str]] = []
    for order, r in enumerate(doc.get("规则") or []):
        rid = str(r.get("编号", f"R{order + 1}"))
        branches = conv.branches(r.get("当") or [], f"规则 {rid}")
        for b in branches:
            columns.extend(v for v in b if v not in columns)
        then = {str(k): v for k, v in (r.get("则") or {}).items()}
        headers.extend(k for k in then if k not in headers)
        parsed.append((rid, r.get("优先级", 1000 + order), branches, then, str(r.get("名称", ""))))

    wb = Workbook()
    ws = wb.active
    ws.title = "输入"
    ws.append(["名称", "匹配", "重复键", "文件", "说明"])
    for s in rs.inputs:
        ws.append([s.name, _match_text(s, primary), None if s.primary else s.dup_policy, s.path, None])
    if rs.settings:
        ws = wb.create_sheet("设置")
        ws.append(["设置项", "值", "说明"])
        for k, v in rs.settings.items():
            ws.append([k, _excel_value(v), None])
    if rs.derived:
        ws = wb.create_sheet("变量定义")
        ws.append(["变量", "定义", "说明"])
        for k, v in rs.derived:
            ws.append([k, v, None])
    if rs.field_groups:
        ws = wb.create_sheet(FIELD_GROUP_KEY)
        ws.append(["名称", "列名", "说明"])
        for k, cols in rs.field_groups.items():
            ws.append([k, ", ".join(cols), None])

    ws = wb.create_sheet("决策表")
    ws.append(["命中策略", rs.hit_policy])
    ws.append(["规则ID", "优先级", *["输入"] * len(columns), *["输出"] * len(headers), "说明"])
    ws.append(["规则ID", "优先级", *columns, *headers, "说明"])
    for rid, priority, branches, then, desc in parsed:
        first = ws.max_row + 1
        for i, b in enumerate(branches):
            cells = [b.get(c) for c in columns]
            if i == 0:
                ws.append([rid, priority, *cells, *[_excel_value(then.get(o)) for o in headers], desc or None])
            else:
                ws.append([None, None, *cells, *[None] * len(headers), None])
        if len(branches) > 1:
            ws.merge_cells(start_row=first, start_column=1, end_row=first + len(branches) - 1, end_column=1)
            ws.merge_cells(start_row=first, start_column=2, end_row=first + len(branches) - 1, end_column=2)

    if rs.tests:
        keys: list[str] = []
        for t in rs.tests:
            keys.extend(k for k in t.inputs if k not in keys)
        ws = wb.create_sheet("测试")
        ws.append(["用例", *["输入"] * len(keys), *["期望"] * len(rs.outputs), "说明"])
        ws.append(["用例", *keys, *rs.outputs, "说明"])
        for t in rs.tests:
            row = [t.name, *[_excel_value(t.inputs.get(k)) for k in keys]]
            ws.append([*row, *[_excel_value(t.expected.get(o)) for o in rs.outputs], None])
    style_rules_workbook(wb)
    return wb


# ---------------------------------------------------------------------------
# 两份规则文件的语义比对
# ---------------------------------------------------------------------------


@dataclass
class DiffResult:
    differences: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)  # 路径之类不算差异的提示
    cases: int = 0  # 交叉验证用了多少条用例

    @property
    def same(self) -> bool:
        return not self.differences


def _norm_cell(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return None if s in BLANK_WORDS else s
    if isinstance(v, datetime):
        return _plain(v)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _norm_given(d: dict[str, Any]) -> dict[str, Any]:
    out = {}
    for k, v in d.items():
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        out[k] = _plain(v) if isinstance(v, (datetime, float)) else v
    return out


def _case_key(name: str) -> str:
    """用例按首个词 (用例 ID) 配对: Excel 的「用例」列只写 T1, YAML 习惯写 "T1 偏差 8% 触发 R1"。"""
    return name.strip().split()[0] if name.strip() else name


def _structure_diff(a: RuleSet, b: RuleSet, res: DiffResult) -> None:
    ia = {s.name: s for s in a.inputs}
    ib = {s.name: s for s in b.inputs}
    for n in sorted(set(ia) | set(ib)):
        if n not in ia or n not in ib:
            res.differences.append(f"输入「{n}」只在 {'B' if n not in ia else 'A'} 里声明")
            continue
        sa, sb = ia[n], ib[n]
        if (sa.primary, sa.keys, sa.dup_policy) != (sb.primary, sb.keys, sb.dup_policy):
            ka, kb = sa.keys or PRIMARY_WORD, sb.keys or PRIMARY_WORD
            res.differences.append(f"输入「{n}」: A 匹配 {ka} / {sa.dup_policy}, B {kb} / {sb.dup_policy}")
        if sa.path != sb.path:
            res.notes.append(f"输入「{n}」的默认文件不同: A {sa.path}, B {sb.path}")
    for k in sorted(set(a.settings) | set(b.settings)):
        va, vb = _norm_cell(a.settings.get(k)), _norm_cell(b.settings.get(k))
        if va != vb:
            (res.notes if k == "输出" else res.differences).append(f"设置「{k}」: A {va}, B {vb}")
    if a.derived != b.derived:
        res.differences.append(f"变量定义不同: A {a.derived}, B {b.derived}")
    if a.field_groups != b.field_groups:
        res.differences.append(f"字段组不同: A {a.field_groups}, B {b.field_groups}")
    if a.hit_policy != b.hit_policy:
        res.differences.append(f"命中策略: A {a.hit_policy}, B {b.hit_policy}")
    if set(a.outputs) != set(b.outputs):
        res.differences.append(f"输出列: A {a.outputs}, B {b.outputs}")
    ra = {r.rule_id: r for r in a.rules}
    rb = {r.rule_id: r for r in b.rules}
    for rid in [r.rule_id for r in a.rules] + [r.rule_id for r in b.rules if r.rule_id not in ra]:
        if rid not in ra or rid not in rb:
            res.differences.append(f"规则 {rid} 只在 {'B' if rid not in ra else 'A'} 里")
            continue
        if ra[rid].priority != rb[rid].priority:
            res.differences.append(f"规则 {rid} 优先级: A {ra[rid].priority:g}, B {rb[rid].priority:g}")
        oa = {c: _norm_cell(v) for c, v in ra[rid].outputs.items() if _norm_cell(v) is not None}
        ob = {c: _norm_cell(v) for c, v in rb[rid].outputs.items() if _norm_cell(v) is not None}
        if oa != ob:
            res.differences.append(f"规则 {rid} 输出: A {oa}, B {ob}")
    ta = {_case_key(t.name): t for t in a.tests}
    tb = {_case_key(t.name): t for t in b.tests}
    keys = [_case_key(t.name) for t in a.tests] + [k for k in (_case_key(t.name) for t in b.tests) if k not in ta]
    for key in keys:
        if key not in ta or key not in tb:
            res.differences.append(f"用例「{key}」只在 {'B' if key not in ta else 'A'} 里")
            continue
        ga, gb = _norm_given(ta[key].inputs), _norm_given(tb[key].inputs)
        if ga != gb:
            res.differences.append(f"用例「{key}」输入不同: A {ga}, B {gb}")
        ea, eb = _norm_given(ta[key].expected), _norm_given(tb[key].expected)
        if ea != eb:
            res.differences.append(f"用例「{key}」期望不同: A {ea}, B {eb}")


def _run_case(rs: RuleSet, tc: TestCase) -> tuple[dict[str, Any], str] | str:
    ns, originals = case_namespace(rs, tc)
    try:
        results, hits = apply_rules(rs, ns, originals)
    except Exception as exc:  # noqa: BLE001
        return str(exc)
    return {col: results[col].iloc[0] for col in rs.outputs}, hits.iloc[0]


def _behavior_diff(a: RuleSet, b: RuleSet, res: DiffResult) -> None:
    """两边的用例合并, 每条在两套规则上各跑一次, 比实际输出与命中 ID (不是比通过与否)。"""
    seen: set[tuple[str, str]] = set()
    cols = [c for c in a.outputs if c in b.outputs]
    for tc in a.tests + b.tests:
        key = (_case_key(tc.name), repr(sorted(_norm_given(tc.inputs).items())))
        if key in seen:
            continue  # 两边同名同输入的用例只跑一次
        seen.add(key)
        res.cases += 1
        oa, ob = _run_case(a, tc), _run_case(b, tc)
        name = _case_key(tc.name)
        if isinstance(oa, str) or isinstance(ob, str):
            ea = oa if isinstance(oa, str) else "正常"
            eb = ob if isinstance(ob, str) else "正常"
            res.differences.append(f"用例「{name}」出错: A {ea}; B {eb}")
            continue
        (va, ha), (vb, hb) = oa, ob
        for col in cols:
            if not same(va[col], vb[col]):
                res.differences.append(f"用例「{name}」输出「{col}」: A {fmt(va[col])}, B {fmt(vb[col])}")
        if ha != hb:
            res.differences.append(f"用例「{name}」命中规则: A {ha or '无'}, B {hb or '无'}")


def diff_rules(path_a: str | Path, path_b: str | Path) -> DiffResult:
    a, b = load_rules(path_a), load_rules(path_b)
    res = DiffResult()
    _structure_diff(a, b, res)
    _behavior_diff(a, b, res)
    return res
